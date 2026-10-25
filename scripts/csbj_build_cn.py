"""Build the Chinese edition (Manuscript_CSBJ_zh.docx) of the CSBJ v2 resubmission.

Reuses the language-neutral machinery of csbj_build_package_v2 (OMML equations,
table builder, reference formatting, figure files); only text is Chinese.
Output: papers/13_CSBJ_v2_20261014/Manuscript_CSBJ_zh.docx (+ _resolved.md).
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import csbj_build_package_v2 as V

O=V.O;F=V.F;REF=V.REF;EQS=V.EQS;element=V.element;table=V.table;reftext=V.reftext

def _east(st,name):
    rpr=st.element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None:
        rf=element('w:rFonts');rpr.append(rf)
    rf.set(qn('w:eastAsia'),name)

def setup_cn():
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
    s.header_distance=s.footer_distance=Inches(.492)
    for name,size,bef,aft,ea in [('Normal',11,0,6,'宋体'),('Title',17,0,12,'黑体'),('Heading 1',14,16,8,'黑体'),('Heading 2',12,12,6,'黑体'),('Heading 3',11,8,4,'黑体'),('Caption',9,4,6,'宋体')]:
        st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.font.bold=name.startswith('Heading') or name=='Title'
        st.paragraph_format.space_before=Pt(bef);st.paragraph_format.space_after=Pt(aft)
        st.paragraph_format.line_spacing=1.25 if name=='Normal' else 1.0
        st.paragraph_format.widow_control=True
        if name.startswith('Heading'):st.paragraph_format.keep_with_next=True
        _east(st,ea)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('MRP-ST  |  ');footer._p.append(element('w:fldSimple',**{'w:instr':'PAGE'}))
    return d

# ---- Chinese tables (numbers identical to the English edition) ----
T1H=['变体','相对 FULL 的改动','保留的信息']
T1R=[['FULL','无移除','有序 + 代理 + 对齐；图像；邻域'],['NOIMAGE','图像表征置零','保留全部损失'],['NOHPV','lambda_hpv = 0','保留有序与代理目标'],['NONEIGH','lambda_neigh = 0','保留其他损失与图像'],['NOPROG','lambda_prog = 0','保留代理与对齐'],['NOSUP','lambda_prog = lambda_sup = 0','保留对齐；无显式目标'],['NOALL','lambda_prog = lambda_sup = lambda_hpv = 0','重建、图像、邻域、方差下限']]
T1T='表 1. 显式消融设计。每行与 FULL 对比。'
T2H=['外部分析','观察到的证据','声明边界']
T2R=[['GeoMx，6 例配对患者','肿瘤 − 良性均值 −2.218；精确 p=0.03125','配对分离；与来源队列评分方向相反'],['GeoMx，Silva 模式','核心均值非单调；7 例患者','无经验证的有序分级或校准'],['HNSCC，8 / 6 张切片','平均 Moran’s I 0.757 / 0.667','空间组织；非疾病判别'],['HNSCC，HPV 状态','样本级 p=0.200 / 0.133','两个队列均无令人信服的状态分离'],['TCGA-CESC，302 例患者','11 基因筛查；均不显著','仅为基因级筛查；非模型级预后']]
T2T='表 5. 外部证据及其支持的解读边界。'
T3H=['判级','判据（条款 C4：跨种子 LODO 探针 tau-b 均值）','解释']
T3R=[['支持','mean τ ≥ 0.35 且 ≥ 2/3 种子 p < 0.05','无标签表征中可复现的信息'],['提示','既非支持亦非不支持','信息弱或种子不稳定；以连续估计为准'],['不支持','|mean τ| ≤ 0.15 且 0/3 种子 p < 0.05','该样本量下证据不足，并非生物学缺失']]
T3T='表 2. 应用于条款 C4（表征信息）的三级定级判据。'
T4H=['水平（f）','主尺度 τ̄ (SD)','显著 run','判级','匹配尺度 τ̄ (SD)','显著 run','判级']
T4R=[['L1 (f = 0)','0.704 (0.009)','9/9','支持','0.457 (0.000)','9/9','支持'],['L2 (f = 0.25)','0.276 (0.154)','6/9','提示','0.394 (0.095)','6/9','支持'],['L3 (f = 0.50)','0.125 (0.122)','1/9','提示','0.286 (0.096)','2/9','提示'],['L4 (f = 0.75)','0.182 (0.104)','3/9','提示','0.177 (0.148)','2/9','提示'],['CTRL（身份混杂）','−0.030 (0.111)','0/9','—','0.042 (0.102)','0/9','—']]
T4T='表 3. C4 留出探针在剂量–反应网格上的操作特性。每格九个 run（3 个生成种子 × 3 个训练种子）。f 为非应答者比例。CTRL 行的显著 run 比例即 α=0.05 下的经验假阳性率（合并 0/18；精确 95% CI 0–18.5%，与名义 5% 相容）。二分类标签上界：τ_max = 0.707（主尺度）与 0.457（匹配尺度）；L1 均值分别达到上界的 99.6% 与 100%。'
T5H=['队列（21 样本子集）','SpaGCN τ (p)','STAGATE τ (p)','GraphST τ (p)','参考探针 τ̄（显著种子）']
T5R=[['L1，90 样本','0.677 (0.0005) *','−0.173 (0.83)','0.491 (0.007) *','0.717 (3/3) *'],['L3，90 样本','0.093 (0.33)','0.173 (0.20)','0.213 (0.14)','0.243 (0/3)'],['L4，90 样本','0.080 (0.36)','−0.465 (0.99)','0.226 (0.14)','0.155 (0/3)'],['L1，35 样本','0.291 (0.073)','−0.045 (0.61)','0.425 (0.005) *','0.425 (3/3) *'],['L3，35 样本','0.291 (0.073)','−0.381 (0.99)','0.336 (0.050)','−0.201 (0/3)'],['L4，35 样本','0.291 (0.073)','−0.336 (0.96)','0.336 (0.050)','−0.194 (0/3)']]
T5T='表 4. 相同 21 样本子集上的统一审计（主尺度 12 正常 + 9 患病；匹配尺度 19 + 2）。已发表方法：无监督嵌入的簇最大值 HPV 通路读出，单次 run，2,000 次标签置换。参考探针：冻结 NOALL 嵌入上的岭回归探针，3 个训练种子（均值 τ̄；α=0.05 下显著种子数），重拟合置换。* p < 0.05。子集二分类标签上界：τ_max = 0.717 与 0.425。支持级仅在 L1 出现（90 样本：SpaGCN 与 GraphST；两尺度：GraphST 与参考探针）。'

CAPTIONS=[
('Fig1_Audit_workflow','图 1. MRP-ST 工作流程。独立的消融对比量化标量评分依赖；留出探针评估冻结 NOALL 嵌入中的信息。样本-标签置换重拟合完整探针。带身份混杂对照的剂量–反应网格测量定级规则的误差率，统一审计将相同判据应用于已发表嵌入方法。外部终点分别评估。编码在协变量上保持转导；宫颈疾病提供演示而非临床验证。'),
('Fig2_Operating_characteristics','图 2. C4 定级规则在剂量–反应网格上的实测操作特性。A：90 样本（圆点）与宫颈匹配 35 样本（方点）尺度下留出探针 tau-b 均值（± SD，九个 run）随非应答者比例 f 的变化；身份混杂对照绘于右侧；点线为二分类标签上界 tau_max（0.707 与 0.457）。B：重拟合 p < 0.05 的 run 比例；点线为名义 α = 0.05，对照柱给出经验假阳性率（0/9 与 0/9）。C：灵敏度（L1 队列）与假阳性率（CTRL 队列）随支持级阈值 θ 的变化；底纹为完美分离窗口 [0.20, 0.675]（主尺度）与 [0.20, 0.45]（匹配尺度）；预设阈值 0.35 位于两个窗口内。'),
('Fig3_Dissociation','图 3. 对照臂解离评分头记忆化、编码器泄漏与表征信息。所有面板使用身份混杂 CTRL 队列，其标签在设计上与共享表达结构独立。A：90 样本尺度。B：宫颈匹配 35 样本尺度。柱为九个 run 的均值；点为单个 run。每尺度三个读出：FULL 头样本内 tau-b（监督拟合；不适用检验）、FULL 嵌入留一探针 tau-b（标签训练转导编码器；全部 p ≤ 1e-4）、NOALL 嵌入留一探针 tau-b（无标签冻结编码器；实心点 p < 0.05）。虚线：二分类标签上界 tau_max（0.707 与 0.457）。只有无标签编码器返回正确的零行为（−0.030 ± 0.111 与 +0.042 ± 0.102）；与标签训练嵌入之间 0.73 与 0.41 个 tau 单位的差距即条款 C4 排除的编码器泄漏量。'),
('Fig4_Methods_benchmark','图 4. 统一审计在相同 21 样本子集上对已发表无监督嵌入方法与参考协议定级。A：90 样本队列。B：宫颈匹配 35 样本队列。每面板显示样本级风险读出与严重程度的 Kendall tau-b：SpaGCN、STAGATE 与 GraphST（单次 run；簇最大值 HPV 通路读出；2,000 次标签置换；实心标记 p < 0.05），以及参考协议在冻结 NOALL 嵌入上的岭回归探针（3 个训练种子的均值 ± SD；重拟合置换；实心菱形：至少一个种子 p < 0.05）。虚线：子集二分类标签上界（0.717 与 0.425）。支持级仅在最强水平 L1 出现；没有任何读出将弱水平与零假设区分开，SpaGCN 的匹配尺度读出在三个水平返回相同 tau（读出压缩）。'),
('Fig5_Cervical_audit','图 5. 宫颈评分依赖与表征信息。A：原始 NPZ 评分按样本池化后与严重程度比较；点为训练种子，柱为均值 ± 样本 SD。NOPROG 有四个种子，其他变体三个。种子集合未完全配对。B：NOALL 留出探针相关及 10,000 次重拟合标签置换的单侧 p 值。三个种子中一个显著，且经 21 项检验 BH 校正后仍如此。'),
('Fig6_Biological_orientation','图 6. 信息移除间的描述性生物学关联。FULL、NOPROG 与 NOSUP 列显示均值相关；NOALL 列展示单种子以暴露符号不稳定。颜色为存储的评分–程序 Spearman 相关。这些是表达衍生的关联，不是独立的机制验证；未做事后符号修正。'),
('Fig7_External_evidence','图 7. 外部性质不可互换。A：FINAL 检查点 GeoMx 评分，在 6 例配对患者的组织类型内聚合。全部肿瘤−良性差为负；精确符号秩 p=0.03125。B：两个头颈部队列的存储平均 Moran’s I，附独立的样本级 HPV 状态检验。两个坐标轴测量不同终点，不合并为验证分。'),
('Fig8_LDI_calibration','图 8. 标签依赖指数校准与门限行为（条款 C3）。A：10,000 次完全样本-标签置换下监督评分的严重度排序分布，标注观测 tau 与 95 分位门限 θ；观测值位于上尾深处。B：注入标签依赖的恢复随注入比例 α 的函数，显示单调校准以及模拟依赖与观测排序匹配的 α。C：七个监督移除变体在 tau 尺度上的均值与跨种子 SD，同一 θ 线。正文报告的 LDI 是该超出量按零分布 SD 的标准化（LDI=3.26 对 θ=1.645）；不使用原始 tau_FULL − tau_NOALL 对比，因为 NOALL 头坍缩为数值不确定的偏移。'),
('Fig9_Local_controls','图 9. 标签依赖、泄漏与阈值选择的局部对照。A：门限 θ 随 α 变化并显示置换种子范围，对照观测 tau_FULL；标签依赖判定在 α ∈ {0.01, 0.05, 0.10, 0.20} 与五个置换种子上稳定。B：七个监督移除变体的均分–严重度 tau 均值与种子 SD（无标签变体为红色）。C：各变体的留一供体外探针相关；用供体标签训练的编码器返回 0.42–0.46 且全部种子显著，无标签编码器返回 0.270 且三个种子中一个显著。面板 B 与 C 将记忆化差距与可迁移信号分开。')]

def build_main_cn():
    text=(O/'Manuscript_CSBJ_zh.md').read_text(encoding='utf-8-sig');order=[]
    def cite(m):
        ks=[k.strip().lstrip('@') for k in m.group(1).split(';')]
        for k in ks:
            if k not in order:order.append(k)
        return '['+', '.join(str(order.index(k)+1) for k in ks)+']'
    text=re.sub(r'\[@([^\]]+)\]',cite,text)
    text=re.sub(r'(?m)^(#{1,3} [^\n]+)\n(?!\n)',r'\1\n\n',text)
    d=setup_cn();eq=0
    for b in text.split('\n\n'):
        b=b.strip()
        if not b:continue
        if b=='[[TABLE1]]':table(d,T1H,T1R,[1450,2900,5010],T1T);continue
        if b=='[[TABLE2]]':table(d,T2H,T2R,[2350,3200,3810],T2T);continue
        if b=='[[TABLE3]]':table(d,T3H,T3R,[1300,3900,4160],T3T);continue
        if b=='[[TABLE4]]':table(d,T4H,T4R,[2050,1500,850,1150,1500,850,1460],T4T);continue
        if b=='[[TABLE5]]':table(d,T5H,T5R,[2200,1800,1800,1800,1760],T5T);continue
        if b=='[[REFERENCES]]':
            for i,k in enumerate(order,1):
                p=d.add_paragraph(f'[{i}] '+reftext(k));p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(5)
                for r in p.runs:r.font.size=Pt(9)
            continue
        if b=='[[FIGURES]]':
            for name,cap in CAPTIONS:
                d.add_page_break();p=d.add_paragraph();p.paragraph_format.keep_with_next=True
                p.add_run().add_picture(str(F/(name+'.png')),width=Inches(6.45));d.add_paragraph(cap,style='Caption')
            continue
        if b.startswith('$$'):
            for j,parts in enumerate(EQS[eq]):
                p=d.add_paragraph();p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(6)
                om=element('m:oMath')
                for part in parts:om.append(part)
                p._p.append(om)
                if j==len(EQS[eq])-1:p.add_run(f'    ({eq+1})').font.size=Pt(10)
            eq+=1;continue
        if b.startswith('### '):d.add_heading(b[4:],2)
        elif b.startswith('## '):d.add_heading(b[3:],1)
        elif b.startswith('# '):d.add_paragraph(b[2:],style='Title')
        else:d.add_paragraph(b)
    try:d.save(O/'Manuscript_CSBJ_zh.docx')
    except PermissionError:
        d.save(O/'Manuscript_CSBJ_zh_new.docx');print('WARN: Manuscript_CSBJ_zh.docx locked by Word; saved as Manuscript_CSBJ_zh_new.docx')
    refs='\n\n'.join(f'[{i}] '+reftext(k) for i,k in enumerate(order,1))
    (O/'References_zh.txt').write_text(refs,encoding='utf-8')
    resolved=text.replace('[[REFERENCES]]',refs).replace('[[FIGURES]]','\n\n'.join(c for _,c in CAPTIONS))
    for marker,head,rows in [('[[TABLE1]]',T1H,T1R),('[[TABLE2]]',T2H,T2R),('[[TABLE3]]',T3H,T3R),('[[TABLE4]]',T4H,T4R),('[[TABLE5]]',T5H,T5R)]:
        resolved=resolved.replace(marker,' | '.join(head)+'\n'+'\n'.join(' | '.join(r) for r in rows))
    (O/'Manuscript_CSBJ_zh_resolved.md').write_text(resolved,encoding='utf-8')
    n_cn=len(re.findall(r'[\u4e00-\u9fff]',resolved));n_en=len(re.findall(r'[A-Za-z]+',resolved))
    return n_cn,n_en,len(order),eq

if __name__=='__main__':
    cn,en,rf,eq=build_main_cn()
    print('CN chars=%d EN words=%d refs=%d eqs=%d'%(cn,en,rf,eq))
