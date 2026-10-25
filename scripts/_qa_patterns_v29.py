from pathlib import Path
import sys,re,json,math,hashlib
from docx import Document
sys.stdout.reconfigure(encoding='utf-8')
P=Path('papers/20_MRP-ST-Repair_Patterns');Q=Path('.qa/patterns_v29')
src=(P/'MRP-ST-Repair_manuscript_Patterns_v29.md').read_text(encoding='utf-8')
d=Document(P/'Manuscript_Patterns_v29.docx');body='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for tab in d.tables for row in tab.rows for c in row.cells)
checks=[]
def ck(name,ok):checks.append({'check':name,'pass':bool(ok)});print('PASS' if ok else 'FAIL',name)
ck('preserve 5 embedded figures',len(d.inline_shapes)==5)
ck('preserve 7 main tables',len(d.tables)==7)
nums=[int(n) for n in re.findall(r'^(\d+)\. ',src.split('## REFERENCES')[1],re.M)]
ck('44 contiguous references including archive',nums==list(range(1,45)))
for n in ['10.5281/zenodo.23640979','pathway-anchored-risk-benchmark','73ad0f1710953651c43055afc693f7a18ce93c6f','not contain','Corrected train-only preprocessing check.']:
 ck('manuscript contains '+n,n in body)
ck('remove unverified replacement repository from manuscript','github.com/WorldBMW/mrp-st-audit' not in body)
ck('remove false full-reproducibility promise','Every number in this paper is reproducible from the released manifests' not in body)
ck('correct code-assisted AI declaration','assisted in drafting and checking analysis scripts' in body and 'Neither tool generated core algorithmic logic' not in body)
sm=src.split('## SUMMARY')[1].split('**Keywords:')[0];ib=src.split('## IN BRIEF')[1].split('## SUMMARY')[0]
hl=re.findall(r'^\d+\. (.+)$',src.split('## HIGHLIGHTS')[1].split('## IN BRIEF')[0],re.M)
ck('Summary <=150 words',len(sm.split())<=150);ck('In brief <=50 words',len(ib.split())<=50)
ck('4 highlights <=85 chars',len(hl)==4 and max(map(len,hl))<=85)
for tag in ['CoverLetter','TitlePage','Supplemental_Information']:
 x=Document(P/f'{tag}_Patterns_v29.docx');full='\n'.join(p.text for p in x.paragraphs)
 ck(tag+' DOI present','10.5281/zenodo.23640979' in full)
a=json.loads(Path('results/outer_refit_v29.json').read_text(encoding='utf-8'))
ck('outer analysis fitted preprocessing in training fold','solely on outer-training' in a['protocol']['preprocessing'])
for r in a['runs']:
 ck('outer dose '+str(r['lam'])+' 30 unique group holdouts',len({f['group'] for f in r['folds']})==30 and sum(f['n_test'] for f in r['folds'])==r['m_units'])
 ck('outer dose '+str(r['lam'])+' matches manuscript rounded results',f"{r['pooled_tau_frozen']:.4f}" in body and f"{r['pooled_tau_outer']:.4f}" in body)
zip_paths=(Q/'zenodo_zip_paths.txt').read_text(encoding='utf-8').splitlines()
ck('frozen public ZIP really lacks current audit runner',not any(p.endswith('exp_regate_v28.py') for p in zip_paths))
out={'checks':checks,'summary_words':len(sm.split()),'in_brief_words':len(ib.split()),'highlight_chars':list(map(len,hl)),'all_pass':all(c['pass'] for c in checks),'scope':'Document integrity, archive coverage and new-run consistency; not acceptance certification or proof of every scientific claim.'}
(P/'AUDIT_v29_checks.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
if not out['all_pass']:raise SystemExit(1)
