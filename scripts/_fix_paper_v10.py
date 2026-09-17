# -*- coding: utf-8 -*-
"""Manuscript_CMPB_v9_中文版.docx -> v10 数据修正版
所有修正均依据 results/ 下真实 JSON 核对结果。
"""
import docx

DST = r"F:\project4_cervix_multimodal\papers\06_CMPB_最终版_v2_20260905\Manuscript_CMPB_v10_中文版_数据修正.docx"

d = docx.Document(DST)
log = []


def set_para_text(p, new):
    """段落级重写：保留首个 run 的字符格式，清空其余 run"""
    if not p.runs:
        return False
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""
    return True


def find_one(key):
    hits = [i for i, p in enumerate(d.paragraphs) if key in p.text]
    if not hits:
        print("  !! 未找到:", key[:40])
        return None
    if len(hits) > 1:
        print("  (多处命中 %d, 取首)" % len(hits), key[:30])
    return hits[0]


# ---- 1) §3.10 TCGA：数字全错 + 遗漏显著阳性结果 ----
i = find_one("将 ST 训练的表达分支")
if i is not None:
    new = ("将 ST 训练的表达分支（GeneEncoder + 风险头，冻结权重）应用于 TCGA-CESC bulk RNA-seq"
           "（n=302，72 例事件），逐样本风险分数的单因素 Cox 回归未达显著"
           "（HR = 1.209，95% CI 0.760–1.924，Wald p = 0.422，C-index = 0.596）。"
           "这一阴性结果有清晰解释：（1）ST spot 级风险捕获的空间异质性在 bulk 平均中丢失；"
           "（2）模型未为生存预测训练，其优化目标是空间风险表征。"
           "值得注意的是，将同一表达分支的 108 基因 HPV 通路均值作为单变量特征输入 Cox 模型时，"
           "得到显著预后关联（HR = 1.688，95% CI 1.055–2.701，Wald p = 0.029，C-index = 0.620，n=302），"
           "表明框架捕获的 HPV 通路信号在独立 bulk 队列中具有预后相关性；"
           "而逐 spot 空间风险分数本身不直接迁移至 bulk 生存预测。"
           "本研究定位为空间风险表征的方法学贡献，非临床预后工具。")
    set_para_text(d.paragraphs[i], new)
    log.append("§3.10 TCGA 重写（n=304→302；p≈0.9→0.422；补回显著阳性结果 HR=1.688 p=0.029）")

# ---- 2) §3.1 Moran's I 口径注释补充 ----
i = find_one("本文出现三个 Moran")
if i is not None:
    t = d.paragraphs[i].text
    new = (t + " 另需说明：结果文件中另存有一套基于全局邻接矩阵计算的 Moran's I"
           "（40 基因 panel 落在 0.544–0.598 区间），与本文「逐样本 k=6 邻接后跨样本均值」的口径"
           "（0.891–0.938）不属于同一计算方式，数值不可混用；本文正文与全部表格统一采用后者。")
    set_para_text(d.paragraphs[i], new)
    log.append("§3.1 Moran's I 口径注释补充（消除 0.549 vs 0.926 双口径疑点）")

# ---- 3) §3.5 CA_HPV_2 全数据 vs LODO 澄清 ----
i = find_one("CA_HPV_2（腺癌）——基因均值 0.67")
if i is not None:
    t = d.paragraphs[i].text
    new = t.replace(
        "模型风险 1.195（高）",
        "模型风险 1.195（高；此为 35 例全数据训练后的 in-sample 分数，"
        "与表 2 中 LODO 留出该样本时的 0.047 分属不同实验设置，二者不可直接比较）")
    set_para_text(d.paragraphs[i], new)
    log.append("§3.5 CA_HPV_2 数值口径澄清（1.195 in-sample vs 0.047 LODO）")

# ---- 4) §3.8 GSE268014 样本数说明 ----
i = find_one("模型在 HPV 相关 HNSCC 样本中产生中等空间一致性")
if i is not None:
    t = d.paragraphs[i].text
    new = (t + " 需说明：GSE268014 为 39 例队列，本文使用其中已完成下载与预处理的 8 例"
           "（GSM8283897–GSM8283904）；GSE281978 的全部 6 例均纳入分析。")
    set_para_text(d.paragraphs[i], new)
    log.append("§3.8 GSE268014 样本来源说明（39 例队列仅用 8 例）")

# ---- 5) §4.4(3) TCGA 局限性同步 ----
i = find_one("TCGA bulk 验证阴性：模型未为生存预测训练")
if i is not None:
    new = ("（3）TCGA bulk 验证部分阴性：逐样本风险分数的 Cox 回归未达显著"
           "（HR = 1.209，95% CI 0.760–1.924，p = 0.422），因模型未为生存预测训练且 ST spot 级"
           "空间异质性在 bulk 中丢失；但以 108 基因 HPV 通路均值为特征的 Cox 模型显著"
           "（HR = 1.688，95% CI 1.055–2.701，p = 0.029，C-index = 0.620），"
           "提示通路信号具备 bulk 预后相关性，需在独立队列进一步验证。")
    set_para_text(d.paragraphs[i], new)
    log.append("§4.4(3) 局限性同步更新")

# ---- 6) 表 3 标题补 AUC 列脚注 ----
i = find_one("表 3. 消融实验结果")
if i is not None:
    new = ("表 3. 消融实验结果（† = 40 基因与 108 基因 panel 不可直接比较绝对值；"
           "* 该列为 HPV 通路内部一致性 AUC（hpv_proxy_auc），因与训练目标共享标签存在循环性，"
           "仅作描述性指标、不作验证证据；主实验的疾病/正常 AUC 另见表 4）")
    set_para_text(d.paragraphs[i], new)
    log.append("表 3 标题补 AUC 列脚注（原列名误标为「疾病/正常 AUC」）")

# ---- 7) 摘要：补 TCGA 阳性结果（可选增强）----
i = find_one("框架在独立宫颈腺癌队列中得到外部验证，并可迁移至 HPV 相关 HNSCC。")
if i is not None:
    t = d.paragraphs[i].text
    if "TCGA" not in t:
        new = t.replace(
            "并可迁移至 HPV 相关 HNSCC。",
            "并可迁移至 HPV 相关 HNSCC；其 HPV 通路特征在 TCGA-CESC bulk 队列中亦显示显著预后关联"
            "（HR = 1.688，95% CI 1.055–2.701，p = 0.029）。")
        set_para_text(d.paragraphs[i], new)
        log.append("摘要补充 TCGA 阳性结果")

d.save(DST)
print("=== 段落重写完成 ===")
for x in log:
    print("  [OK]", x)
print("\n已保存:", DST)
