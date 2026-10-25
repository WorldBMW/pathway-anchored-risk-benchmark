# -*- coding: utf-8 -*-
"""Build SI Table S15 (sensitivity of the feasibility threshold to the retained fraction c)."""
import io
import os
import shutil

from docx import Document

PD = r'G:\project4_cervix_multimodal\papers\20_MRP-ST-Repair_Patterns'

cohorts = [('Shared-program simulation', 5.8, 'no accepted cell by construction'),
           ('Cervical (35 sections)', 2.4, 'no accepted cell'),
           ('Head and neck (16)', 2.18, '0.46-0.73'),
           ('Prostate (52)', 2.0, '0.67-0.76'),
           ('HAR, wearable activity', 13.8, 'no accepted cell'),
           ('Gas sensor drift', 11.0, 'no accepted cell')]

cs = [0.5, 0.6, 0.7, 0.8, 0.9]
thr = {c: 1.645 / c for c in cs}

hdr = ['Cohort', 'tau_base / sigma'] + ['c = %.1f (threshold %.2f)' % (c, thr[c]) for c in cs] + \
      ['Realized retention tau_LDO / tau_base at leakage-passing cells']

rows = []
for name, ratio, ret in cohorts:
    row = [name, '%.2f' % ratio]
    for c in cs:
        if ratio < thr[c] - 0.06:
            row.append('feasible')
        elif ratio < thr[c]:
            row.append('feasible (marginal)')
        else:
            row.append('empty')
    row.append(ret)
    rows.append(row)

DOCX = os.path.join(PD, 'Supplemental_Information_Patterns.docx')
shutil.copyfile(DOCX, DOCX.replace('.docx', '_v8pre_20261024.docx'))
doc = Document(DOCX)

p = doc.add_paragraph()
p.add_run('Table S15. ').bold = True
r = p.add_run('Sensitivity of the feasibility threshold to the retained fraction c.')
r.bold = True
p2 = doc.add_paragraph()
p2.add_run(
    'Proposition 1 gives a threshold of 1.645/c for the pre-repair ratio τ_base/σ̂, so the feasibility verdict is '
    'indexed by c, the fraction of the baseline association a deployer requires the repair to retain. c is a '
    'protocol choice, not a constant: this table gives the threshold and the resulting verdict for five values, so '
    'that a reader who requires a different tolerance can read the corresponding column without re-deriving the '
    'proposition. The last column records what retention was actually realized at cells that passed the leakage '
    'axis, which is the empirical basis for quoting c = 0.7 as the value at which the fidelity axis binds without '
    'being empty: c = 0.9 admits no cell in any real cohort we ran, and c = 0.5 would certify cells retaining less '
    'than half the baseline association. "Marginal" marks a ratio within 0.06 of the threshold.')

tb = doc.add_table(rows=1, cols=len(hdr))
tb.style = 'Table Grid'
h = tb.rows[0].cells
for i, x in enumerate(hdr):
    h[i].text = ''
    rr = h[i].paragraphs[0].add_run(x)
    rr.bold = True
for row in rows:
    cs_ = tb.add_row().cells
    for i, x in enumerate(row):
        cs_[i].text = x

doc.save(DOCX)
print('appended Table S15; SI tables now %d' % len(Document(DOCX).tables))

TP = os.path.join(PD, 'TitlePage_Patterns.md')
s = io.open(TP, encoding='utf-8').read()
old = '14 supplemental tables (S1–S14;'
new = '15 supplemental tables (S1–S15;'
if old in s:
    io.open(TP, 'w', encoding='utf-8').write(s.replace(old, new))
    print('title page count updated')
else:
    print('WARN title page pattern not found')

old2 = 'S14 carries the per-cell dose-grid σ̂ drift and gate-robustness check behind the ' \
       'pre-computability scope condition)'
new2 = 'S14 carries the per-cell dose-grid σ̂ drift and gate-robustness check behind the ' \
       'pre-computability scope condition, and S15 carries the sensitivity of the feasibility threshold to the ' \
       'retained fraction c)'
if old2 in s:
    s = s.replace(old2, new2)
    io.open(TP, 'w', encoding='utf-8').write(s)
    print('title page S15 note added')
else:
    print('WARN S15 note pattern not found')
