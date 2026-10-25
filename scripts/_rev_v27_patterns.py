"""Evidence-constrained editorial revision; never modifies raw results."""
from pathlib import Path
import re,json,shutil,sys
from docx import Document
sys.stdout.reconfigure(encoding='utf-8')
P=Path('papers/20_MRP-ST-Repair_Patterns'); R=Path('results')
old=(P/'MRP-ST-Repair_manuscript_Patterns_v26.md').read_text(encoding='utf-8')
t=old
title='Auditing label dependence and sibling-unit inflation in group-structured representations'
t=re.sub(r'^# .*', '# '+title,t,count=1)
def section(start,end,new):
 global t
 a=t.index(start); b=t.index(end,a+len(start)); t=t[:a]+new.strip()+'\n\n'+t[b:]
section('## HIGHLIGHTS','## GRAPHICAL ABSTRACT', '''## HIGHLIGHTS

1. Paired unit and group holdouts reveal sibling-unit inflation
2. Refit permutation audits measure label association after intervention
3. A fixed-scale bound exposes conflict in a remove-and-preserve gate
4. Repair reduces identity decoding without demonstrated endpoint benefit

## IN BRIEF

Bai and colleagues audit label dependence in grouped representations using refit permutation inference and paired unit and group holdouts. Spatial transcriptomic controls and two non-biological datasets reveal evaluation inflation and limits of adversarial repair. Reduced label association alone does not establish removal of identity leakage or improved prediction.

## SUMMARY

Repeated observations from the same patient, subject, or batch can inflate evaluation when related units cross validation folds. We examine this problem through paired unit-level and group-level holdouts, a label-dependence index calibrated by refitting a ridge probe under label permutation, and adversarial interventions on learned representations. Spatial transcriptomic simulations and public cohorts reveal that unit-level evaluation can remain optimistic after group-level transfer deteriorates. Two non-biological datasets show smaller but consistent sibling-unit inflation. A fixed-null-scale inequality identifies when low label dependence and retention of the same association are incompatible; it does not guarantee a reachable repair dose. Identity decoding decreases only partially, and no accepted intervention demonstrates improved downstream prediction. The resulting workflow separates label association, sibling-unit inflation, identity decoding, and retained utility, with explicit limits on causal interpretation and cross-domain transfer.

**Keywords:** grouped data; evaluation leakage; permutation inference; representation auditing; adversarial training; spatial transcriptomics
''')
section('## BIGGER PICTURE','## INTRODUCTION', '''## BIGGER PICTURE

Many datasets contain repeated measurements from the same person, institution, or instrument. Randomly splitting these measurements can place related observations in both training and evaluation sets. Performance on such a split may say little about performance on a new person or site.

This study examines how to audit representations in that setting. We compare predictions when one observation is held out with predictions when its entire group is held out, and express label association relative to a refitted permutation reference. We then examine what changes when an adversarial training objective discourages recognition of group identity. These measurements answer different questions: label association is not itself evidence of leakage, and weak association is not evidence that identity has been erased.

The experiments include spatial transcriptomic controls and public biological cohorts, plus exploratory audits of wearable activity and chemical sensor data. Repair reduced some measures of identity decoding but did not demonstrate improved downstream prediction. An algebraic bound also shows that requiring both a low value of an association statistic and preservation of that same statistic can be contradictory. The practical contribution is an evidence-traceable workflow for reporting these trade-offs, rather than a universally effective repair or a guarantee of deployment performance.
''')
section('## INTRODUCTION','## RESULTS', '''## INTRODUCTION

Patients contribute multiple tissue measurements, subjects contribute many sensor windows, and instruments generate repeated readings. These observations are dependent. If related units appear on both sides of an evaluation split, models can exploit information shared within a group. Group-disjoint evaluation is therefore established practice when the intended use is prediction on previously unseen groups.[7,39,40] Leakage is a recognized source of irreproducibility across machine-learning applications,[1,22] and shortcut learning can preserve benchmark performance without solving the intended task.[4–6] Repeated units create an opportunity for leakage; they do not guarantee that a model will exploit it or that every random split will be optimistic.

An existing representation can be examined without retraining its encoder by fitting probes under explicitly declared holdouts. Such an analysis is conditional on the representation: if the encoder was fitted using all observations, holding groups out only from a probe does not make the complete pipeline externally validated. Moreover, a probe's association with a target label is distinct from its ability to decode group identity. Strong label association may reflect a shared biological program, a nuisance shortcut, or both. A low association may reflect successful nuisance reduction or destruction of useful information. These alternatives require separate measurements and controls.

We study an audit workflow combining a ridge-probe label-dependence index (LDI), paired leave-one-unit-out and leave-one-group-out associations, group-identity decoding, and retained signal or downstream endpoints. LDI expresses the unit-level cross-validated label association relative to its refit permutation scale. It is not a direct estimate of the fraction of predictive performance caused by identity. Where one unit corresponds to one group, unit and group holdouts coincide; where groups contribute siblings, we report both rather than treating them as interchangeable. The unit-versus-group difference is the statistic that directly describes sibling-unit inflation in this study.

Adversarial representation learning is an established approach to domain invariance and sensitive-attribute suppression.[2,3,26–28,35,36] We use a standard gradient-reversal group classifier as an intervention, not as a novel operator. Group-DRO[9] and invariant risk minimization[10] provide additional comparisons on the recorded configurations. Batch-integration methods address related but different objectives,[11,12,32–34] and linear projection is a relevant inexpensive alternative.[31] We make no claim that this intervention outperforms the wider family of deconfounding methods.

The study asks three questions. First, can unit-level association remain high while group-level transfer deteriorates? We examine a controlled multi-unit simulation and two public non-biological datasets. Second, can a label-association gate conflict with retention of the same association? We derive a fixed-scale interval condition and distinguish that algebraic condition from existence of a reachable repair dose. Third, what does the intervention preserve or damage in spatial transcriptomics? Planted controls assess signal retention and detector behavior, while three public cohorts provide dose responses and downstream readouts. Cross-cohort contrasts are descriptive, because label semantics, cohort composition, and signal strength change together.

The scope is deliberately limited. Spatial transcriptomics supplies the planted controls and biological readouts. Wearable activity recognition and gas-sensor drift supply exploratory audits of sibling-unit inflation and partial identity suppression, without evidence of downstream benefit. The other domains mapped in Table 4 are untested. No accepted repair improved a measured downstream endpoint. The contribution is the explicit separation and traceability of these measurements, together with the failure modes revealed by their disagreement.
''')
# Correct categorical claims throughout retained Results and captions.
t=t.replace('unit-level evaluation is optimistic whenever a group contributes more than one unit','unit-level evaluation can be optimistic when groups contribute sibling units')
t=t.replace('label–identity coupling determines repairability','label–identity coupling and observed repairability')
t=t.replace('A random-label control on the same cohort establishes the coupling as causal.','A random-label control checks sensitivity to the observed assignment, without isolating a causal mechanism.')
t=t.replace('**A random-label control isolates the cause.**','**A random-label control tests the observed assignment.**')
t=t.replace('breaking every label–identity coupling','breaking the observed label assignment while preserving the section-level class counts')
t=t.replace('The contrast is causal:', 'The contrast is consistent with assignment-dependent structure:')
t=t.replace('Findings 3 to 8','Findings 3 to 7')
t=t.replace('Sci. Rep. 16, 43717 (2026)','Sci. Rep. 16, 13246 (2026)')
# The published GeoMx paper distinguishes specimens, patients, ROIs and compartments.
t=t.replace('186 regions of interest from 7 cases','186 compartment measurements from 93 ROIs across 7 specimens (6 patients)')
t=t.replace('186 ROIs, 7 cases','186 compartments, 7 specimens, 6 patients')
t=t.replace('186 ROIs from 7 cases','186 compartments from 93 ROIs across 7 specimens (6 patients)')
t=t.replace('Baseline Δ was +0.0075, +0.3736 and +0.3955 on the three arms','Baseline Δ was +0.3736, +0.3955 and +0.0075 for risk, Silva pattern and cell type, respectively')
t=t.replace('τ_base/σ̂ was 8.49, 9.34 and 10.39','τ_base/σ̂ was 9.34, 10.39 and 8.49 for risk, Silva pattern and cell type, respectively')
# Different archived experimental configurations are not one interchangeable cohort.
t=t.replace('so each donor contributes one unit.', 'so each section contributes one unit. The archive reports 35 section units and 34 donor groups; donor-disjoint and section-disjoint holdouts must therefore be distinguished.')
t=t.replace('In the validation domain that number is one section per donor, so the trap is dormant and the two statistics coincide;', 'In the head-and-neck configuration there is one section per group. The cervical archive instead reports 35 sections and 34 donor groups, so equality cannot be assumed there;')
t=t.replace('Our repairability claims are therefore computed at k = 1, where the two statistics coincide by construction,', 'Our main repairability calculations use section-level embeddings; unit and group holdouts coincide only when the recorded section-to-group map is one-to-one,')
t=t.replace('every donor contributes one unit and leave-one-unit-out is identical to leave-one-group-out','every section is assigned its own group in this control and leave-one-unit-out is identical to leave-one-group-out')
t=t.replace('each donor contributes K sibling units','each section contributes K sibling units')
# Replace the statistical interpretation with the implementation actually archived.
section('### Notation, probes, and the instrument','### Proposition 1', '''### Notation, probes, and the instrument

For m analysis units with embedding matrix X and numerical label vector y, τ_in is the in-sample Kendall association of a ridge probe with y. τ_LOO holds out one unit from probe fitting; τ_LGO (also written τ_LDO) holds out all units of a recorded group. Training-fold standardization and ridge regularization are applied within the probe. These are probe-level holdouts on representations, not a claim that the complete encoder-training pipeline was refitted in every outer fold.

The archived implementation computes LDI = (τ_LOO − 0.001)/σ̂, with σ̂ estimated from B = 2,000 unit-label permutations, refitting the probe for every permutation (seed 20261010). R = τ_LOO/τ_in is the fraction of in-sample rank association retained under unit-level holdout; it is group-transferable only when units and groups coincide. Sibling-unit inflation is reported separately as Δ = τ_LOO − τ_LGO. Group-identity decoding is a separate classifier and is not what the label probe measures.

The 1.645 cutoff is a one-sided standard-normal reference. Normality and exchangeability are assumptions, not consequences of estimating a permutation standard deviation. Unit-label permutation is not automatically valid for clustered observations. Where labels are constant within groups, inference requires group-preserving permutation; where labels vary within groups, a null must be chosen for the specific scientific question. The archived multi-unit results use the unit-permutation reference and are descriptive under that reference, not calibrated group-level 5% tests. Crossing below 1.645 is a reporting convention and cannot establish absence of leakage or equivalence to a clean representation. The binary-label ceiling is sqrt(n₀n₁ / C(m,2)) when predictions have no ties (0.7303 for 8 versus 8). Multiclass activity and gas identifiers are nominal: their arbitrary numerical coding makes Kendall τ encoding-dependent; class-aware metrics and recoding sensitivity are needed for a task-level conclusion.

### Relation of LDI to existing calibration procedures

Refit permutation inference is established practice.[38] We use its scale to report label association, not as a novel test or causal attribution. For a fixed empirical null, larger observed τ corresponds to a smaller upper-tail permutation p-value; the numerical approximation LDI ≈ Φ⁻¹(1 − p) additionally requires a suitably centered Gaussian null and is not guaranteed by B = 2,000. The fixed 0.001 floor is not an estimated null mean. LDI remains sensitive to sample size, label coding, probe capacity, and the permutation scheme, so a common numerical scale alone does not establish transportability across cohorts. In particular LDI is the calibrated association, not the calibrated memorization gap τ_in − τ_LOO.

The supplementary resolution ladder changes the number of analysis units while retaining the source sections. Its σ̂ falls as unit count grows; this illustrates resolution sensitivity rather than an independent increase in donor-level information. The archived ladder assigns section identifiers as groups (35 groups), whereas the main cervical archive reports 34 donor groups. It therefore tests section blocking, and does not verify donor-disjoint evaluation.
''')
# Explicit scope qualifications without changing archived numerical outcomes.
insert='''**Interpretation of Proposition 1.** The result below describes overlap of two intervals for the same statistic τ at a fixed positive σ̂, positive τ_base and 0 < c ≤ 1. Non-empty overlap is necessary for those numerical requirements to coexist; it is not sufficient for an encoder or dose to attain a point in the overlap. A baseline-only exclusion is conditional on holding σ̂ fixed, or on a prespecified bound on its post-intervention value. When the leakage axis uses τ_LOO and the fidelity axis uses a different τ_LGO, the one-dimensional equivalence does not apply. In the archived cross-domain and GeoMx runners both gates use τ_LOO; τ_LGO retention is reported separately. For the external section cohorts the published Table 2 retention counts use the cohort-mean baseline, rather than each seed's baseline. These conventions must not be mixed. The exact critical retention is (0.001 + 1.645 σ̂)/τ_base; tables retain the near-zero-floor approximation for correspondence with the archive.

'''
t=t.replace('### Proposition 1 (precondition test for the dual gate)\n\n','### Proposition 1 (precondition test for the dual gate)\n\n'+insert)
t=t.replace('whether any dose can remove identity and keep signal','whether the fixed-scale numerical gate permits overlapping association values')
t=t.replace('The boundary\'s *existence* transported; its *placement* did not','This experiment does not map a group-count boundary outside the validation domain; its numerical placement did not transport')
t=t.replace('Units were 300 observations per domain','The requested audit size was 300 observations per domain (299 realized for HAR and 300 for gas drift)')
t=t.replace('Audit units: 300 observations drawn','Audit units: 299 HAR and 300 gas-drift observations drawn')
t=t.replace('The gap is not sampling variation.', 'The observed seed ranges are distinct; seeds and doses are repeated analyses of the same cohorts, not independent population samples.')
t=t.replace('and a bootstrap 95% interval on the difference of means is +0.024 to +0.063 using baseline cells and +0.032 to +0.047 using all 15 cells, neither interval containing zero.', 'and the archived cell bootstrap intervals are descriptive only: repeated doses and seeds do not supply independent cohort replication.')
t=t.replace('1.645 (one-sided 5%)','1.645 (one-sided normal reference)')
# Correct an explicit false definition in the statistics subsection.
t=t.replace('σ̂ is the standard deviation, across the B permutations, of the leave-one-group-out (equivalently leave-one-donor-out) Kendall τ recomputed with the probe refit inside every permutation.', 'σ̂ in the archived audit is the standard deviation of the leave-one-unit-out Kendall τ across unit-label permutations, with probe refitting. It equals a group-holdout null only where the unit-to-group mapping is one-to-one; the multi-unit null is not group-preserving.')
t=t.replace('applied to LDI = (τ_LGO − 0.001)/σ̂','applied in the archived implementation to LDI = (τ_LOO − 0.001)/σ̂')
section('## DISCUSSION','## EXPERIMENTAL PROCEDURES', '''## DISCUSSION

The experiments demonstrate why label association, group-level transfer and identity decoding must be reported separately. Unit-level association remained at its ceiling in 45 simulated multi-unit runs while group-level transfer deteriorated. Both non-biological cohorts showed positive unit-versus-group differences across their recorded seed–dose cells, although the absolute effect was small in HAR. These findings support paired evaluation under a declared grouping map. They do not show that every grouped dataset is affected or quantify the causal contribution of identity to an inherited model's published accuracy.

The planted controls establish behavior within their specified generators: the strong shared program survived 75 runs, while 45 identity-confounding runs remained below the selected label-association reference. These counts are empirical control outcomes, not a general false-positive rate. An identity-confounded representation can have low held-out label association; indeed that is why low LDI alone cannot certify removal. Independent identity decoding and retained utility are essential to interpretation.

The fixed-scale feasibility inequality exposes a conflict in the acceptance criterion itself. If the same positive association must lie below a low-dependence threshold and above a retained fraction of baseline, the intervals may not overlap. This is an algebraic diagnostic of that criterion, not an identifiability theorem, a fairness impossibility result, or a prediction of which dose an optimizer can reach. Low LDI measures low label association, which can reflect loss of genuine signal. A defensible intervention therefore reports the separate axes rather than treating threshold crossing as successful deconfounding.

The external-cohort results are limited. Under the cohort-mean baseline convention and the reported dose subset, three of eighteen seed–dose cells pass both recorded axes: one head-and-neck cell and two prostate cells. Different baseline conventions can change these counts. No accepted repair improved a measured downstream endpoint. The head-and-neck AUROC declined from 0.906 ± 0.063; prostate comparisons have only four benign sections and do not resolve small differences. Lower identity decoding is therefore not evidence of clinical benefit. The descriptive contrasts among severity, HPV status and specimen provenance do not isolate label semantics from cohort size, source, class balance or signal strength. Randomizing labels changes their relation to both identity and useful signal and cannot by itself establish a causal coupling mechanism.

Group-DRO and IRM behaved differently from the adversarial intervention on the archived configurations. The deliverable of this paper is the auditing protocol, not the repair operator. We make no claim that the adversarial operator outperforms other deconfounding methods. Untested projection, contrastive and reweighting families could yield different trade-offs. The simulation boundary between 16 and 32 groups is a configuration-specific observation, not a transferable minimum sample size. Gas drift retained substantial group-level association with ten batches, and no outside-domain phase grid establishes a corresponding boundary.

### Scope beyond spatial transcriptomics

Table 4 gives structural mappings. Spatial transcriptomics supplies planted controls and domain readouts; HAR and gas drift supply exploratory audits only. Neither non-biological cohort demonstrates downstream utility or complete non-blindness calibration. The six remaining mappings are untested. A third-party pipeline, a class-aware analysis of nominal labels and valid group-preserving inference are needed before broader validation claims are justified.

'''+old[old.index('**Table 4.'):old.index('Two caveats keep the mapping honest.')]+'''

### Limitations

The representation is generally learned before the probe holdouts. Consequently these are conditional representation audits, not fully nested estimates of encoder performance on unseen groups. Encoder refitting, preprocessing fitted only within outer training folds, and independent external evaluation are required for that stronger claim. Package versions and hardware listed with archived experiments describe provenance; they have not been independently demonstrated here to reproduce every result from raw data.

The implemented LDI and null use unit-level holdouts and unit-label permutations. In multi-unit data they are not equivalent to group-level inference. The nominal multiclass labels in HAR and gas drift have no intrinsic rank; numerical recoding can change Kendall τ. A one-sided 1.645 reference does not prove absence of identity information, and a nonlinear probe can decode information missed by the linear ridge probe. Re-audit with a more expressive probe wherever the deployment is high-stakes, using a null appropriate to that probe and grouping scheme.

The group identifiers also bound the conclusions. The main cervical archive reports 35 section units and 34 donors, whereas the resolution ladder uses 35 section groups. The repeated donor must be resolved from source metadata before a donor-disjoint certification is possible. The GeoMx source reports seven specimens from six patients and 93 ROIs split into 186 epithelial/stromal compartment measurements.[42] The archived runner groups by seven case identifiers. Its holdout therefore does not establish patient-disjoint evaluation, and the paired compartment structure must be preserved when constructing an appropriate null. Source-study ROI selection also limits representativeness.

Seed variability measures optimization variability on fixed cohorts, not uncertainty across patients, sites or datasets. Repeated doses share data and are not independent replicates. Cell-level bootstrap intervals should therefore be read descriptively; population inference would require resampling independent groups with the analysis repeated. The strong-dose instability, small positive class in prostate, and restricted head-and-neck subset limit precision. We report no dose optimum.

The relative fidelity gate is uninformative near a null baseline: the spurious simulation arm yields accepted cells without a useful transferable program. An absolute utility requirement should be prespecified, and failure to reject an association null should not be treated as an equivalence test. The archived results did not implement that additional requirement. No accepted repair improved a downstream endpoint; no clinical deployment benefit or privacy guarantee is established.

### Practical reporting workflow

1. Record the observational unit, independent grouping variable, intended deployment target and any repeated specimens from the same patient.
2. Distinguish the encoder's training data from probe training folds. Use fully nested group splits if claiming generalization of the complete pipeline.
3. Report τ_LOO, τ_LGO and their difference, alongside direct identity decoding and an independently defined utility endpoint. LDI alone does not measure identity removal.
4. Choose a group-preserving permutation scheme and a task-appropriate metric. Report nominal-label coding sensitivity and empirical null quantiles; state whether a normal reference is only descriptive.
5. Apply the interval inequality only to the same statistic and a declared fixed null scale. Do not infer that non-empty overlap guarantees a reachable dose.
6. Report both intervention costs and retained signal on controls, with baseline conventions and the tested dose subset specified. Do not read acceptance as erasure.
7. Release the exact code, configurations, data joins and result-to-panel mapping before claiming public reproducibility. Resolve patient identifiers before patient-level inference.
''')
# Prominent reader-facing account of archive limitations, also governing Results.
note='''**Interpretation of archived results.** The numerical results below retain the original experimental configurations. LDI is calibrated on unit-level τ_LOO, whereas sibling inflation is the separate difference τ_LOO − τ_LGO. The two coincide only for a one-to-one unit/group map. In the multi-unit audits, the published dual gate uses τ_LOO for both the dependence and retention axes; it is not a group-level acceptance certificate. The fixed-scale feasibility calculation describes interval overlap, not existence of a reachable dose. Nominal-label rank coding, non-group-preserving permutations, and representation fitting before probe holdouts limit inferential claims. These restrictions apply to every result, figure and table below; details are in Experimental Procedures and Discussion.

'''
t=t.replace('## RESULTS\n\n','## RESULTS\n\n'+note)
# Status claims require accessible deposits, not a proposed URL.
for pat in [r'\*\*Code\.\*\*[^\n]+',r'(?m)^All code[^\n]+']:
 t=re.sub(pat, '**Code.** The local submission package contains the analysis scripts and archived results. The proposed public repository is https://github.com/WorldBMW/cervix-st-multimodal and the intended release is v1.0.0. Public accessibility and an immutable archive must be verified before submission; a proposed URL is not evidence of a released artifact.',t)
t=t.replace('and deposited in Zenodo before publication, whose DOI is reported in the version of record','with public archiving pending verification')
t=t.replace('Every run reported here, simulated or real, is produced by a script that reads one JSON configuration manifest and writes one JSON result file; no analysis step accepts an interactive argument that is not recorded in the manifest.', 'The archived runners use script arguments and record parameter dictionaries in result JSON files. A universal manifest-only interface was not verified.')
t=t.replace('The archive also contains a checksum index mapping each table row and figure panel to the manifest and result file from which it was computed, which is the artefact a reader needs in order to re-derive any reported value without correspondence with the authors.', 'A complete row-level and panel-level checksum mapping must accompany any public reproducibility claim; its completeness has not been verified.')
t=t.replace('so every number in every table is traceable to one manifest and one result file','to support traceability; completeness of the mapping requires verification')
t=t.replace('a pre-computed condition saying whether any dose can remove identity and keep signal','a fixed-scale numerical compatibility condition')
out=P/'MRP-ST-Repair_manuscript_Patterns_v27.md';out.write_text(t,encoding='utf-8')
for f in ['Manuscript_Patterns.docx','CoverLetter_Patterns.docx','TitlePage_Patterns.docx','Supplemental_Information_Patterns.docx']:
 dst=P/(Path(f).stem+'_v27pre_20261025.docx')
 if not dst.exists():shutil.copy2(P/f,dst)
# Use the existing deterministic manuscript converter.
sys.path.insert(0,str(Path(__file__).resolve().parent))
from _md2docx import convert
convert(str(out),str(P/'Manuscript_Patterns.docx'))
# Preserve supplement tables and images, correct only the explanatory prose.
d=Document(P/'Supplemental_Information_Patterns.docx')
for p in d.paragraphs:
 s=p.text
 s=s.replace('186 regions of interest (ROIs) from 7 cases','186 compartment measurements from 93 ROIs across 7 specimens (6 patients)')
 s=s.replace('the group-label coupling kappa falls from','the group-label coupling kappa is')
 s=s.replace('every donor contributes exactly one unit','every section is assigned its own group and contributes exactly one unit')
 s=s.replace('each donor contributes K sibling units','each section contributes K sibling units')
 s=s.replace('the identity classifier on the same embedding retains a large lift over chance throughout','the identity classifier is not evaluable as a repeated-group decoding check at K = 1 (the archived lift is zero); at K > 1 repeated section groups permit decoding')
 if s!=p.text:p.text=s
q=d.add_paragraph('Interpretation of archived supplemental analyses')
q.style='Heading 1'
d.add_paragraph('LDI and its refit permutation scale are based on leave-one-unit-out label association. For multi-unit cohorts they are not calibrated leave-one-group-out tests. Tables S13 and S17 use unit-level association in both archived gate axes; group-level retention is reported separately. Table S16 blocks by 35 section identifiers, not the 34 donor identifiers of the main cervical archive. The source GeoMx study reports seven specimens from six patients; the seven-case holdout is not patient-disjoint, and the 186 measurements are epithelial/stromal compartments paired within 93 ROIs. Normal-reference thresholds, unit-label permutations and nominal-label rank coding limit the conclusions. No new inference or experiment is implied by this editorial correction.')
d.save(P/'Supplemental_Information_Patterns.docx')
wc=lambda s:len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'+/-]*",s))
intro=t[t.index('## INTRODUCTION'):t.index('## EXPERIMENTAL PROCEDURES')]
stats={'title_characters':len(title),'main_text_words_including_inline_tables':wc(intro),'summary_words':len(t.split('## SUMMARY\n\n')[1].split('**Keywords:')[0].split()),'in_brief_words':len(t.split('## IN BRIEF\n\n')[1].split('## SUMMARY')[0].split())}
tp=f'''# Title page

**Article type:** Research Article, Patterns

**Title:** {title}

**Authors:** Mingwei Bai¹, Zhengmin Gu², Ye Zhang², Zhiping Ni¹, Lang Guo³

**Affiliations:** ¹JiMei Technology Group Co., Ltd., Shenyang 110002, China; ²Medical Record Management Center, The First Hospital of China Medical University, Shenyang 110002, China; ³School of Pharmacy, Shenyang Medical College, Shenyang 110034, China

**Corresponding author and lead contact:** Lang Guo, School of Pharmacy, Shenyang Medical College, 146 Huanghe North Avenue, Shenyang 110034, China. guolang1016@symc.edu.cn; +86 13555707001

**Contents:** Four highlights; In brief ({stats['in_brief_words']} words); Summary ({stats['summary_words']} words); Bigger picture; five main figures supplied separately; six main tables including the key resources table; seventeen numbered supplemental tables (some contain multiple Word table objects); 43 references.

**Keywords:** grouped data; evaluation leakage; permutation inference; representation auditing; adversarial training; spatial transcriptomics

**Declarations:** See the manuscript for author contributions, competing interests, and AI-assisted editorial preparation.
'''
(P/'TitlePage_Patterns.md').write_text(tp,encoding='utf-8');convert(str(P/'TitlePage_Patterns.md'),str(P/'TitlePage_Patterns.docx'))
cl=f'''# Cover letter

25 October 2026

Dear Editors of Patterns,

Please consider our manuscript, “{title},” as a Research Article.

The study concerns evaluation of representations learned from repeated observations within patients, subjects or batches. It combines paired unit-level and group-level probe holdouts, refit permutation reference scales, direct identity decoding, planted signal controls and downstream readouts. Its relevance to Patterns is methodological: these measurements separate label association from identity decoding and from the performance differences caused by the evaluation split.

The evidence consists of spatial transcriptomic simulations and public cohorts, together with exploratory wearable-activity and gas-sensor audits. A controlled simulation shows unit-level association remaining high after group-level transfer deteriorates. A fixed-scale interval inequality explains when requiring both low association and retention of that same association is contradictory. This inequality is not a guarantee that an intervention can attain a feasible point. The adversarial operator is standard, suppression is partial, and no accepted repair demonstrates improved downstream prediction. Cross-domain rank metrics depend on nominal-label coding and the archived multi-unit permutation scheme is not group-preserving; these limitations are stated explicitly.

Related work disclosure. A related manuscript by the same group is available locally at papers/19_ISCLS_20261016/Manuscript_ISCLS.md and should accompany submission as a related file. It also discusses label-dependence indexing and uses overlapping public cervical data. It is not cited as evidence for this manuscript. Its current editorial status and any overlap in numerical results, text or figures require author verification before this letter is uploaded. Public data overlap is disclosed and should not be represented as evidence that substantive overlap is absent.

Reproducibility materials are prepared locally. The proposed public repository and intended release are recorded in the manuscript, but public access and immutable archiving must be verified before submission. This letter does not assert that a release or Zenodo deposit already exists. The independent audit records outstanding issues concerning donor identifiers, patient-disjoint inference, null calibration, figure interpretations and package provenance.

The manuscript and figures have received AI-assisted editorial review, with the use declared in the manuscript. Author approval, exclusive submission, related-manuscript status and the final availability statement must be confirmed by the corresponding author before upload.

Sincerely,

Lang Guo, on behalf of the authors

School of Pharmacy, Shenyang Medical College, Shenyang 110034, China

guolang1016@symc.edu.cn
'''
(P/'CoverLetter_Patterns.md').write_text(cl,encoding='utf-8');convert(str(P/'CoverLetter_Patterns.md'),str(P/'CoverLetter_Patterns.docx'))
(P/'AUDIT_v27_statistics.json').write_text(json.dumps(stats,indent=2),encoding='utf-8')
print(stats)
