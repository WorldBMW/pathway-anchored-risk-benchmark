#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""_rev_v28_patterns.py — v27b -> v28：把六项重做实验的结果写进正文。

Manuscript_Patterns.docx 内嵌五幅图，不能用 md -> docx 重建（会丢图），因此全部
就地 patch：替换段落只改首个 run 的文本并保留其格式，新增段落用 deepcopy 继承
目标位置的段落样式，新增表格先建在文末再移动到目标位置。

**所有文本一律从 MRP-ST-Repair_manuscript_Patterns_v28.md 读取**，脚本内不重复
写第二份，避免 md 与 docx 漂移，也保证 σ̂ / τ_LOO / κ / Δ / λ 等符号与全稿一致。

同步的六处：
  1. Limitations 前新增 "Remediation experiments" 小节（含表 7）
  2. Limitations 前三段改写
  3. Practical reporting workflow 第 4 条改写
  4. STAR Methods 新增 "Group-preserving null and the remediated statistics."
  5. 图注 "Symbols and interpretation." 段改写
  6. 表 7（从 md 的 markdown 表格解析）

用法：python scripts/_rev_v28_patterns.py [--dry]
"""
from __future__ import annotations

import argparse
import copy
import re
import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph

P = Path(__file__).resolve().parents[1] / "papers" / "20_MRP-ST-Repair_Patterns"
DOC = P / "Manuscript_Patterns.docx"
SRC = P / "MRP-ST-Repair_manuscript_Patterns_v28.md"

BOLD_RE = re.compile(r"^\*\*(.+?)\*\*\s*(.*)$", re.S)


# ------------------------------------------------------------------ md 读取
def md_blocks():
    """把 v28 md 中新增小节切成 (kind, payload) 序列。

    kind: 'heading' | 'para' (prefix, text) | 'table' (header, rows)
    """
    md = SRC.read_text(encoding="utf-8")
    seg = md.split("### Remediation experiments")[1].split("### Limitations")[0]
    out = []
    for raw in seg.split("\n\n"):
        b = raw.strip()
        if not b:
            continue
        if b.startswith("### "):
            out.append(("heading", b[4:].strip()))
        elif b.startswith("|"):
            lines = [l.strip() for l in b.split("\n") if l.strip().startswith("|")]
            cells = []
            for l in lines:
                row = [c.strip() for c in l.strip("|").split("|")]
                if all(set(c) <= set("-: ") for c in row):
                    continue
                cells.append(row)
            out.append(("table", (cells[0], cells[1:])))
        else:
            m = BOLD_RE.match(b)
            if m:
                out.append(("para", (m.group(1), m.group(2))))
            else:
                out.append(("para", (None, b)))
    return out


def md_para(prefix):
    """按开头前缀取出 v28 md 中的整段文本。"""
    md = SRC.read_text(encoding="utf-8")
    for raw in md.split("\n\n"):
        b = raw.strip()
        if b.startswith(prefix):
            return b
    # 编号列表项之间是单换行，整块不匹配时退回逐行查找
    for line in md.split("\n"):
        s = line.strip()
        if s.startswith(prefix):
            return s
    raise SystemExit(f"paragraph not found in md: {prefix[:40]}")


# ------------------------------------------------------------------ docx 工具
def find(doc, pred):
    hits = [i for i, p in enumerate(doc.paragraphs) if pred(p.text.strip())]
    if len(hits) != 1:
        raise SystemExit(f"anchor matched {len(hits)} paragraphs")
    return doc.paragraphs[hits[0]]


def set_text(p, new):
    runs = p.runs
    if not runs:
        p.add_run(new)
        return
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""


def clone_like(ref_p, text, bold_prefix=None):
    el = copy.deepcopy(ref_p._p)
    p = Paragraph(el, ref_p._parent)
    runs = p.runs
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    if not runs:
        runs = [p.add_run("")]
    if bold_prefix:
        runs[0].text = bold_prefix if bold_prefix.endswith(" ") else bold_prefix + " "
        runs[0].bold = True
        r2 = p.add_run(text)
        if runs[0]._element.rPr is not None:
            r2._element.insert(0, copy.deepcopy(runs[0]._element.rPr))
        r2.bold = False
    else:
        runs[0].text = text
    return el


def make_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(8)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            rr = cells[i].paragraphs[0].add_run(str(v))
            rr.font.size = Pt(8)
    return t


# ------------------------------------------------------------------ 主流程
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry", action="store_true")
    args = ap.parse_args()

    # 小节标题本身被 split 掉了，这里补回（md 中它是 ### 级标题）
    blocks = [("heading", "Remediation experiments")] + md_blocks()
    limit1 = md_para("The representation is generally learned")
    limit2 = md_para("The remediation runs above replace")
    limit3 = md_para("The group identifiers still bound")
    item4 = md_para("4. Choose a group-preserving")
    star = md_para("**Group-preserving null and the remediated statistics.**")
    symb = md_para("**Symbols and interpretation.**")

    doc = Document(DOC)
    lim_h = find(doc, lambda t: t == "Limitations")
    lim1 = find(doc, lambda t: t.startswith("The representation is generally learned"))
    lim2 = find(doc, lambda t: t.startswith("The implemented LDI and null"))
    lim3 = find(doc, lambda t: t.startswith("The group identifiers"))
    it4 = find(doc, lambda t: t.startswith("4. Choose a group-preserving"))
    star_p = find(doc, lambda t: t.startswith("Definition of"))
    symb_p = find(doc, lambda t: t.startswith("Symbols and interpretation"))

    if args.dry:
        print(f"md blocks: {len(blocks)}")
        for k, v in blocks:
            if k == "para":
                print("  para  ", (v[0] or "(no bold prefix)")[:50], "|",
                      v[1][:60].replace("\n", " "))
            elif k == "table":
                print("  table ", f"{len(v[1])} rows, {len(v[0])} cols")
            else:
                print("  heading", v)
        print("anchors ok")
        return

    bak = P / f"Manuscript_Patterns_v28pre_{date.today():%Y%m%d}.docx"
    if not bak.exists():
        shutil.copy2(DOC, bak)

    cursor = lim_h._p
    body_ref = lim1
    for kind, payload in blocks:
        if kind == "heading":
            cursor.addprevious(clone_like(lim_h, payload))
        elif kind == "para":
            prefix, text = payload
            cursor.addprevious(clone_like(body_ref, text, bold_prefix=prefix))
        else:
            header, rows = payload
            tbl = make_table(doc, header, rows)
            cursor.addprevious(tbl._tbl)
    cursor.addprevious(clone_like(body_ref, ""))

    set_text(lim1, limit1)
    set_text(lim2, limit2)
    set_text(lim3, limit3)
    set_text(it4, item4)

    m = BOLD_RE.match(star)
    star_p._p.addnext(clone_like(star_p, m.group(2), bold_prefix=m.group(1)))

    m = BOLD_RE.match(symb)
    set_text(symb_p, f"{m.group(1)} {m.group(2)}")

    doc.save(DOC)
    d2 = Document(DOC)
    print(f"patched {DOC.name}: tables now {len(d2.tables)}, "
          f"paragraphs {len(d2.paragraphs)}; backup {bak.name}")


if __name__ == "__main__":
    main()
