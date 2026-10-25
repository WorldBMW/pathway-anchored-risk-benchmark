#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""_qa_v16.py — v16 稿件 QA：禁字符、表数、字数、STAR-Methods 框架与措辞统一。"""
import re
from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parent.parent / "papers" / "20_MRP-ST-Repair_Patterns"
d = Document(D / "Manuscript_Patterns.docx")
ps = [p.text for p in d.paragraphs]
full = "\n".join(ps)
for t in d.tables:
    for r in t.rows:
        full += "\n" + " | ".join(c.text for c in r.cells)

WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9'+/-]*")
fails = []

n_tbl = len(d.tables)
print("em-dash", full.count("\u2014"), "| literal **", full.count("**"),
      "| main tables", n_tbl)
if full.count("\u2014") or full.count("**"):
    fails.append("forbidden character present")
if n_tbl != 6:
    fails.append(f"main table count {n_tbl} != 6")

print("words(all):", len(WORD.findall(full)))
ii = [i for i, x in enumerate(ps) if x.strip().upper() == "INTRODUCTION"][0]
jj = [i for i, x in enumerate(ps)
      if x.strip().upper().startswith("EXPERIMENTAL PROCEDURES")][0]
body = "\n".join(ps[ii:jj])
print("main text words (Introduction -> Experimental Procedures):",
      len(WORD.findall(body)))

# ---- Highlights: 5 条，每条 <= 85 字符
hl = [x.strip() for x in ps[ps.index("HIGHLIGHTS") + 1:ps.index("IN BRIEF")]
      if x.strip()]
print(f"highlights: {len(hl)}")
for h in hl:
    flag = "OK " if len(h) <= 85 else "FAIL"
    if len(h) > 85:
        fails.append(f"highlight too long ({len(h)}): {h}")
    print(f"  {flag} {len(h):3d}  {h}")
if len(hl) != 5:
    fails.append(f"highlight count {len(hl)} != 5")

# ---- In brief <= 50 words, Summary <= 150 words
ib = "\n".join(ps[ps.index("IN BRIEF") + 1:ps.index("SUMMARY")]).strip()
sm = "\n".join(x for x in ps[ps.index("SUMMARY") + 1:ps.index("GRAPHICAL ABSTRACT")]
               if not x.strip().startswith("Keywords:")).strip()
print("in brief words:", len(ib.split()), "| summary words:", len(sm.split()))
if len(ib.split()) > 50:
    fails.append(f"in brief {len(ib.split())} > 50")
if len(sm.split()) > 150:
    fails.append(f"summary {len(sm.split())} > 150")

checks = [
    # v14 / v15 内容仍在
    "Control: the channel is opened by coupling",
    "Table S17", "Table S16", "44. [Dataset] Axelrod", "GSE316098",
    "Validation is end-to-end in spatial transcriptomics alone",
    "no accepted repair improved a downstream endpoint, which the framework predicts in advance",
    "Not tested", "structural mappings only, for which we claim no transfer",
    "We therefore make no claim that the adversarial operator used here outperforms",
    "The deliverable of this paper is the auditing protocol, not the repair operator",
    "Do not read acceptance as erasure",
    "Degenerate regime, to be read before any accept verdict is used",
    "placements for this configuration, not transferable cut-offs",
    "nonlinear probe can decode",
    # v16 新增：STAR-Methods 框架
    "Lead contact. Requests for further information, resources, or code",
    "Materials availability. This study did not generate new unique reagents",
    "Quantification and statistical analysis",
    "Key resources table",
    "What each n counts",
    "B = 10,000",
    "standard errors of the mean",
    "we report no dose optimum",
    "requirements-patterns.txt",
    # v16 新增：三级证据 standing qualifier
    "Standing qualifier for this section",
    "audited for theoretical predictions only",
    # v16 新增：协议与局限
    "Re-audit with a more expressive probe wherever the deployment is high-stakes",
    "Do not trust an accept output obtained while the baseline association is near null",
    "future extension rather than a requirement for the present manuscript",
    "it is not a detector of every form in which group identity can be encoded",
    "we report no optimal",
    "the prostate acceptance is a single operating point",
    # v16 新增：非擦除统一措辞
    "it does not mean group identity is fully unrecoverable",
    # 代码可访问性
    "releases/tag/v1.0.0",
]
for c in checks:
    n = full.count(c)
    print(f"  {'OK ' if n >= 1 else 'MISS'} {n}  {c[:70]}")
    if n < 1:
        fails.append(f"missing: {c[:60]}")

for bad in ["validated across multiple domains", "validated in multiple domains",
            "improves downstream performance", "improves prediction performance",
            "our repair improves", "fully removed", "group identity is removed"]:
    n = full.lower().count(bad.lower())
    print(f"  {'OK ' if n == 0 else 'WARN'} {n}  banned: {bad}")
    if n:
        fails.append(f"banned phrase: {bad}")

ri = [i for i, x in enumerate(ps) if x.strip().upper().startswith("REFERENCES")][0]
nref = len([p for p in ps[ri:] if re.match(r"^\d+\. ", p.strip())])
print("ref count:", nref)
print("SI table objects:",
      len(Document(D / "Supplemental_Information_Patterns.docx").tables))
for p in Document(D / "TitlePage_Patterns.docx").paragraphs:
    if "Manuscript statistics" in p.text or "approximately" in p.text:
        print("  TP:", p.text[:200])

print()
print("RESULT:", "ALL OK" if not fails else f"{len(fails)} FAILURES")
for f in fails:
    print("  -", f)
