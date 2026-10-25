#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""_rev_v13_s16.py — 把单位分辨率阶梯实验写成 SI Table S16，并更新封面补充表计数。

Table S16. Unit-resolution control: the sibling channel does not open when the
unit is redefined. 同一宫颈队列 / 同一嵌入 / 同一供体级标签，
仅把 unit 从「整张切片」改成「切片内的 spot 区块」，使每个供体贡献 K=1,3,8 个
sibling unit。
"""
from __future__ import annotations

import io
import json
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
RES = ROOT / "results"
D = ROOT / "papers" / "20_MRP-ST-Repair_Patterns"

payload = json.loads((RES / "unit_ladder_cervical.json").read_text(encoding="utf-8"))
runs = payload["runs"]
summary = payload["summary"]

# ------------------------------------------------------------------ 表格正文
doc = Document(D / "Supplemental_Information_Patterns.docx")

p = doc.add_paragraph()
r = p.add_run("Table S16. Unit-resolution control: redefining the unit does not "
              "open a sibling channel.")
r.bold = True

p = doc.add_paragraph()
p.add_run(
    "The cervical cohort, the spot embedding, the grouping variable and the "
    "donor-level severity label are held fixed; only the definition of the unit "
    "changes. At K = 1 the unit is the whole section mean-pooled over its spots, "
    "so every donor contributes exactly one unit and leave-one-unit-out is "
    "identical to leave-one-group-out, making the sibling-unit inflation zero by "
    "construction. At K = 3 and K = 8 the spots of each section are partitioned "
    "at random into K blocks and each block is mean-pooled into a unit, so each "
    "donor contributes K sibling units that share its label. The instrument, the "
    "dose grid, the seeds and the permutation protocol are those of Table S13. "
    "The last two columns are the point of the control: the identity classifier "
    "on the same embedding retains a large lift over chance throughout, so the "
    "absence of a sibling channel is not explained by the embedding having lost "
    "donor identity. sigma_hat is the null scale at the declared unit "
    "resolution and therefore falls as the number of units rises; LDI is "
    "comparable across cohorts at a fixed resolution, not across resolutions of "
    "the same data."
)

hdr = ["K (units per donor)", "Units", "Groups", "lambda", "seed",
       "tau_LOO", "tau_LGO", "Delta (tau_LOO - tau_LGO)", "sigma_hat", "LDI",
       "group-ID accuracy", "group-ID lift"]
tb = doc.add_table(rows=1, cols=len(hdr))
tb.style = "Table Grid"
for j, h in enumerate(hdr):
    tb.rows[0].cells[j].text = h

for row in runs:
    K = row["units_per_group"]
    cells = tb.add_row().cells
    vals = [str(K), str(row["m_units"]), str(row["G"]),
            "%.1f" % row["lam"], str(row["seed"]),
            "%+.4f" % row["tau_loo"], "%+.4f" % row["tau_lgo"],
            "%+.4f" % row["inflation"], "%.4f" % row["sigma"],
            "%.3f" % row["ldi"],
            "%.3f" % row["group_id_acc"], "%.1fx" % row["group_id_lift"]]
    for j, v in enumerate(vals):
        cells[j].text = v

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Summary by resolution. ").bold = True
p.add_run(
    "; ".join(
        "K = %s: %d units, tau_base = %.4f, sigma_hat = %.4f, LDI_base = %.2f, "
        "inflation over all doses %.4f to %.4f, dual gate accepted %d of %d dosed cells"
        % (k.lstrip("K"), s["units"], s["tau_base"], s["sigma_hat"], s["ldi_base"],
           s["inflation_range"][0], s["inflation_range"][1],
           s["dual_gate_pass_dosed"], s["dosed_cells"])
        for k, s in summary.items())
)
doc.save(D / "Supplemental_Information_Patterns.docx")

# ------------------------------------------------------------------ 封面计数
tp = D / "TitlePage_Patterns.md"
t = io.open(tp, encoding="utf-8").read()
t = t.replace("S1\u2013S15", "S1\u2013S16").replace("S1-S15", "S1-S16")
t = t.replace("15 supplemental tables", "16 supplemental tables")
t = t.replace("fifteen supplemental tables", "sixteen supplemental tables")
io.open(tp, "w", encoding="utf-8").write(t)

print("Table S16 appended; SI tables now =", len(Document(D / "Supplemental_Information_Patterns.docx").tables))
for k, s in summary.items():
    print("  %s units=%d tau_base=%.4f sigma=%.4f LDI=%.2f infl=[%+.4f,%+.4f] dual=%d/%d"
          % (k, s["units"], s["tau_base"], s["sigma_hat"], s["ldi_base"],
             s["inflation_range"][0], s["inflation_range"][1],
             s["dual_gate_pass_dosed"], s["dosed_cells"]))
