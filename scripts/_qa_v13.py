#!/usr/bin/env python
# -*- coding: utf-8 -*-
import re
from docx import Document
from pathlib import Path

D = Path(__file__).resolve().parent.parent / "papers" / "20_MRP-ST-Repair_Patterns"
d = Document(D / "Manuscript_Patterns.docx")
ps = [p.text for p in d.paragraphs]
full = "\n".join(ps)
for t in d.tables:
    for r in t.rows:
        full += "\n" + " | ".join(c.text for c in r.cells)

print("em-dash", full.count("\u2014"), "| literal **", full.count("**"),
      "| tables", len(d.tables))
print("words(all):", len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'+/-]*", full)))
ii = [i for i,x in enumerate(ps) if x.strip().upper()=="INTRODUCTION"][0]
jj = [i for i,x in enumerate(ps) if x.strip().upper().startswith("EXPERIMENTAL PROCEDURES")][0]
a = len("\\n".join(ps[:ii]))
b = len("\\n".join(ps[:jj]))
print("main text words (approx):",
      len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'+/-]*", full[a:b])))
for k in ["Control: the absence of the channel", "Table S16", "Yagis",
          "percentage points on four datasets", "one object with two components",
          "defined test rather than a fortunate run", "294 papers",
          "in practice it is the common one"]:
    print("%-48s %d" % (k, full.count(k)))
print()
print("refs:", len([p for p in ps if re.match(r"^\d{1,2}\. [A-Z\[]", p) and len(p) > 60]))
print("last ref:", [p for p in ps if re.match(r"^43\. ", p)][0][:90])
print()
print("SI tables:", len(Document(D / "Supplemental_Information_Patterns.docx").tables))
for p in Document(D / "TitlePage_Patterns.docx").paragraphs:
    if "Supplemental" in p.text or "Main text" in p.text:
        print("  TP:", p.text[:150])
print()
for p in Document(D / "CoverLetter_Patterns.docx").paragraphs:
    if "ten hospitals" in p.text or "sibling channel only because" in p.text:
        print("  CL:", p.text[:120], "...")
