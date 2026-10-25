# -*- coding: utf-8 -*-
"""Build SI Table S14 (dose-grid sigma drift / gate robustness) and append it to the SI docx."""
import io
import json
import os
import shutil

from docx import Document
from docx.shared import Pt

ROOT = r'G:\project4_cervix_multimodal'
PD = os.path.join(ROOT, 'papers', '20_MRP-ST-Repair_Patterns')
GATE = 1.645
TAU_PERP = 0.001

files = [('Head and neck (GSE268014)', 'repair_e4_hnscc_stable.json'),
         ('Prostate (GSE278936)', 'repair_prostate_gse278936_stable.json')]

rows = []
for label, fn in files:
    d = json.load(io.open(os.path.join(ROOT, 'results', fn), encoding='utf-8-sig'))
    cells = {}
    for k, v in d['scan'].items():
        p = k.split('_lam')
        cells[(p[0], float(p[1]))] = v
    seeds = sorted({s for s, l in cells})
    lams = sorted({l for s, l in cells})
    base = sum(cells[(s, 0.0)]['sigma_null'] for s in seeds) / len(seeds)
    for s in seeds:
        b = cells[(s, 0.0)]
        for l in lams:
            v = cells[(s, l)]
            sig = v['sigma_null']
            ldi_own = (v['tau_lodo'] - TAU_PERP) / sig
            ldi_base = (v['tau_lodo'] - TAU_PERP) / base
            dsig = 100.0 * (sig - base) / base
            dtau = 100.0 * (v['tau_lodo'] - b['tau_lodo']) / b['tau_lodo']
            rows.append([label, s.replace('s', 'seed '), ('%.1f' % l), '%.4f' % sig, '%.3f' % ldi_own,
                         '%.3f' % ldi_base, '%+.1f%%' % dsig, '%+.0f%%' % dtau,
                         'no' if (ldi_own < GATE) == (ldi_base < GATE) else 'YES'])

hdr = ['Cohort', 'Seed', 'lambda', 'sigma(run)', 'LDI (run sigma)', 'LDI (baseline sigma)',
       'sigma drift vs baseline', 'tau_LOO change vs baseline', 'Gate decision changed']

# ---- append to the SI docx
DOCX = os.path.join(PD, 'Supplemental_Information_Patterns.docx')
shutil.copyfile(DOCX, DOCX.replace('.docx', '_v7pre_20261024.docx'))
doc = Document(DOCX)

p = doc.add_paragraph()
p.add_run('Table S14. ').bold = True
r = p.add_run('Dose-grid drift of the null scale σ̂, and the gate-robustness check behind the pre-computability claim.')
r.bold = True
p2 = doc.add_paragraph()
p2.add_run(
    'Every seed × dose cell of the two external cohorts (30 cells). σ̂(run) is the permutation null scale estimated for that run '
    '(B = 2,000 permutations, seed 20261010); σ̂(0) is the mean baseline σ̂ of Table 1 (0.2717 head and neck, 0.1406 prostate). '
    'Column 5 is the LDI as reported throughout the paper, computed against the run\'s own σ̂; column 6 recomputes the same cell '
    'against the baseline σ̂, which is what a pre-computable reading of Proposition 1 would give. Column 9 asks whether the leakage '
    'axis of the gate (LDI < 1.645) changes between the two: it does not, for any cell. The smallest distance from a recomputed '
    'value to the gate is 0.053 (head and neck, seed 1, λ = 2) and 0.101 (prostate, seed 1, λ = 1). Column 8 gives the change in '
    'τ_LOO over the same interval; at every cell whose LDI falls below the gate the absolute change in τ_LOO exceeds the absolute '
    'change in σ̂, so no acceptance verdict is carried by the null scale rather than by the association. This table concerns '
    'dose-grid drift only; the permutation-depth recheck (B = 2,000 → 10,000, σ̂ change ≤ 4.4%, 0 of 15 runs changed decision) is a '
    'separate perturbation and is reported in Table S4.')

tb = doc.add_table(rows=1, cols=len(hdr))
tb.style = 'Table Grid'
h = tb.rows[0].cells
for i, x in enumerate(hdr):
    h[i].text = ''
    rr = h[i].paragraphs[0].add_run(x)
    rr.bold = True
for row in rows:
    cs = tb.add_row().cells
    for i, x in enumerate(row):
        cs[i].text = x

doc.save(DOCX)
print('appended Table S14 (%d data rows); tables now %d' % (len(rows), len(Document(DOCX).tables)))

# ---- title page count
TP = os.path.join(PD, 'TitlePage_Patterns.md')
s = io.open(TP, encoding='utf-8').read()
old = '13 supplemental tables (S1–S13; S13 carries the per-seed cross-domain inflation values behind the uncertainty interval reported in the text)'
new = ('14 supplemental tables (S1–S14; S13 carries the per-seed cross-domain inflation values behind the uncertainty interval '
       'reported in the text, and S14 carries the per-cell dose-grid σ̂ drift and gate-robustness check behind the '
       'pre-computability scope condition)')
if old in s:
    io.open(TP, 'w', encoding='utf-8').write(s.replace(old, new))
    print('title page updated')
else:
    print('WARN: title page pattern not found')
