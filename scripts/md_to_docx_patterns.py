#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Convert the Patterns submission markdown files to .docx (Cell Press style).

Handles: #/##/### headings, **bold**, *italic*, > blockquotes, - and N. lists,
--- separators, and math-subscript tokens (τ_base -> tau with subscript base).
Also prints word-count statistics for the title page.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# Active font selection (set per convert() call). cjk=None keeps the original
# Latin-only behavior; when set, East-Asian script runs get the CJK typeface.
_FONT = {"latin": "Arial", "cjk": None}


def _style_run(r):
    r.font.name = _FONT["latin"]
    if _FONT["cjk"]:
        rf = r._element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:ascii"), _FONT["latin"])
        rf.set(qn("w:hAnsi"), _FONT["latin"])
        rf.set(qn("w:eastAsia"), _FONT["cjk"])

BASE = Path(__file__).resolve().parents[1] / "papers" / "20_MRP-ST-Repair_Patterns"
FILES = {
    "MRP-ST-Repair_manuscript_Patterns_v1.md": "Manuscript_Patterns.docx",
    "TitlePage_Patterns.md": "TitlePage_Patterns.docx",
    "CoverLetter_Patterns.md": "CoverLetter_Patterns.docx",
}

SUB_PREFIX = set("τλσnRφα")
SUB_SUFFIX = {"base", "LOO", "in", "LDO", "lodo", "null", "samples"}
SUB_RE = re.compile(rf"([{ ''.join(SUB_PREFIX) }]̂?)_(\w+)")


def is_sub(match_prefix, match_suffix):
    return match_suffix in SUB_SUFFIX


def add_runs(par, text, base_bold=False):
    """Parse **bold**, *italic*, and subscripts into runs."""
    # tokenize bold first
    tokens = re.split(r"(\*\*.+?\*\*)", text)
    for tok in tokens:
        if not tok:
            continue
        bold = base_bold
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            tok = tok[2:-2]
            bold = True
        # italics
        for sub_tok in re.split(r"(\*[^*]+\*)", tok):
            if not sub_tok:
                continue
            italic = False
            if sub_tok.startswith("*") and sub_tok.endswith("*") and len(sub_tok) > 2:
                sub_tok = sub_tok[1:-1]
                italic = True
            # subscript tokens
            pos = 0
            for m in SUB_RE.finditer(sub_tok):
                if m.group(2) not in SUB_SUFFIX:
                    continue
                head = sub_tok[pos:m.start()]
                _emit(par, head, bold, italic)
                r = par.add_run(m.group(1))
                r.bold = bold
                r.italic = italic
                _style_run(r)
                r2 = par.add_run(m.group(2))
                r2.bold = bold
                r2.italic = italic
                r2.font.subscript = True
                _style_run(r2)
                pos = m.end()
            _emit(par, sub_tok[pos:], bold, italic)


def _black(par):
    for r in par.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)


def _emit(par, text, bold, italic):
    if not text:
        return
    r = par.add_run(text)
    r.bold = bold
    r.italic = italic
    _style_run(r)


def convert(md_path, docx_path, latin="Arial", cjk=None):
    global _FONT
    _FONT = {"latin": latin, "cjk": cjk}
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = latin
    style.font.size = Pt(11)
    if cjk:
        rf = style.element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:ascii"), latin)
        rf.set(qn("w:hAnsi"), latin)
        rf.set(qn("w:eastAsia"), cjk)
        for hs in ("Heading 1", "Heading 2"):
            hf = doc.styles[hs].element.get_or_add_rPr().get_or_add_rFonts()
            hf.set(qn("w:ascii"), latin)
            hf.set(qn("w:hAnsi"), latin)
            hf.set(qn("w:eastAsia"), cjk)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(1)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line.strip():
            continue
        if line.strip() == "---":
            continue
        if line.startswith("### "):
            doc.add_heading("", level=2)
            add_runs(doc.paragraphs[-1], line[4:].strip())
            _black(doc.paragraphs[-1])
        elif line.startswith("## "):
            doc.add_heading("", level=1)
            add_runs(doc.paragraphs[-1], line[3:].strip())
            _black(doc.paragraphs[-1])
        elif line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, line[2:].strip(), base_bold=True)
            for r in p.runs:
                r.font.size = Pt(14)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            add_runs(p, line[2:].strip())
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, line)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:].strip())
        else:
            p = doc.add_paragraph()
            add_runs(p, line.strip())

    doc.save(docx_path)
    print(f"[docx] {docx_path.name} ({len(doc.paragraphs)} paragraphs)")


def word_counts():
    md = (BASE / "MRP-ST-Repair_manuscript_Patterns_v1.md").read_text(
        encoding="utf-8")
    summary = re.search(r"## SUMMARY\n\n(.*?)\n\n", md, re.S).group(1)
    print(f"Summary words: {len(summary.split())}")

    def count_between(a, b):
        seg = md.split(a, 1)[1]
        seg = seg.split(b, 1)[0]
        # strip markdown decorations
        seg = re.sub(r"[*#>`_]", " ", seg)
        return len(seg.split())

    intro = count_between("## INTRODUCTION", "## RESULTS")
    res = count_between("## RESULTS", "## DISCUSSION")
    disc = count_between("## DISCUSSION", "## EXPERIMENTAL PROCEDURES")
    proc = count_between("## EXPERIMENTAL PROCEDURES", "## FIGURES")
    total = intro + res + disc + proc
    print(f"Intro {intro} + Results {res} + Discussion {disc} + "
          f"Exp.Proc {proc} = {total} words (main text)")


def main():
    for md_name, docx_name in FILES.items():
        convert(BASE / md_name, BASE / docx_name)
    word_counts()


if __name__ == "__main__":
    main()
