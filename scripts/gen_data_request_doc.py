"""生成宫颈空间转录组研究数据需求清单 Word 文档（v2：含数据来源科室、数据表明细、参数详情）。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


def set_cell_bg(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): color_hex,
    })
    tc_pr.append(shd)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    return h


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_bg(cell, "1F4E79")
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    if col_widths:
        for j, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[j].width = Cm(w)
    return table


def main():
    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    title = doc.add_heading("宫颈空间转录组多模态多任务学习研究\n数据需求清单", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    add_para(doc, "用途：向合作医院各科室申请研究数据（含数据来源科室、数据表明细、参数规格）", size=10)
    add_para(doc, "研究方向：HPV 通路锚定的宫颈空间转录组基因表达-病理图像多模态多任务学习", size=10)
    add_para(doc, "拟投稿期刊：Computer Methods and Programs in Biomedicine (CMPB, Q2)", size=10)
    add_para(doc, "生成日期：2026-09-06", size=10)
    doc.add_paragraph()

    # ========== 一、数据需求总览（含来源科室） ==========
    add_heading(doc, "一、数据需求总览（含数据来源科室）", level=1)
    add_para(doc, "本研究需向医院以下科室协同申请数据：病理科、检验科、妇产科、放射/影像科、信息科。")
    doc.add_paragraph()

    add_table(doc,
        ["序号", "数据类型", "来源科室", "优先级", "最小量", "理想量", "用途概述"],
        [
            ["1", "宫颈组织空间转录组(ST)", "病理科 + 妇产科", "★★★ 必选", "16例", "38例",
             "模型训练/LODO泛化验证/多阶段进展分析"],
            ["2", "临床病理信息表", "病理科 + 妇产科 + 信息科", "★★★ 必选", "同ST", "同ST",
             "样本标注/生存分析/HPV分层"],
            ["3", "匹配 bulk RNA-seq", "病理科 + 检验科/中心实验室", "★★ 建议", "15例", "30例",
             "ST信号bulk迁移验证/外部预后验证"],
            ["4", "HPV 检测原始数据", "检验科 / 病理科", "★★ 建议", "同ST", "同ST",
             "HPV通路锚定生物学真实性验证"],
            ["5", "连续切片 H&E 全玻片(WSI)", "病理科", "★ 加分", "同ST", "同ST",
             "组织病理学精细注释/可解释性"],
        ],
        col_widths=[1.0, 3.5, 3.5, 1.8, 1.2, 1.2, 4.8])
    doc.add_paragraph()

    # ========== 二、空间转录组数据 ==========
    add_heading(doc, "二、空间转录组(ST)原始数据 — 核心必选", level=1)

    add_heading(doc, "2.1 来源科室", level=2)
    add_para(doc, "• 病理科：负责组织取材、切片、H&E 染色、ST 实验（10x Visium）、Space Ranger 输出")
    add_para(doc, "• 妇产科：提供患者临床信息、随访数据、治疗记录")
    add_para(doc, "• 信息科：协助数据导出与去标识化")
    doc.add_paragraph()

    add_heading(doc, "2.2 样本类型与数量", level=2)
    add_para(doc, "平台：10x Genomics Visium（FF 新鲜冻存优先，FFPE 可接受但需注明）。每例须含匹配 H&E 图像。")
    doc.add_paragraph()

    add_table(doc,
        ["疾病阶段", "CIN分级", "最小例数", "理想例数", "说明"],
        [
            ["正常宫颈 Normal", "0", "5", "10", "含 HPV+ 正常与 HPV- 正常对照"],
            ["LSIL (CIN1)", "1", "2", "5", "当前完全缺失，必须补"],
            ["CIN2", "2", "2", "5", "当前完全缺失，必须补"],
            ["HSIL (CIN3)", "3", "2", "5", "当前仅1例"],
            ["浸润性鳞癌 SCC", "4", "3", "8", "当前2例"],
            ["浸润性腺癌 ADC", "4", "2", "5", "当前仅1例"],
            ["合计", "—", "16", "38", "含正常对照"],
        ],
        col_widths=[4.5, 2.0, 2.0, 2.0, 6.0])
    doc.add_paragraph()

    add_heading(doc, "2.3 每例样本文件清单（数据表明细）", level=2)
    add_para(doc, "标准 10x Visium Space Ranger 输出目录结构，每例一个目录：")
    doc.add_paragraph()

    add_table(doc,
        ["文件/目录", "数据内容/参数", "格式", "必选", "大小参考"],
        [
            ["filtered_feature_bc_matrix.h5",
             "spots×genes 稀疏表达矩阵（含 barcode、feature、counts）",
             "HDF5 (.h5)", "✅", "5–30 MB"],
            ["raw_feature_bc_matrix.h5",
             "原始未过滤矩阵（含全部 barcode）",
             "HDF5 (.h5)", "⭕", "50–200 MB"],
            ["spatial/tissue_hires_image.png",
             "H&E 高分辨率组织图像", "PNG", "✅", "10–100 MB"],
            ["spatial/tissue_lowres_image.png",
             "H&E 低分辨率组织图像", "PNG", "✅", "1–5 MB"],
            ["spatial/scalefactors_json.json",
             "像素↔微米映射、spot 半径等参数", "JSON", "✅", "<1 KB"],
            ["spatial/tissue_positions_list.csv",
             "每 spot 坐标：barcode,in_tissue,array_row,array_col,pxl_row,pxl_col",
             "CSV", "✅", "1–5 MB"],
            ["metrics_summary.csv",
             "Space Ranger 质控指标：总 spots、中位 genes/spot、reads/spot 等",
             "CSV", "⭕", "<1 MB"],
            ["web_summary.html", "质控网页报告", "HTML", "⭕", "1–2 MB"],
        ],
        col_widths=[4.5, 5.0, 2.0, 1.2, 2.3])
    doc.add_paragraph()

    add_heading(doc, "2.4 数据参数与训练用途", level=2)
    add_table(doc,
        ["训练目标", "使用的文件/参数", "用途说明"],
        [
            ["L_hpv (HPV通路一致性)",
             "h5 表达矩阵中 HPV 通路基因 + HPV 状态标签",
             "使高风险区域具 HPV 生物学解释"],
            ["L_sup (病变监督)",
             "histology_subtype / cin_grade 临床标注",
             "监督模型学习病变 vs 正常区分"],
            ["L_prog (进展有序)",
             "cin_grade (0→4) 有序标签",
             "学习 Normal→CIN→癌的进展梯度"],
            ["L_recon (自监督重建)",
             "h5 表达矩阵全基因",
             "约束隐空间保留表达信息"],
            ["L_neigh (空间邻域)",
             "tissue_positions_list.csv 坐标",
             "相邻 spot 风险分数平滑，提升 Moran's I"],
            ["H&E 形态编码",
             "tissue_hires_image.png + scalefactors",
             "UNI 基础模型提取形态特征，与基因表达交叉注意力融合"],
            ["LODO 泛化验证",
             "全部 ST 样本逐例 holdout",
             "评估模型对 unseen 疾病样本的风险排序"],
        ],
        col_widths=[3.5, 4.5, 7.0])
    doc.add_paragraph()

    # ========== 三、临床病理信息表 ==========
    add_heading(doc, "三、临床病理信息表 — 核心必选", level=1)

    add_heading(doc, "3.1 来源科室", level=2)
    add_para(doc, "• 病理科：组织学亚型、CIN 分级、FIGO 分期、病理报告")
    add_para(doc, "• 妇产科：年龄、治疗状态、随访生存数据")
    add_para(doc, "• 检验科：HPV 检测结果")
    add_para(doc, "• 信息科：数据整合、去标识化、导出")
    doc.add_paragraph()

    add_heading(doc, "3.2 数据表字段明细（参数清单）", level=2)
    add_para(doc, "格式：CSV 或 Excel（UTF-8），每例一行，sample_id 须与 ST 目录名一致。")
    doc.add_paragraph()

    add_table(doc,
        ["字段名", "数据类型", "来源科室", "必选", "取值/说明"],
        [
            ["sample_id", "字符串", "病理科", "✅", "与 ST 目录名完全一致"],
            ["patient_id", "字符串", "信息科", "✅", "去标识化患者 ID"],
            ["age", "整数", "妇产科", "✅", "年龄（岁）"],
            ["histology_subtype", "枚举", "病理科", "✅", "Normal/LSIL/HSIL/SCC/ADC"],
            ["cin_grade", "0–4 整数", "病理科", "✅", "0=正常,1=CIN1,2=CIN2,3=CIN3,4=浸润癌"],
            ["figo_stage", "枚举", "病理科", "⭕", "I/II/III/IV"],
            ["hpv_status", "+/-", "检验科", "✅", "HPV 阳性/阴性"],
            ["hpv_type", "字符串", "检验科", "⭕", "16/18/31/33/52/58/其他"],
            ["tissue_site", "枚举", "病理科", "⭕", "ectocervix/endocervix/transformation zone"],
            ["treatment_naive", "yes/no", "妇产科", "⭕", "是否治疗前样本"],
            ["os_months", "浮点数", "妇产科", "⭕", "总生存期（月）"],
            ["os_event", "0/1", "妇产科", "⭕", "1=死亡, 0=删失"],
            ["pfs_months", "浮点数", "妇产科", "⭕", "无进展生存期（月）"],
            ["pfs_event", "0/1", "妇产科", "⭕", "1=进展/死亡, 0=删失"],
            ["pathology_report", "文本", "病理科", "⭕", "病理报告摘要"],
        ],
        col_widths=[3.2, 1.8, 2.5, 1.2, 6.3])
    doc.add_paragraph()

    add_heading(doc, "3.3 字段用途", level=2)
    add_table(doc,
        ["用途", "用到的字段", "分析内容"],
        [
            ["样本标注", "histology_subtype, cin_grade", "L_sup、L_prog 损失监督信号"],
            ["HPV 分层", "hpv_status, hpv_type", "HPV 通路锚定验证、按型别分层"],
            ["生存分析", "os_months, os_event", "Cox 比例风险模型"],
            ["预后分析", "pfs_months, pfs_event", "无进展生存分析"],
            ["多变量校正", "age, figo_stage", "年龄/分期校正 Cox"],
        ],
        col_widths=[3.0, 4.0, 8.0])
    doc.add_paragraph()

    # ========== 四、bulk RNA-seq ==========
    add_heading(doc, "四、匹配 bulk RNA-seq — 强烈建议", level=1)

    add_heading(doc, "4.1 来源科室", level=2)
    add_para(doc, "• 病理科：组织取材、RNA 提取")
    add_para(doc, "• 检验科 / 中心实验室：测序（Illumina 平台）")
    add_para(doc, "• 信息科：表达矩阵导出")
    doc.add_paragraph()

    add_heading(doc, "4.2 数据表与参数", level=2)
    add_table(doc,
        ["数据项", "参数/格式要求", "说明"],
        [
            ["样本来源", "与 ST 同队列患者的术前活检或术后组织", "匹配优先"],
            ["最小/理想例数", "15 例 / 30 例", "—"],
            ["测序平台", "Illumina NovaSeq / HiSeq", "标准 RNA-seq"],
            ["表达矩阵格式", "gene × sample，CSV/TSV", "raw counts / TPM / FPKM"],
            ["基因 ID", "Gene Symbol 或 Ensembl ID", "须可与 ST 基因对齐"],
            ["质控报告", "FastQC / MultiQC、比对率、基因检出数", "评估数据质量"],
            ["临床匹配表", "同第三节（含生存数据更佳）", "sample_id 须一致"],
        ],
        col_widths=[3.5, 6.0, 5.5])
    doc.add_paragraph()

    add_heading(doc, "4.3 数据链与用途", level=2)
    add_table(doc,
        ["数据链", "分析内容"],
        [
            ["ST 模型表达分支 → bulk 打分",
             "冻结 ST 模型 GeneEncoder+risk head，对 bulk 表达按基因面板对齐后直接打分，Cox 生存分析（不翻转方向）"],
            ["HPV 通路活性签名 → bulk 预后",
             "在 bulk 内做 HPV+ vs HPV− 差异分析构建 352 基因签名，验证 HPV 通路活性的独立预后价值"],
            ["ST 信号复现性验证",
             "比较 ST per-spot 风险分数与 bulk 风险分数的一致性"],
        ],
        col_widths=[5.5, 9.5])
    doc.add_paragraph()

    # ========== 五、HPV 检测 ==========
    add_heading(doc, "五、HPV 检测原始数据 — 建议", level=1)

    add_heading(doc, "5.1 来源科室", level=2)
    add_para(doc, "• 检验科：HPV 分型 PCR、E6/E7 mRNA 检测")
    add_para(doc, "• 病理科：p16 IHC 免疫组化")
    doc.add_paragraph()

    add_heading(doc, "5.2 数据表与参数", level=2)
    add_table(doc,
        ["检测类型", "数据表/格式", "参数/字段", "优先级"],
        [
            ["HPV 分型 PCR", "Excel/CSV",
             "sample_id, hpv_type(16/18/31/33/52/58/...), hpv_status(+/-), ct_value", "★★★"],
            ["HPV E6/E7 mRNA", "Excel/CSV",
             "sample_id, e6_expression, e7_expression, hpv_type", "★★★"],
            ["HPV DNA 载量", "Excel/CSV",
             "sample_id, viral_load(copies/μL), hpv_type", "★★"],
            ["p16 IHC", "扫描图像 + 报告",
             "sample_id, p16_score(0/1+/2+/3+), 染色图像", "★★"],
        ],
        col_widths=[3.0, 3.0, 7.0, 1.5])
    doc.add_paragraph()

    add_heading(doc, "5.3 用途", level=2)
    add_para(doc, "验证 HPV 通路锚定的生物学真实性；按 HPV 亚型（16/18/其他）分层分析模型风险分数；"
                  "区分 HPV 驱动与非 HPV 驱动肿瘤的通路差异。")
    doc.add_paragraph()

    # ========== 六、H&E WSI ==========
    add_heading(doc, "六、连续切片 H&E 全玻片图像(WSI) — 加分项", level=1)

    add_heading(doc, "6.1 来源科室", level=2)
    add_para(doc, "• 病理科：切片制备、H&E 染色、全玻片扫描、病理注释")
    doc.add_paragraph()

    add_heading(doc, "6.2 数据表与参数", level=2)
    add_table(doc,
        ["数据项", "参数/格式", "说明"],
        [
            ["来源", "与 ST 切片相邻的连续切片", "保证组织区域对应"],
            ["格式", "SVS / NDPI / TIFF / MRXS", "标准 WSI 格式"],
            ["扫描倍率", "20× 或 40×", "—"],
            ["数量", "与 ST 样本同例数", "—"],
            ["病理区域标注", "XML/JSON 或注释软件导出",
             "肿瘤区/基质区/免疫区/坏死区/正常区（如有）"],
        ],
        col_widths=[3.5, 5.0, 6.5])
    doc.add_paragraph()

    add_heading(doc, "6.3 用途", level=2)
    add_para(doc, "更精细的组织病理学注释；模型可解释性分析（风险分数 × 病理区域叠加）；"
                  "验证高风险区域与病理形态的对应关系。")
    doc.add_paragraph()

    # ========== 七、交付规格汇总 ==========
    add_heading(doc, "七、数据交付规格汇总", level=1)
    add_table(doc,
        ["数据项", "来源科室", "数量", "格式", "大小预估", "优先级"],
        [
            ["ST 样本", "病理科+妇产科", "16–38 例", "10x Visium 目录", "1–10 GB", "★★★"],
            ["临床病理表", "病理科+妇产科+信息科", "1 份", "CSV/Excel", "<1 MB", "★★★"],
            ["bulk RNA-seq", "病理科+检验科", "15–30 例", "表达矩阵 CSV", "50–200 MB", "★★"],
            ["HPV 检测", "检验科+病理科", "同 ST 例数", "Excel/CSV", "<1 MB", "★★"],
            ["H&E WSI", "病理科", "同 ST 例数", "SVS/NDPI/TIFF", "50–200 GB", "★"],
        ],
        col_widths=[2.5, 3.5, 2.5, 2.5, 2.0, 1.5])
    doc.add_paragraph()

    # ========== 八、科室协同分工 ==========
    add_heading(doc, "八、各科室协同分工说明", level=1)
    add_table(doc,
        ["科室", "提供数据", "负责内容"],
        [
            ["病理科", "ST 原始数据、H&E 图像、WSI、病理诊断、p16 IHC",
             "组织取材、切片、ST 实验、病理诊断、图像扫描"],
            ["妇产科", "年龄、治疗记录、随访生存数据、样本采集",
             "患者入组、临床信息、随访"],
            ["检验科", "HPV 分型、E6/E7 mRNA、bulk RNA-seq 测序",
             "HPV 检测、RNA 提取、测序"],
            ["中心实验室", "bulk RNA-seq（如无检验科）", "测序与质控"],
            ["信息科", "数据整合、去标识化、导出",
             "数据脱敏、格式转换、伦理合规"],
            ["放射/影像科", "如需影像组学（可选）", "MRI/CT 影像（本研究暂不需）"],
        ],
        col_widths=[3.0, 5.5, 6.5])
    doc.add_paragraph()

    # ========== 九、伦理与隐私 ==========
    add_heading(doc, "九、伦理与隐私说明", level=1)
    add_para(doc, "1. 所有数据严格去标识化，不含患者姓名、身份证号、住院号等直接标识符。")
    add_para(doc, "2. 仅用于学术研究，不用于商业用途。")
    add_para(doc, "3. 数据存储于加密服务器，访问受密码与权限控制。")
    add_para(doc, "4. 可配合提交伦理审查申请材料。")
    doc.add_paragraph()

    # ========== 十、联系方式 ==========
    add_heading(doc, "十、联系方式", level=1)
    add_para(doc, "申请人：________________________")
    add_para(doc, "单位：________________________")
    add_para(doc, "邮箱：________________________")
    add_para(doc, "电话：________________________")
    add_para(doc, "日期：2026 年 9 月 6 日")

    out = r"f:\project4_cervix_multimodal\papers\06_CMPB_最终版_v2_20260905\数据需求清单_医院合作_v2.docx"
    doc.save(out)
    print(f"Saved -> {out}")


if __name__ == "__main__":
    main()
