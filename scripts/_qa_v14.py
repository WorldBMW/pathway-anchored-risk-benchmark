#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""_qa_v14.py — v14 稿件 QA：禁字符、表数、引用、新实验段、字数。"""
import re
from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parent.parent / "papers" / "20_MRP-ST-Repair_Patterns"
d = Document(D / "Manuscript_Patterns.docx")
ps = [p.text for p in d.paragraphs]
full = "\n".join(ps)
for t in d.tables:
    for r in t.rows:
        full += "\n" + " | ".join(c.text for c in r.cells)

WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9'+/-]*")
print("em-dash", full.count("\u2014"), "| literal **", full.count("**"),
      "| main tables", len(d.tables))
print("words(all, incl. SI captions and refs):", len(WORD.findall(full)))

ii = [i for i, x in enumerate(ps) if x.strip().upper() == "INTRODUCTION"][0]
jj = [i for i, x in enumerate(ps)
      if x.strip().upper().startswith("EXPERIMENTAL PROCEDURES")][0]
body = "\n".join(ps[ii:jj])
print("main text words (Introduction -> Experimental Procedures):",
      len(WORD.findall(body)))

checks = ["Control: the channel is opened by coupling",
          "Table S17", "Table S16", "44. [Dataset] Axelrod", "GSE316098",
          "6.55x chance in all three",
          "0 of 12 dosed seed x dose cells passed it in each",
          "structural answer to why",
          "three accepted cells among our 18 dosed clinical cells",
          "Cervix ROI-level, GSE316098"]
for k in checks:
    print("  %-58s %d" % (k, full.count(k)))

print("\nref count:", len([p for p in ps
                           if re.match(r"^\d{1,2}\. (\[Dataset\]|[A-Z])", p)
                           and len(p) > 60]))
print("ref 44:", [p for p in ps if p.strip().startswith("44. [Dataset]")][0][:100])

# Table 1 三行
for t in d.tables:
    hdr = " | ".join(c.text for c in t.rows[0].cells)
    if "Critical retention" in hdr:
        print("\nTable 1 rows:", len(t.rows))
        for r in t.rows[1:]:
            print("   ", " | ".join(c.text[:38] for c in r.cells))
        break

si = Document(D / "Supplemental_Information_Patterns.docx")
print("\nSI table objects:", len(si.tables))
print("SI S17 header present:",
      any("Within-group label variation" in p.text for p in si.paragraphs))
for p in Document(D / "TitlePage_Patterns.docx").paragraphs:
    if "supplemental tables" in p.text:
        print("TP stats:", p.text[:200])
