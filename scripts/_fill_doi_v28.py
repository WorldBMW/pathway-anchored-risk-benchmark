#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""_fill_doi_v28.py — 把已分配的 Zenodo DOI 写进 Manuscript_Patterns.docx。

Manuscript_Patterns.docx 内嵌五幅图，不能用 md -> docx 重建（会丢图），因此三处
全部就地 patch。文本一律从 v28 md 读取，避免 md 与 docx 漂移。

同步的三处：
  1. Resource availability 的 **Code.** 段
  2. Key resources table 的 Software 行（identifier 列）
  3. DATA AND CODE AVAILABILITY 段

用法：python scripts/_fill_doi_v28.py [--dry]
"""
from __future__ import annotations

import argparse
import copy
from datetime import date
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

P = Path(__file__).resolve().parents[1] / "papers" / "20_MRP-ST-Repair_Patterns"
DOC = P / "Manuscript_Patterns.docx"
SRC = P / "MRP-ST-Repair_manuscript_Patterns_v28.md"

CODE_PREFIX = "Code."
KRT_NEEDLE = "github.com/WorldBMW"
DATA_START = "The public source accessions used for the transcriptomic analyses"


# ------------------------------------------------------------------ md 读取
def md_para(prefix):
    md = SRC.read_text(encoding="utf-8")
    for raw in md.split("\n\n"):
        b = raw.strip()
        if b.startswith(prefix):
            return b.replace("**", "")
    for line in md.split("\n"):
        s = line.strip()
        if s.startswith(prefix):
            return s.replace("**", "")
    raise SystemExit(f"paragraph not found in md: {prefix[:40]}")


def md_krt_row():
    md = SRC.read_text(encoding="utf-8")
    for line in md.split("\n"):
        if line.startswith("| Software |") and KRT_NEEDLE in line:
            cells = [c.strip() for c in line.strip("|").split("|")]
            return cells
    raise SystemExit("KRT software row not found in md")


# ------------------------------------------------------------------ docx 工具
def set_para_text(p, text, prefix=None):
    """替换段落文本；prefix 给定时保留首个 run 的格式作为粗体前缀。"""
    runs = p.runs
    if prefix and runs and runs[0].text.strip() == prefix:
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
        r2 = p.add_run(text[len(prefix):].lstrip())
        if runs[0]._element.rPr is not None:
            r2._element.insert(0, copy.deepcopy(runs[0]._element.rPr))
        r2.bold = False
        return
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def find_one(doc, pred, what):
    hits = [i for i, p in enumerate(doc.paragraphs) if pred(p.text.strip())]
    if len(hits) != 1:
        raise SystemExit(f"{what}: anchor matched {len(hits)} paragraphs")
    return doc.paragraphs[hits[0]]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry", action="store_true")
    args = ap.parse_args()

    code_txt = md_para("**Code.**")
    data_txt = md_para(DATA_START)
    krt_cells = md_krt_row()

    doc = Document(DOC)
    code_p = find_one(doc, lambda t: t.startswith(CODE_PREFIX + " The analysis scripts"),
                      "Code paragraph")
    data_p = find_one(doc, lambda t: t.startswith(DATA_START), "Data availability paragraph")

    krt_hits = []
    for ti, t in enumerate(doc.tables):
        for ri, r in enumerate(t.rows):
            for ci, c in enumerate(r.cells):
                if KRT_NEEDLE in c.text:
                    krt_hits.append((ti, ri, ci))
    if len(krt_hits) != 1:
        raise SystemExit(f"KRT identifier cell matched {len(krt_hits)} cells: {krt_hits}")

    if args.dry:
        print("Code   :", code_txt[:150])
        print("Data   :", data_txt[:150])
        print("KRT    :", krt_cells)
        print("cells  :", krt_hits)
        return

    bak = P / f"Manuscript_Patterns_v28doi_{date.today():%Y%m%d}.docx"
    if not bak.exists():
        import shutil
        shutil.copy2(DOC, bak)

    set_para_text(code_p, code_txt, prefix=CODE_PREFIX)
    set_para_text(data_p, data_txt)

    ti, ri, ci = krt_hits[0]
    cell = doc.tables[ti].rows[ri].cells[ci]
    cp = cell.paragraphs[0]
    runs = cp.runs
    if runs:
        runs[0].text = krt_cells[-1]
        for r in runs[1:]:
            r.text = ""
    else:
        cp.add_run(krt_cells[-1])
    for extra in cell.paragraphs[1:]:
        for r in extra.runs:
            r.text = ""

    doc.save(DOC)

    d2 = Document(DOC)
    txt = "\n".join(p.text for p in d2.paragraphs)
    txt += "\n" + "\n".join(c.text for t in d2.tables for r in t.rows for c in r.cells)
    print(f"patched {DOC.name}: images {len(d2.inline_shapes)}, tables {len(d2.tables)}")
    print("  DOI occurrences in docx:", txt.count("10.5281/zenodo.23640979"))
    print("  'DOI-minting' left:", txt.count("DOI-minting"))
    print("  'pending verification' left:", txt.count("pending verification"))
    print("  backup:", bak.name)


if __name__ == "__main__":
    main()
