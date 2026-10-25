# -*- coding: utf-8 -*-
"""gen_submission_v10.py — 投稿套件：Cover Letter / Highlights / Title Page
目标期刊：Computer Methods and Programs in Biomedicine (CMPB, Elsevier；中科院医学 2 区 / JCR Q1)。
数字读 results/paper_v10_numbers.json，与正文/图完全一致。
输出到 papers/07_CSBJ_v10/。
"""
import json, os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path

ROOT = Path(r"f:\project4_cervix_multimodal")
OUT = ROOT / "papers" / "07_CSBJ_v10"
N = json.load(open(ROOT / "results" / "paper_v10_numbers.json", encoding="utf-8"))
A, BIO, EXT, RC, SCR, TC, CO = (N['ablation'], N['biology'], N['external'],
                                N['random_control'], N['scrna'], N['tcga'], N['cohort'])

TITLE = ("A label-dependence ablation benchmark for pathway-anchored risk representation in "
         "extreme low-sample spatial transcriptomics: a cervical-lesion label-free validation "
         "protocol and cross-platform transfer case study")
RUNNING = "A label-dependence ablation benchmark for low-sample ST risk"

def _rf(run, size=11, bold=False, italic=False, en='Times New Roman'):
    run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    run.font.name = en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def para(doc, t="", size=11, bold=False, italic=False, align=None, after=6, indent=None, sa=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(sa if sa is not None else after)
    p.paragraph_format.line_spacing = 1.15
    if indent: p.paragraph_format.first_line_indent = Cm(indent)
    if t: _rf(p.add_run(t), size, bold, italic)
    return p

def base_doc():
    d = Document(); s = d.sections[0]
    s.top_margin = s.bottom_margin = Cm(2.4); s.left_margin = s.right_margin = Cm(2.6)
    return d

# ---------------------------------------------------------------- Highlights
def build_highlights():
    hl = [
        "Benchmark of label dependence for ST risk at n=4 diseased samples (cervix)",
        "MRP-ST: 5-item minimum reporting protocol with supported/suggestive grading",
        "Label removal drops Kendall tau to p=0.029-0.056 (3/4 seeds); suggestive only",
        "GeoMx 6/6 tumor<benign (sign inverted) at n=6 floor (p=0.031); no grade",
        "Image branch: coherence gain supported, ranking gain unsupported; no clinical utility",
    ]
    # 校验 Elsevier 规范：每条 <=85 字符（含空格）
    for h in hl:
        assert len(h) <= 85, f"highlight too long ({len(h)}): {h}"
    d = base_doc()
    para(d, "Highlights", 14, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    for h in hl:
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        _rf(p.add_run("•  " + h), 11)
    fp = OUT / "Highlights_v10.docx"; d.save(str(fp))
    # 同步存一份 txt 便于投稿系统粘贴
    (OUT / "Highlights_v10.txt").write_text("\n".join(hl) + "\n", encoding="utf-8")
    print("Highlights ->", fp, "| max chars =", max(len(h) for h in hl))

# ---------------------------------------------------------------- Title page
def build_titlepage():
    d = base_doc()
    para(d, TITLE, 15, True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    para(d, "Running title: " + RUNNING, 10.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    para(d, "Authors: Mingwei Bai^a, ... , [Corresponding Author]^a*", 11, after=4)
    para(d, "^a [Department / Institution, City, Country  - TO BE CONFIRMED]", 10.5, after=4)
    para(d, "* Corresponding author: reserved / to be assigned - [Name], [email], ORCID: [____]", 10.5, after=14)
    para(d, "Keywords: spatial transcriptomics; cervical lesion; HPV pathway; small-sample learning; "
            "label-dependence ablation; minimum reporting protocol; cross-platform transfer; methodological benchmark",
         11, after=10)
    para(d, "Manuscript type: Research Article (Computational Methods / Medical Informatics)", 11, after=4)
    para(d, "Target journal: Computer Methods and Programs in Biomedicine (CMPB)", 11, after=10)
    para(d, "Counts:", 11, bold=True, after=2)
    para(d, "Abstract: structured, 442 words; Main text: ~4,800 words (excluding references and tables); 5 figures; 2 tables; 48 references; 1 Supplementary Information file (Tables S1-S2, Notes S1-S3).", 11, after=10)
    para(d, "Data & code: all training scripts, configurations, result JSONs and fused embeddings are "
            "released at GitHub (URL: https://github.com/WorldBMW/pathway-anchored-risk-benchmark); data are from public databases (GEO Visium; GeoMx "
            "GSE316098; TCGA-CESC; GEO scRNA-seq GSE308792).", 11, after=6)
    para(d, "Declaration: original work, not under consideration elsewhere; all authors approved; no competing interests. First author: Mingwei Bai (CRediT: conceptualization, methodology, software, formal analysis, investigation, data curation, visualization, writing - original draft). Corresponding author reserved / to be assigned. Funding: this research received no external funding.", 11)
    fp = OUT / "TitlePage_v10.docx"; d.save(str(fp)); print("Title page ->", fp)

# ---------------------------------------------------------------- Cover letter
def build_cover():
    d = base_doc()
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    para(d, "Date: [submission date]", 11, after=12)
    para(d, "To the Editor-in-Chief and Editorial Board,", 11, after=2)
    para(d, "Computer Methods and Programs in Biomedicine", 11, bold=True, after=12)

    para(d, "Dear Editor,", 11, after=10)
    para(d, "We are pleased to submit our manuscript entitled “%s” for consideration as a "
            "Research Article in Computer Methods and Programs in Biomedicine." % TITLE,
         11, align=J, after=10)

    para(d, "Why this study. Spatial transcriptomics (ST) combined with pathology foundation "
            "models is rapidly being used to map cancer risk in situ. A methodological hazard is, "
            "however, almost never tested: public ST cohorts contain very few disease samples. In "
            "cervix, only four diseased Visium samples (one HSIL, three squamous/adenocarcinoma) "
            "are obtainable from public repositories. At this scale a supervised risk model can "
            "simply memorize the disease labels and report near-perfect in-cohort separation that "
            "does not generalize—and most published small-sample ST risk studies do not include "
            "the controls that would reveal this.", 11, align=J, after=8)

    para(d, "Honest framing before the findings. This is a methods paper built on an extreme case, not a positive "
            "diagnostic-performance paper: the cervical Visium cohort contains only four diseased samples (one HSIL, "
            "three SCC/ADC). At n=4 the credibility ceiling of any disease-ordering claim is set by the data itself, "
            "and we have graded every conclusion accordingly (supported / suggestive / unsupported). We state up "
            "front that the score has no current clinical decision utility; the value of the work is the evaluation "
            "protocol and the cautionary case, both fully reproducible.", 11, align=J, after=8)

    para(d, "What we contribute. We use an HPV-pathway-anchored expression-based ST representation (the UNI pathology "
            "image branch is included a priori as an exploratory, ablation-tested module, not a claimed contribution) "
            f"as the vehicle for a strictly unified, multi-seed label-dependence benchmark ({CO['n_samples']} samples, "
            f"{CO['n_spots']:,} spots; one cohort and one protocol for every run), and we formalize it as a minimum "
            "reporting protocol (MRP-ST) with a decision tree and a three-level conclusion-grading scheme. The key results are:", 11, align=J, after=6)
    for b in [
        f"the decisive evidence is correctly bounded, not oversold: with a true permutation test that shuffles the "
        f"severity vector across all 35 samples (B=10,000; ~2.1e5 attainable labelings-resolution is NOT limited to "
        f"4!=24), FULL disease ordering is strong (Kendall tau {A['FULL']['tau_mean']:.2f}, p~1e-4), whereas removing "
        f"the severity-label loss (NO-PROG) drops tau to {A['NOPROG']['tau_mean']:.2f}+/-{A['NOPROG']['tau_sd']:.3f} "
        f"with permutation p={A['NOPROG']['tau_perm_p_min']:.3f}-{A['NOPROG']['tau_perm_p_max']:.3f}: {A['NOPROG']['tau_perm_nsig']} of {A['NOPROG']['n_seeds']} seeds are nominally "
        f"significant and the remaining seed sits at the 0.05 boundary, so a single seed can flip the nominal verdict. Because only four independent disease events drive the effect, we grade label anchoring as "
        "SUGGESTIVE rather than confirmed-exactly the wording discipline the protocol enforces;",
        f"at label-free inference the expression score retains strong tumor-benign rank discrimination cross-platform "
        f"on GeoMx DSP using the corrected GEO SOFT Silva annotation, but the calibration sign does NOT transfer (tumor "
        f"{EXT['geomx_tumor_mean']:.2f} below benign {EXT['geomx_benign_mean']:.2f}; "
        f"core Mann-Whitney p={EXT['geomx_tn_core_p']:.1e}; {EXT['geomx_wp_n']}/{EXT['geomx_wp_n']} within-patient pairs "
        f"concordant, paired Wilcoxon p={EXT['geomx_wp_p']:.3f}). We report that at n=6 this p (2/2^6={EXT['geomx_wp_exact_two_sided_floor_n6']:.3f}) is the "
        f"smallest attainable two-sided value-6/6 is the strongest mathematically possible outcome-while the exact "
        f"95% interval is [{EXT['geomx_wp_ci_lo']:.2f},1.00], so this is strong rank discrimination but wide-uncertainty evidence with an inverted sign, graded suggestive. A post-hoc "
        "stratified analysis excludes compartment-imbalance artifact (balanced epi/stromal counts per Silva pattern), "
        f"localizes the pattern-A-lowest signal to epithelial ROIs (epi A {EXT['geomx_silva_epi_A']:.2f} < B "
        f"{EXT['geomx_silva_epi_B']:.2f} < C {EXT['geomx_silva_epi_C']:.2f}; A vs B p={EXT['geomx_silva_epi_AB_p']:.1e}; no stromal gradient), and flags the mechanism "
        "as unverified (no HPV-ISH/p16 in the dataset). Cross-organ on 14 HNSCC sections the score shows genuine "
        f"spatial organization (Moran's I {EXT['hnscc_moran_mean']:.2f}/{EXT['hnscc_gse281978_moran_mean']:.2f}; coordinate permutation "
        f"p<=0.001; KEGG panel above expression-matched random sets on all sections, sign p={EXT['rg268_exprm_signp']:.3f}/{EXT['rg281_exprm_signp']:.3f}), "
        "graded supported, but every section-level discriminative contrast is negative; the score is also NOT a grade "
        f"(Silva non-monotonic with A lowest; pathology High/Low n.s. in the risk-axis direction, p={EXT['geomx_risk_core_p']:.3f});",
        f"the exploratory image branch acts on two axes and is not claimed as multimodal fusion: it raises spatial "
        f"coherence (FULL Moran's I {A['FULL']['moran_mean']:.3f} vs NO-IMAGE {A['NOIMAGE']['moran_mean']:.3f}) but "
        f"leaves disease ranking unchanged (tau {A['FULL']['tau_mean']:.3f} vs {A['NOIMAGE']['tau_mean']:.3f}), so the "
        f"coherence gain is supported and the ranking/endpoint gain unsupported; "
        f"fully-unsupervised LODO is partial ({EXT['unsup_lodo_n_above_median']}/4, mean percentile "
        f"{EXT['unsup_lodo_mean_pct']:.2f}, ADC fold {EXT['unsup_lodo_min_pct']:.2f}) and the semi-supervised protocol gives "
        f"{EXT['lodo_n_above_median']}/4 (mean {EXT['lodo_mean_pct']:.2f}), but the exact 95% lower bounds {EXT['lodo_4of4_ci_lo']:.2f} vs "
        f"{EXT['unsup_3of4_ci_lo']:.2f} overlap, so even this contrast is reported as suggestive. TCGA-CESC Cox is "
        f"null for all 11 spatial DEGs (min p={TC['min_p']:.3f}).",
    ]:
        p = d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        _rf(p.add_run("–  " + b), 11)

    para(d, "Why it matters for CMPB readers. CMPB's scope covers formal computing methods and "
            "their application in biomedical research and medical practice, and this work is "
            "squarely a methodological contribution: an operational minimum reporting protocol "
            "(MRP-ST, five required items, an explicit trigger at n_dis<=10, a decision tree, and "
            "supported/suggestive/unsupported grading), demonstrated on a strictly unified "
            "multi-seed benchmark, all released as code and results. The message is constructive: "
            "we convert a discussion-section warning about label memorization into a line-item "
            "checklist that reviewers can apply to any small-sample ST risk manuscript, and we show "
            f"what happens when it is omitted-in-sample AUC=1.00 collapses to boundary-level "
            f"evidence (p={A['NOPROG']['tau_perm_p_min']:.3f}-{A['NOPROG']['tau_perm_p_max']:.3f}) after the label-removal ablation. We deliberately position "
            "the null results as part of the product: the image branch adds spatial coherence but no "
            "disease ranking, TCGA survival is null, the score does not transfer its sign and is not "
            "a grade, and it has no current "
            "clinical decision utility. For a readership that builds and evaluates computational "
            "methods, an honest benchmark with openly released code, exact-test resolution stated "
            "up front, and explicitly graded claims is, we believe, precisely the reusable "
            "contribution CMPB seeks.", 11, align=J, after=8)

    para(d, "Standards. Novelty: label leakage itself is not new [Vabalas et al. 2019; "
            "Boulesteix 2015, cited]; the novelty is operationalizing these principles as a "
            "network-agnostic reporting protocol for spatial transcriptomics, together with the "
            "first label-dependence ablation benchmark in the extreme n=4 ST regime and "
            "label-free cross-platform/cross-organ pathway transfer. Soundness: all runs share one "
            f"cohort and protocol across seeds; every number traces to released results JSON/npz "
            "files; permutation tests use the full 35-sample label vector, exact-test floors and "
            "Clopper-Pearson intervals are reported for every tiny-n count, and power-limited "
            "claims are downgraded rather than stretched. Reproducibility: all data are public and "
            "all code, configs, result JSONs and fused embeddings are released on GitHub (https://github.com/WorldBMW/pathway-anchored-risk-benchmark). "
            "Significance: the work defines how small-sample ST risk studies should be reported "
            "and directly guards against the overclaiming most likely to mislead this field.",
         11, align=J, after=8)

    para(d, "This manuscript is original, has not been published, and is not under consideration "
            "elsewhere. All authors have read and approved the submission and declare no competing "
            "interests. We thank the reviewers and editor for their time and welcome any suggestions "
            "for additional analyses.", 11, align=J, after=14)

    para(d, "Sincerely,", 11, after=2)
    para(d, "[Corresponding author name]", 11, after=2)
    para(d, "on behalf of all co-authors", 10.5, italic=True, after=2)
    para(d, "[Affiliation, email, ORCID]", 10.5)
    fp = OUT / "CoverLetter_v10.docx"; d.save(str(fp)); print("Cover letter ->", fp)

if __name__ == "__main__":
    build_highlights()
    build_titlepage()
    build_cover()
    print("DONE.")
