from pathlib import Path
import sys,json,re
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
import submission_reaudit_20261016 as M
t='Here C and D count concordant and discordant sample pairs, respectively; Tₛ and Tᵧ count pairs tied only in score and only in annotation. Pairs tied in both variables are excluded from these counts.'
for path in M.OUT.rglob('Manuscript_*.docx'):
    d=Document(path)
    if not any(p.text==t for p in d.paragraphs):
        p=next(p for p in d.paragraphs if p.text.startswith('The label-free NOALL scalar head is nearly constant'))
        node=OxmlElement('w:p');p._p.addprevious(node);Paragraph(node,p._parent).add_run(t);d.save(path)
        md=path.with_suffix('.md');s=md.read_text(encoding='utf-8');s=s.replace('The label-free NOALL scalar head is nearly constant',t+'\n\nThe label-free NOALL scalar head is nearly constant');md.write_text(s,encoding='utf-8')
        print(path.name,flush=True)
stats=M.read(M.OUT/'_审查记录/build_stats.json')
for v in stats:
    # Reset against the known pre-definition count, avoiding double increments.
    old={'BiB':3391,'BMC':3584,'CBM':3391,'CSBJ':3390,'GPB':3392,'ISCLS':3390}[v['code']]
    v['body_words']=old+len(re.findall(r'\b[\w]+(?:[–-][\w]+)*\b',t))
    guide=M.OUT/v['folder']/'投稿文件说明.md';s=guide.read_text(encoding='utf8');s=s.replace(f'正文约 {old} 词',f"正文约 {v['body_words']} 词");guide.write_text(s,encoding='utf8')
for p in [M.QA/'build_stats.json',M.OUT/'_审查记录/build_stats.json']:p.write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf8')
