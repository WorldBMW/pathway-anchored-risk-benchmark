"""Rebuild JBI submission from audited results; original generators are preserved."""
from pathlib import Path
import json,re,types,os,shutil,copy,zipfile,hashlib
import numpy as np
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];R=ROOT/'results';P=ROOT/'papers/08_JBI';Q=P/'_qa_20260924';Q.mkdir(exist_ok=True)
D=json.loads((R/'jbi_submission_audit/rank_perm_recomputed.json').read_text())
N=json.loads((R/'jbi_submission_audit/numbers.json').read_text())
for v,a in N['ablation'].items():
 a.setdefault('moran_global_mean',a['moran_mean']);a.setdefault('moran_global_sd',a['moran_sd'])
 a['moran_mean']=a['within_moran_mean'];a['moran_sd']=a['within_moran_sd']
(R/'jbi_submission_audit/numbers.json').write_text(json.dumps(N,indent=2),encoding='utf-8')
A=N['ablation'];E=N['external'];B=N['baseline']
TITLE='Auditing label dependence in small-sample spatial transcriptomics: a pathway-anchored benchmark in cervical lesions'
TITLE_CN='小样本空间转录组中的标签依赖审计：以宫颈病变为例的通路锚定基准'
def paragraphs(doc):
 yield from doc.paragraphs
 for t in doc.tables:
  for row in t.rows:
   for c in row.cells:yield from c.paragraphs
def settext(p,text):
 if p.text==text:return
 prop=copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear();r=p.add_run(text)
 if prop is not None:r._r.insert(0,prop)
def change(doc,prefix,text,required=True):
 found=[p for p in doc.paragraphs if p.text.startswith(prefix)]
 if not found and prefix.startswith('图 '):found=[p for p in doc.paragraphs if p.text.startswith(prefix.replace(' ',''))]
 if required and len(found)!=1:raise ValueError((prefix,len(found)))
 for p in found:settext(p,text)
def rep(doc,a,b):
 for p in paragraphs(doc):
  if a in p.text:settext(p,p.text.replace(a,b))
def remove(p):p._element.getparent().remove(p._element)
EN={
'Objective.':"Objective. Spatial-transcriptomic risk models can fit many spots from very few diseased samples. We developed an audit of label dependence that separates disease ordering from spatial autocorrelation, using cervical lesions as a small-sample case study.",
'Methods.':"Methods. We analyzed 35 cervical Visium sections (31 normal, four diseased; 46,820 spots) using pathway-anchored expression and frozen morphology features. Seven multi-seed variants removed ordinal supervision, an expression-derived proxy, the pathway regularizer, or individual model components. Analyses included section-level Kendall τ, fixed-score label permutation, within-section Moran's I, paired leave-one-disease-out evaluation, random-gene controls, and external GeoMx, head-and-neck, bulk RNA-sequencing, and single-cell datasets.",
'Results.':"Results. Mean Kendall τ decreased from 0.460 with all terms active to 0.251 after removing ordinal labels (one-sided permutation p=0.030–0.053), 0.210 after also removing proxy supervision (p=0.019–0.155), and −0.105 after additionally removing pathway regularization (p=0.623–0.929). Spatial autocorrelation persisted: mean within-section Moran's I was 0.744 for FULL and 0.890 for NO-ALL. The morphology branch produced little change in section-level ordering or within-section autocorrelation. GeoMx tumor scores were lower than matched benign scores in six paired patients, reversing the training direction. Head-and-neck transfer retained spatial structure but showed no significant discrimination in the assessed contrasts. Exploratory TCGA-CESC analyses provided no prognostic evidence.",
'Conclusion.':"Conclusion. Spatial coherence and in-sample separation do not establish disease-risk generalization. The audit exposes dependence on supervision and separates exploratory transfer findings from clinical claims. We propose MRP-ST as a five-item reporting checklist; its thresholds and broader applicability require validation. The evaluated score is not ready for clinical use.",
'A systematically underestimated':"A central methodological problem is the mismatch between the number of spots and independent tissue samples. The public cervical Visium datasets assembled here contained four diseased sections (one high-grade squamous intraepithelial lesion, two squamous-cell carcinomas, and one adenocarcinoma) and 31 normal sections. Thousands of spots do not create additional independent disease samples. An in-sample section-level AUC of 1.00 therefore cannot establish generalization. Removing supervision can reveal target dependence, although dependence alone does not prove memorization.",
'This paper makes four':"This study combines four audit components. First, seven multi-seed variants separate spatial autocorrelation from disease ordering, including sequential removal of ordinal supervision, an expression-derived proxy, and pathway regularization. Second, fixed-score section-level permutation, paired supervised and target-free hold-outs, and random-gene controls expose different sources of optimistic interpretation. Third, external GeoMx cervical adenocarcinoma [16,17] and head-and-neck ST [18,19] analyses distinguish score direction, group separation, and ordinal pattern association. Fourth, MRP-ST summarizes these checks in five reporting items and grades conclusions as supported, suggestive, or unsupported by the available evidence. The contribution is this combined evaluation workflow, demonstrated in one architecture and disease setting.",
'The cervical Visium cohort':"The cervical Visium cohort comprises 35 sections from three sources: four GSE208654 sections (N_HPV, normal; HSIL_HPV, HSIL; CA_HPV_1, squamous-cell carcinoma; CA_HPV_2, adenocarcinoma), 30 normal ectocervix sections from the GSE217237 SuperSeries, and the 10x Genomics FFPE cervical-cancer demonstration section (cervical_cancer_1, squamous-cell carcinoma). Severity was encoded as 0/1/2 for normal/HSIL/cancer. The analysis retained 46,820 spots and a unified 4,991-gene input panel. The 108-gene KEGG-HPV list had 107 detected genes in the exported expression subset. The loader selects variable genes and applies log1p to counts. Within-section z-standardization was used for downstream associations, not as a generic model-input preprocessing step. Supplementary Table S5 lists the section identifiers and spot counts.",
'Registration quality control':"Registration quality control. The loader maps full-resolution capture coordinates into the high-resolution image using tissue_hires_scalef before extracting patches. Previously identified coordinate-scaling errors motivated this correction. Frozen UNI feature checks are reported in Supplementary Note S1. H&E panels provide histologic context without score overlays. Unique patch counts detect gross extraction failures but do not independently establish histologic registration accuracy.",
'The model has three routes':"The model combines expression and frozen UNI features [12] (Figure 1b). Its loss contains reconstruction (weight 0.2), score-neighbor agreement (0.3), KEGG-HPV consistency [15] (1.0), proxy binary cross-entropy on both the score head and a fused-representation probe (5.0 each), severity regression (1.5), and a score-variance floor (0.1). The last term penalizes batch score standard deviations below 0.2; it is not a latent-distribution prior. The proxy thresholds each spot's mean detected HPV-panel expression at the median of the training spots. It is neither an HPV-infection measurement nor an independently collected disease label. The full-cohort median uses all 46,820 spots; LODO computes it after excluding the held-out section. AdamW used learning rate 5×10⁻⁴, weight decay 10⁻⁴, batch size 64, and up to 40 epochs. Early stopping monitors active training supervision/probe objectives with patience six; target-free arms use the full epoch budget. There is no held-out validation split. Seeds assess optimization variability and do not add biological replication.",
'Architecturally,':"The expression encoder is a multilayer perceptron (4,991 → 256 → 128; ReLU; dropout 0.2). Frozen 1,024-dimensional UNI features are adapted to 128 dimensions. Four-head attention combines a gene query with an image key/value, followed by a residual connection and layer normalization. The score head is 128 → 64 → 1 and decoder 128 → 256 → 4,991. NO-IMAGE substitutes a zero image representation while retaining fusion layers. This tests the branch in this implementation, not a capacity-matched comparison of all possible morphology models.",
'The benchmark is a fixed':"The audit uses seven variants on the same cohort. FULL activates all terms. NO-IMAGE zeroes the image representation; NO-HPV sets λ_hpv=0; NO-NEIGH sets λ_neigh=0. The nested cascade comprises NO-PROG (λ_prog=0), NO-SUP (λ_prog=λ_sup=0), and NO-ALL (λ_prog=λ_sup=λ_hpv=0). NO-PROG retains the proxy, and NO-SUP retains pathway regularization. NO-ALL retains reconstruction, neighborhood smoothing, the score-variance floor, morphology, and all expression inputs. It removes explicit disease-related targets and pathway regularization, not all biological information. FULL uses seeds 0, 3, and 4; NO-PROG uses 0–3; other variants use 0–2. Comparisons involving FULL are not fully seed-paired. The fourth NO-PROG seed assessed sensitivity near the significance threshold. This is a component-ablation study, not a full factorial experiment.",
'Whole-sample hold-out protocol.':"Whole-section hold-out protocol. LODO excludes each diseased section before training and gene-panel construction. Each fold trains on 31 normal and three diseased sections. The supervised protocol retains ordinal and proxy losses; the target-free protocol sets both to zero while retaining the pathway term. The contrast therefore combines removal of two supervision terms, not a dose of ordinal labels alone. The held-out disease score is ranked against 31 training-normal scores; these are not independent test controls. Four folds share most training samples and are statistically dependent.",
'Why each contrast is necessary':"Interpretation of contrasts. Removing a target tests dependence of observed ordering on that target under the specified procedure. A smaller τ does not establish memorization, and a non-significant positive-direction test does not establish equivalence to zero. Coherence is evaluated separately because smooth scores may lack disease ordering. NO-ALL lacks a disease-based orientation constraint; its score sign is not identifiable as clinical risk. Negative τ is reported without post-hoc sign reversal.",
'Two output axes':"Two output axes are reported. Moran's I [38] is computed separately on each section's six-nearest-neighbor graph after within-section centering. We average across 35 sections, then summarize these means across seeds. Legacy pooled values use global centering and include between-section mean differences; they are retained as sensitivity analyses (Supplementary Table S3). Kendall τ-b [39] relates section mean scores to severity. A one-sided positive-ordering null permutes the 31 normal, one HSIL, and three cancer labels against fixed fitted scores (B=10,000; seed 20260924). All 209,440 assignments are also enumerated as a check. Monte Carlo p=(b+1)/(B+1). Conditional p-values describe fitted-score alignment; when labels influenced fitting they are not tests of the learning algorithm or evidence of generalization. That requires null refitting or independent validation. Disease-versus-normal comparisons use one-sided Mann–Whitney U [40]. GeoMx patient differences use two-sided paired Wilcoxon tests; core/ROI p-values are exploratory because observations are nested within patients. Rank-based enrichment [41], within-section z-standardization, Spearman correlation, and Benjamini–Hochberg adjustment [42] characterize 11 programs descriptively at spot level. Analyses use NumPy [43], scikit-learn [44], and SciPy [45].",
'Null controls.':"Null controls. Thirty retrained random 108-gene panels are compared with a separate KEGG reference run (λ_rec=1.0, λ_sup=2.0, λ_prog=1.0, λ_hpv=0.5). They are size-matched draws from the training gene union, not expression-matched panels. External head-and-neck controls instead use untrained panel means and expression-decile-matched random genes. Coordinate permutation (B=1,000) tests structure within each external section. Separate morphology-only LOSO baselines predict expression-derived spot labels or reconstruct expression; their AUC is a different endpoint from section-level disease AUC. External expression-only scoring uses checkpoint_FULL_DEPLOY_FAST_seed42.pt without labels at inference. Its weights were trained with supervision; inference is not unsupervised learning. GeoMx labels follow GEO SOFT metadata, not TMA core-name prefixes.",
'All experiments run':"Training used an NVIDIA RTX 4090 (24 GB). Frozen image features can be reused, but ablation and hold-out runs require retraining. The audit recomputes statistics from saved NPZ score arrays and run JSONs; a separate manifest records provenance. The main matrix contains 22 runs and LODO adds eight; the 30 random-panel retraining runs are additional experiments. The complete audit is therefore not negligible-cost post-processing.",
'From this case we formalize':"MRP-ST is a proposed five-item reporting checklist for models with few independent diseased samples. Its trigger (≤10 diseased samples or ≤5 in a class) is pragmatic, not an established adequacy threshold. Items are (1) supervision/proxy/pathway ablations with multiple seeds; (2) section-level permutation distinguishing fixed-score from refitted tests; (3) size- and expression-matched random-gene controls where feasible, with departures disclosed; (4) whole-section hold-outs with training-only preprocessing; and (5) evidence grading and a clinical-utility statement. Supported denotes evidence within the stated design; suggestive denotes exploratory or unstable evidence; unsupported denotes a claim not established here, not proof of absence (Supplementary Table S2).",
}
EN.update({
'Figure 3.':"Figure 3. Ablation from current saved scores. A: mean within-section Moran's I, averaging 35 section-specific values per seed. B: section-level Kendall τ-b. Points denote seeds; bars show mean ± seed SD. Labels give one-sided fixed-score permutation p ranges (B=10,000). These describe alignment, not generalization. Legacy globally centered values appear in Supplementary Table S3.",
'Table 3.':"Table 3. Ablation on 35 sections and 46,820 spots. Moran's I averages section-specific statistics, then summarizes across seeds. Seed IDs and exact permutation checks appear in Supplementary Table S1. NO-ALL removes ordinal/proxy targets and pathway regularization, retaining the other losses and all inputs. P-values are one-sided conditional alignment diagnostics.",
'Removing severity-label regression':"Removing ordinal supervision reduced Kendall τ from 0.460 to 0.251±0.018. NO-PROG values were 0.266/0.229/0.244/0.266, with recomputed permutation p=0.030/0.053/0.043/0.030. Exact enumeration gave p=0.0304/0.0553/0.0440/0.0304, the same three-of-four nominal count. Removing proxy supervision as well gave NO-SUP τ=0.210±0.071 (p=0.155/0.090/0.019; one of three nominal). NO-ALL gave τ=−0.105±0.084 (p=0.713/0.929/0.623). The cascade is consistent with dependence on explicit objectives but cannot distinguish useful supervision from memorization. Non-significance does not prove absent ordering, and target-free score orientation is unconstrained. Seeds measure optimization stability, not biological replication. The small cohort and exploratory design support a suggestive interpretation.",
'Coherence is largely':"Spatial autocorrelation persisted after target removal, but its magnitude depended on the estimand. Mean within-section Moran's I was 0.744 for FULL, 0.734 for NO-IMAGE, 0.813 for NO-PROG, 0.837 for NO-SUP, and 0.890 for NO-ALL. Legacy globally centered FULL/NO-IMAGE values were 0.931/0.859. Their difference of 0.072 is not a within-section morphology effect. The corrected difference was 0.009; three incompletely paired seeds do not establish an incremental benefit. Disease ordering was also similar (τ=0.460 versus 0.457). The frozen-feature analysis independently documents spatial structure (Supplementary Note S1), but not improved prediction.",
'In-cohort metrics do not':"A learning-free mean of 107 detected HPV-panel genes gave section-level in-sample disease AUC 0.75 versus 1.00 for FULL. A separate morphology-only LOSO experiment yielded mean within-section proxy-label AUC 0.497±0.021 for ridge (16/35 folds evaluable) and approximately 0.497 for an MLP; other sections had one proxy class. These are spot-level expression-proxy endpoints, not held-out disease classification. They must not be interpreted as the same model's disease AUC falling from 1.00 to 0.50. Morphology-to-expression LOSO gave mean per-gene Pearson r=−0.015 and MSE 0.597. LODO supplies the available disease-direction assessment, subject to its training-normal reference limitation.",
'GeoMx: rank discrimination':"GeoMx: group separation with reversed score direction. Verified expression-only deployment scores cover seven patients, 93 cores, and 186 ROIs. Mean tumor-core scores were −12.98 versus −10.52 for benign cores. The patient comparison is primary: tumor means were lower than matched benign means in all six paired patients (two-sided Wilcoxon p=0.03125; descriptive concordance interval 0.54–1.00). Core-level p=2.0×10⁻⁶ and ROI-level p=1.8×10⁻⁸ are exploratory because cores/ROIs share patients. Silva means were non-monotonic (A −14.25, B −12.07, C −12.73); core A-versus-B p=0.008 and High/Low pathology-risk p=0.100 are also descriptive. Silva classifies invasion architecture, not conventional histologic grade. The findings suggest group separation but not transportable high-score/high-risk direction, probability calibration, or clinical grading. Platform, histology, normalization, and composition effects cannot be separated.",
'Head and neck:':"Head and neck: spatial structure without demonstrated discrimination. Mean model-score Moran's I was 0.76 across eight GSE268014 sections and 0.67 across six GSE281978 sections, exceeding coordinate-permutation nulls (p≤0.001). Plain KEGG-panel means exceeded expression-matched random-panel means in all sections (8/8 and 6/6; nominal two-sided sign p=0.008/0.031). These are section-level findings; donor clustering, especially in paired primary/metastatic material, limits independent replication claims. Discrimination contrasts were non-significant: malignant versus non-malignant (p=0.195), front versus core (p=0.055), HPV status (p=0.20/0.13), and primary versus lymph-node metastasis (p=1.00; Supplementary Note S3).",
'Labels stabilize a partial direction.':"LODO provided limited held-out directional evidence. Target-free folds placed three of four held-out disease sections above the training-normal median, with mean percentile 0.79; one exceeded all normals. With ordinal and proxy targets active, all four exceeded the median, three exceeded all normals, and the mean was 0.93. The adenocarcinoma fold increased from 0.26 to 0.71. Values were checked against current fold JSONs. The contrast changes two targets, shares training samples, and uses in-training normals; it is not an independent diagnostic-accuracy estimate. Histology-specific transfer is a plausible limitation, consistent with reported SCC/ADC microenvironment differences [46,47].",
'In-cohort pathway specificity':"The separate in-cohort random-panel experiment did not establish KEGG specificity. Its trained KEGG reference had globally centered Moran's I 0.962 versus mean 0.949 across 30 size-matched, retrained random-panel models (range 0.922–0.964; empirical p=0.097). These are model outputs, not plain panel scores, and this protocol differs from the main matrix. Mean disease-ordering τ was 0.462 for random runs. Proxy-AUC empirical p=0.548 is not a common disease endpoint because panel-specific targets differ across runs. External untrained, expression-matched panel comparisons address a separate question.",
'No prognostic or decision utility.':"Exploratory survival analyses provided no prognostic evidence. None of 11 spatially selected genes was significant in the univariate gene-level TCGA-CESC Cox analyses [48] (302 patients, 72 events; smallest p=0.114 for PSME2). Gene-level bulk associations do not validate the spatial score as a prognostic model. Overall, the data document spatial autocorrelation, suggest target dependence of ordering, and show limited external separability. They do not establish an incremental morphology benefit, clinical grading, prognosis, or cross-organ disease discrimination.",
'The benchmark produces':"The principal contribution is an operational audit separating frequently conflated claims. Target removal reduced positive disease ordering, while within-section spatial autocorrelation persisted. This demonstrates target dependence in this model, not proven memorization. GeoMx separation reversed the score direction; head-and-neck analyses retained structure without significant discrimination. The corrected morphology contrast was small. Separate estimands are therefore more informative than a single in-sample accuracy or coherence statistic.",
'Four lessons transfer':"Four lessons may extend to other informatics settings. First, spatial autocorrelation requires an explicit level: global centering can include section-mean differences even with section-specific graph edges. Second, fixed-score permutation after supervised fitting cannot replace null retraining or independent validation. Third, AUC values are comparable only when targets, units, and evaluation designs match. Fourth, external separation, score orientation, probability calibration, and clinical discrimination are distinct properties. MRP-ST operationalizes these distinctions alongside existing guidance [32–34,36], but its broader value requires independent evaluation.",
'(1) Power:':"(1) Sample size: four diseased sections and six paired GeoMx patients limit precision. Seeds do not add samples, LODO folds share training data, and ROIs/cores share patients. (2) Confounding: disease, cohort, histology, and platform are partly confounded. (3) Evaluation: no unbiased disease-classification AUC for the full model on independent normal and disease test samples is available. LODO uses training normals; fixed-score permutation does not test null refitting. (4) Multiple testing: analyses are exploratory; spot-level p-values do not fully account for spatial dependence, and most external tests lack multiplicity adjustment. (5) Ablations: seed sets are incompletely matched, stopping rules depend on active targets, and NO-ALL has no clinical sign anchor. The cascade cannot uniquely assign causal contributions. (6) Scope: one architecture/pathway was evaluated; random-panel retraining uses another protocol and only size matching. MRP-ST thresholds are proposals, not validated operating limits.",
'The work is an audit methodology':"Clinical translation requires independent evaluation. This audit does not establish a diagnostic, grading, or prognostic tool. Future studies should prospectively specify the task and analysis plan, recruit independent patients across sites and histologies, assess registration and platform effects, and compare against appropriate baselines. Sample size should follow precision or power calculations, not an arbitrary section count. Decision-impact studies are appropriate only after external validity and calibration are demonstrated for a defined use.",
'We presented a label-dependence':"We developed a seven-variant label-dependence audit and demonstrated it in four diseased cervical sections. Ordering depended on supervision and pathway regularization, while within-section structure persisted. External analyses suggested GeoMx separability but did not establish transportable score direction or clinical utility. Metric definitions, endpoint comparability, and result provenance are essential to an interpretable audit.",
'Applicability.':"Applicability. MRP-ST is a proposed checklist for models trained on many spatial observations from few independent samples. Adaptation to other spatial-omics or imaging tasks requires task-specific validation. The ≤10-diseased-sample trigger is a reporting aid, not a threshold above which routine cross-validation is sufficient. Ablation and hold-out components require retraining; the checklist does not replace external validation or established guidance.",
'Ethics statement.':"Ethics statement. This study reanalyzed public, de-identified datasets (GSE208654, GSE217237, GSE316098, GSE268014, GSE281978, GSE308792, TCGA-CESC, and the 10x demonstration dataset). No new participants were recruited, interventions performed, or identifiable private information accessed. Original data-collection approval and consent procedures are described by the source studies.",
'Data and code availability.':"Data and code availability. Source data are public under the accessions in Table 2 and the 10x Genomics public dataset portal. Project code: https://github.com/WorldBMW/pathway-anchored-risk-benchmark. The accompanying numerical source-data archive contains the section-level audit statistics, external summaries, exact/Monte Carlo permutation checks, and source-file manifest used for this revision. Raw data and third-party model weights remain subject to their original access conditions.",
})
def english(doc):
 settext(doc.paragraphs[0],TITLE)
 for k,v in EN.items():change(doc,k,v)
 for t in list(doc.tables):
  if t.cell(0,0).text=='Line of work':t._element.getparent().remove(t._element)
 for p in list(doc.paragraphs):
  if p.text.startswith('Table 1. Positioning'):remove(p)
 for a,b in [('Table 1 positions this work','Supplementary Table S4 positions this work'),('Table 2','Table 1'),('Table 3','Table 2'),('calibration sign','score direction'),('calibration direction','score direction'),('label memorization','label dependence'),('histologic grading system','invasion-pattern classification'),('prespecified','specified'),('pre-specified','specified'),('cross-platform validation','cross-platform evaluation'),('a priori exploratory','exploratory'),('spatial self-organization','spatial autocorrelation'),('P001–P030','30 analyzed section IDs'),('label- and prior-free','target- and pathway-regularizer-free'),('fully-unsupervised','target-free'),('fully unsupervised','target-free'),('fully label-free','target-free'),('caliber','analysis')]:rep(doc,a,b)
 t=doc.tables[0]
 for i,s in enumerate(["Many spots from very few diseased sections can yield optimistic assessments.","In-sample fit does not establish generalization; spatial dependence, cohort effects, and endpoint mismatch complicate evaluation.","A seven-variant audit combines target removal, section-level statistics, hold-outs, and external controls. It distinguishes fitted-score alignment from generalization and proposes a five-item checklist.","Researchers, reviewers, and pathology teams evaluating small-sample spatial models."],1):settext(t.cell(i,1).paragraphs[0],s)
 t=doc.tables[-1]
 readings={'FULL':'All terms active','NO-IMAGE':'Within-section I and τ close to FULL','NO-HPV':'Pathway consistency off','NO-NEIGH':'Neighborhood term off','NO-PROG':'Ordinal targets off; proxy retained','NO-SUP':'Targets off; pathway term retained','NO-ALL':'Targets and pathway term off; sign unanchored'}
 for row in t.rows[1:]:
  if row.cells[0].text in readings:settext(row.cells[-1].paragraphs[0],readings[row.cells[0].text])
 for a,b in [('4.2 Label-removal ablation: ordered ranking is label-anchored (suggestive)','4.2 Target-removal ablation and disease ordering'),('4.3 Sources of coherence, and the two-axis morphology verdict','4.3 Within-section autocorrelation and morphology ablation'),('4.5 Cross-platform and cross-organ label-free transfer','4.5 Cross-platform and cross-organ evaluation'),('4.6 Label dose-response, null controls and the clinical-utility boundary','4.6 Hold-outs, random controls, and clinical scope')]:rep(doc,a,b)
 # Fix figure 5 caption's inferential unit.
 change(doc,'Figure 5.',"Figure 5. External evaluation. A: section-level spatial autocorrelation of model scores, untrained KEGG panels, expression-matched random panels, and coordinate-permutation controls. B: GeoMx core scores; patient-level paired p=0.03125 is primary, while core-level p is exploratory. C: descriptive Silva pattern and pathology-risk comparisons; cores share patients. D: held-out disease percentiles relative to training normals under supervised and target-free LODO. The two protocols differ in both ordinal and proxy supervision.")
 rep(doc,'whether it reflects learning or label anchoring is adjudicated by the nested NO-PROG/NO-SUP/NO-ALL cascade','its dependence on explicit training targets is examined by the nested NO-PROG/NO-SUP/NO-ALL cascade')
def cover(doc):
 ps=['September 25, 2026','The Editor-in-Chief','Journal of Biomedical Informatics','Dear Editor,',f'Please consider our original research article, “{TITLE}”, for publication in the Journal of Biomedical Informatics.',
 'The manuscript addresses a biomedical-informatics evaluation problem: many spatial observations can conceal a scarcity of independent disease samples. We demonstrate a reproducible audit in 35 cervical sections, including four diseased sections.',
 'The audit combines seven ablation variants, section-level fixed-score permutation, within-section spatial statistics, whole-section hold-outs, and external analyses. It separates dependence on training targets from generalization and spatial coherence from disease ordering. MRP-ST summarizes the approach in five reporting items whose broader validity remains to be tested.',
 'Mean Kendall τ changed from 0.460 with all terms active to 0.251 after ordinal-label removal, 0.210 after additional proxy removal, and −0.105 after pathway regularization was removed. Spatial structure persisted. GeoMx tumor–benign separation reversed the training score direction; the assessed cross-organ discrimination and prognostic analyses did not establish utility.',
 'The contribution is an evaluation workflow and methodological case study with explicit treatment of statistical units, endpoint comparability, metric definitions, and uncertainty. The score is not presented as a clinical prediction tool.','Thank you for considering our manuscript.','Sincerely,','Lang Guo, on behalf of all authors','School of Pharmacy, Shenyang Medical College, 146 Huanghe North Avenue, Shenyang 110034, China','E-mail: guolang1016@symc.edu.cn']
 for p in list(doc.paragraphs):remove(p)
 for txt in ps:
  p=doc.add_paragraph(txt);p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1.05
  for r in p.runs:r.font.size=Pt(11);r.font.name='Times New Roman'
def highlights(doc):
 for p in list(doc.paragraphs):remove(p)
 doc.add_paragraph('Highlights')
 for h in ['A seven-variant audit separates disease ordering from spatial autocorrelation','Removing explicit targets weakens positive disease ordering','Within-section spatial structure persists after target removal','GeoMx tumor–benign separation reverses the training score direction','MRP-ST proposes five reporting checks for small-sample spatial models']:
  assert len(h)<=85
  doc.add_paragraph(h,style='List Bullet')
CN={
'目的。':'目的。空间转录组风险模型可在独立疾病样本极少的情况下拟合大量 spot。本研究以宫颈病变为例，建立区分疾病排序与空间自相关的标签依赖审计。',
'方法。':'方法。分析35张宫颈 Visium 切片（31张正常、4张病变；46,820个 spot），使用通路锚定表达与冻结形态学特征。七个多种子变体分别关闭有序监督、表达派生代理、通路正则或模型组件，结合切片级 Kendall τ、固定评分标签置换、切片内 Moran’s I、配对留一疾病样本、随机基因对照，以及 GeoMx、头颈部ST、bulk RNA和单细胞外部数据。',
'结果。':'结果。平均 Kendall τ由 FULL 的0.460降至去除有序标签后的0.251（单侧置换p=0.030–0.053），再去除代理后为0.210（p=0.019–0.155），再关闭通路正则后为−0.105（p=0.623–0.929）。空间自相关仍存在：FULL与NO-ALL的切片内平均 Moran’s I 分别为0.744和0.890。形态学分支对排序及切片内自相关的影响均较小。GeoMx六对患者中肿瘤评分均低于配对良性评分，与训练方向相反。头颈部保留空间结构，但所检验的判别对比未显著；TCGA-CESC探索分析未建立预后证据。',
'结论。':'结论。空间一致性和样本内分离不能证明疾病风险泛化。审计揭示训练目标依赖，并区分探索性迁移结果与临床主张。MRP-ST为五条目报告建议，其阈值和更广泛适用性仍需验证。当前评分尚不能用于临床。',
'一个被系统性低估':'核心方法学问题是spot数量与独立样本数量不匹配。纳入的公开宫颈Visium数据包含4张病变切片（1张HSIL、2张鳞癌、1张腺癌）和31张正常切片。数千spot不会增加独立疾病样本数，因此切片级样本内AUC=1.00不能建立泛化证据。去除监督可揭示训练目标依赖，但依赖本身不等于已经证明标签记忆。',
'本文做出四项':'本研究组合四类审计：（1）七个多种子变体分别评估空间自相关与疾病排序，包括有序目标、表达代理、通路正则的嵌套去除；（2）固定评分切片级置换、配对监督/目标关闭留出及随机基因对照；（3）GeoMx宫颈腺癌[16,17]与头颈部ST[18,19]外部评估，区分评分方向、组间分离和有序模式关联；（4）将流程归纳为MRP-ST五项报告建议及证据分级。贡献是组合评估流程，目前仅在一个架构和疾病场景中示范。',
'宫颈 Visium 队列汇总':'宫颈Visium队列来自三个来源：GSE208654的N_HPV、HSIL_HPV、CA_HPV_1、CA_HPV_2（分别为正常、HSIL、鳞癌、腺癌），GSE217237的30张正常宫颈外口切片，以及10x FFPE宫颈癌演示切片cervical_cancer_1（鳞癌）。严重度编码为0/1/2。共46,820个spot，统一输入4,991个基因；108基因KEGG-HPV列表中107个见于导出表达子集。加载器选择变异基因并对计数作log1p；切片内z标准化用于下游关联分析，并非模型输入的通用预处理。补充表S5列出切片ID及spot数。',
'配准质量控制':'配准质量控制。加载器使用tissue_hires_scalef将全分辨率坐标映射到高分辨率图像后提取patch，以修复先前的坐标缩放错误。冻结UNI特征检查见补充说明S1。H&E图仅展示组织学背景，不叠加模型评分。patch唯一性能够发现明显提取错误，不能独立证明组织学配准精度。',
'模型含三条路径':'模型结合表达和冻结UNI特征[12]（图1b）。损失包含重建（权重0.2）、邻域评分一致性（0.3）、KEGG-HPV一致性[15]（1.0）、评分头与融合表征探针各自的代理二元交叉熵（各5.0）、严重度回归（1.5）及评分方差下限（0.1）。方差项惩罚批内评分标准差低于0.2，不是隐空间分布先验。代理以训练spot的HPV面板平均表达中位数二分，不是HPV感染检测或独立疾病标签。完整队列使用46,820个spot；LODO先排除留出切片再计算。AdamW学习率5×10⁻⁴、权重衰减10⁻⁴、批量64、最多40轮；有监督目标时按训练目标/探针早停（耐心6），目标关闭臂训练完整轮数。无独立验证集。种子只反映优化波动。',
'架构上，':'表达编码器为4,991→256→128的MLP（ReLU、dropout 0.2）。冻结UNI的1,024维特征映射到128维，四头注意力使用基因query和图像key/value，加残差与层归一化。评分头128→64→1，解码器128→256→4,991。NO-IMAGE将图像表征置零但保留融合层；它检验当前实现中的分支，而不是所有形态模型的容量匹配比较。',
'基准是固定的':'七个变体在同一队列上运行。FULL激活全部项；NO-IMAGE将图像表征置零；NO-HPV关闭λ_hpv；NO-NEIGH关闭λ_neigh。嵌套级联为NO-PROG（λ_prog=0）、NO-SUP（λ_prog=λ_sup=0）、NO-ALL（再置λ_hpv=0）。NO-PROG保留代理，NO-SUP保留通路正则。NO-ALL仍有重建、邻域、评分方差下限、形态与全部表达输入，并非零生物信息。FULL种子为0、3、4，NO-PROG为0–3，其余为0–2；与FULL的比较并非完全配对。第四个NO-PROG种子用于考察边界稳定性。本设计是组件消融，不是全析因实验。',
'全样本留出方案。':'全切片留出。LODO先排除病变切片，再构建训练基因面板。每折训练31张正常和3张病变。监督方案保留有序和代理损失，目标关闭方案同时去除二者但保留通路项；不能解释为单独改变有序标签剂量。留出疾病评分与31张训练正常比较，它们不是独立测试对照。四折共享大部分训练数据，统计上并不独立。',
'每个对比为何必要':'对比解释。去除目标检验特定训练流程对目标的依赖。τ下降不等于证明记忆，正向单侧检验不显著不等于与零等效。NO-ALL缺乏疾病方向约束，其评分正负不能解释为已知临床风险方向；负τ直接报告，不事后翻转。',
'两条输出轴分开报告。':'分别报告两条轴。Moran’s I[38]在每张切片六近邻图上作切片内中心化后计算，先对35张切片取均值，再汇总种子。旧的全局中心化数值包含切片间均值差异，仅作为补充表S3敏感性分析。Kendall τ-b[39]关联切片均值与严重度；固定评分、置换31个正常/1个HSIL/3个癌标签，单侧检验正向排序（B=10,000，种子20260924），并穷举209,440种排列核验。Monte Carlo p=(b+1)/(B+1)。训练标签影响评分时，该条件检验只描述拟合后对齐，不能检验学习算法泛化；后者需要置换重训或独立验证。疾病/正常用单侧Mann–Whitney[40]；GeoMx患者用双侧配对Wilcoxon，芯/ROI检验仅探索性。基因程序采用实现中的秩富集[41]、切片内z标准化、Spearman及BH校正[42]，spot级p仅描述性。软件为NumPy[43]、scikit-learn[44]、SciPy[45]。',
'零模型对照。':'零模型对照。30个重新训练的随机108基因面板与单独KEGG参考运行比较（λ_rec=1.0、λ_sup=2.0、λ_prog=1.0、λ_hpv=0.5）；它们从训练基因并集中等大小抽样，未匹配表达。头颈部外部对照则采用未经训练的面板均值及表达十分位匹配随机基因。每个外部切片坐标置换B=1,000。形态学LOSO预测spot表达代理或重建表达，AUC终点不同于切片级疾病AUC。外部表达分支使用checkpoint_FULL_DEPLOY_FAST_seed42.pt，推理不输入标签但模型曾有监督训练。GeoMx病理标签由GEO SOFT获取，不能由芯名称前缀推断。',
'全部实验在':'训练在NVIDIA RTX 4090（24 GB）上完成。冻结特征可复用，但消融和留出需要重训。投稿审计从保存NPZ和JSON重算统计并记录来源。主矩阵22次运行、LODO8次，另有30次随机面板重训；不能称整套审计为几乎无成本的训练后处理。',
'基于本案例，':'MRP-ST为五项报告建议，≤10病变样本或任一类≤5是务实触发条件，并非已验证的充分性阈值：（1）多种子监督/代理/通路消融；（2）区分固定评分与置换重训；（3）尽可能进行大小和表达匹配随机对照并披露偏离；（4）训练内预处理的全切片留出；（5）证据分级和临床效用声明。supported只限于定义的设计，suggestive表示探索性或不稳定证据，unsupported表示本研究未建立该主张而非证明不存在（补充表S2）。',
'图 3  ':'图3  当前保存分数重算的消融结果。A为先汇总35张切片的切片内Moran’s I，再汇总种子；B为切片级Kendall τ-b。点为种子、柱为均值±种子SD，标注单侧固定评分置换p范围。旧全局中心化值见补充表S3；这些p不检验泛化。',
'表 3  ':'表3  35张切片、46,820个spot的消融。Moran’s I为切片内统计，先对切片平均再汇总种子。种子ID及穷举检验见补充表S1。NO-ALL保留其他损失及全部输入，p为条件对齐诊断。',
'去除严重度标签回归':'关闭有序监督后τ从0.460降至0.251±0.018；四种子τ=0.266/0.229/0.244/0.266，重算置换p=0.030/0.053/0.043/0.030。穷举p=0.0304/0.0553/0.0440/0.0304，均得到3/4名义显著。NO-SUP为0.210±0.071（p=0.155/0.090/0.019，1/3显著），NO-ALL为−0.105±0.084（p=0.713/0.929/0.623）。结果与训练目标依赖一致，但不能区分有效监督与记忆；不显著不等于排序不存在，无目标评分方向未锚定。种子是优化稳定性而非生物学重复，结论保持提示性。',
'一致性主要是数据属性。':'去除目标后空间自相关仍存在，但数值依赖指标定义。切片内均值：FULL 0.744、NO-IMAGE 0.734、NO-PROG 0.813、NO-SUP 0.837、NO-ALL 0.890。旧全局中心化FULL/NO-IMAGE为0.931/0.859，差0.072不能解释为切片内形态学效应。校正后差仅0.009，三组未完全配对种子不足以建立增益；τ也接近（0.460/0.457）。冻结特征分析证明空间结构，不证明预测改善。',
'队列内指标不能外推。':'107个检出HPV基因的免学习均值，其切片级样本内疾病AUC=0.75，FULL=1.00。另一个形态学LOSO实验中，ridge的切片内spot代理AUC=0.497±0.021（16/35折可评价），MLP约0.497，其余折仅一个代理类别。这不是留出疾病分类，不能解释为同一模型疾病AUC从1.00降到0.50。表达重建LOSO的平均逐基因r=−0.015、MSE=0.597。LODO提供疾病方向评估，但以训练正常为参照。',
'图 5  ':'图5  外部评估。A为切片级模型/纯KEGG/表达匹配随机面板/坐标零模型；B为GeoMx芯评分，主要推断为患者配对p=0.03125，芯级p为探索性；C为描述性Silva和病理风险比较；D为留出疾病对训练正常的百分位，两方案同时改变有序与代理目标。',
'GeoMx：':'GeoMx：组间分离伴评分方向反转。核验后的部署评分包括7位患者、93个芯、186个ROI。肿瘤芯均值−12.98，良性−10.52。六对患者肿瘤均低于良性（双侧配对Wilcoxon p=0.03125，描述性一致比例区间0.54–1.00）。芯p=2.0×10⁻⁶、ROI p=1.8×10⁻⁸因患者内相关仅供探索。Silva均值A−14.25/B−12.07/C−12.73不单调，A/B芯p=0.008、High/Low p=0.100也仅描述性。Silva是浸润模式而非传统组织学分级。结果未建立原高分高风险方向、概率校准或临床分级的迁移；平台、组织学、归一化与组成效应不可分离。',
'头颈部：':'头颈部：存在空间结构，但未建立判别效用。8张GSE268014与6张GSE281978切片模型Moran’s I均值0.76/0.67，坐标置换p≤0.001。纯KEGG面板高于表达匹配随机面板（8/8、6/6；名义双侧符号p=0.008/0.031）。这是切片级证据，同供者特别是原发/转移配对限制独立重复主张。恶性/非恶性p=0.195、前沿/中心p=0.055、HPV状态p=0.20/0.13、原发/转移p=1.00均未显著（补充说明S3）。',
'标签将局部方向':'LODO仅提供有限的留出方向证据。目标关闭方案3/4折高于训练正常中位、1/4高于全部正常，均值百分位0.79；监督方案为4/4、3/4和0.93。腺癌折0.26→0.71，已核对当前逐折JSON。两种监督项共同改变、折间训练重叠、正常为训练参照，故不是独立诊断准确性估计。组织学迁移局限与已有SCC/ADC微环境报道相容[46,47]。',
'队列内通路特异性':'独立随机面板实验未建立KEGG特异性。训练后KEGG参考模型全局中心化Moran’s I=0.962；30个等大小随机面板重训均值0.949（0.922–0.964，经验p=0.097）。这不是纯表达面板，协议也不同于主矩阵。随机运行平均τ=0.462。代理AUC经验p=0.548因各面板代理目标不同不能视为共同疾病终点。外部未经训练、表达匹配的面板比较回答另一问题。',
'无预后或决策效用。':'探索性生存分析未建立预后证据。11个空间选定基因在TCGA-CESC单基因单变量Cox分析[48]（302位患者、72事件）均未显著，最小PSME2 p=0.114。bulk基因关联不等于空间评分的预后验证。整体支持空间自相关、提示目标依赖及有限外部分离，但未建立形态学增益、临床分级、预后或跨器官判别。',
'本基准给出':'本研究的主要贡献是区分常被混用主张的操作性审计。去除训练目标削弱正向疾病排序，而切片内自相关持续存在；这证明当前模型目标依赖，不等于已证实记忆。GeoMx分离方向反转，头颈部保留结构但未建立判别。校正的形态学效应较小。应报告不同估计目标，而非用一个样本内准确性或一致性指标概括模型。',
'四条经验可推广':'四条经验值得进一步检验：（1）全局中心化即使边只连切片内spot，仍可包含切片均值差异；（2）监督拟合后的固定评分置换不能替代置换重训或独立验证；（3）AUC须预测目标、统计单位和评估设计一致才能比较；（4）外部分离、评分方向、概率校准及临床判别是不同性质。MRP-ST与既有指南[32–34,36]结合落实这些区分，但需独立验证。',
'（1）功效：':'局限性：（1）4张疾病切片、6对患者限制精度；种子不是新增样本、LODO训练重叠、芯/ROI患者内相关。（2）疾病、队列、组织学和平台部分混杂。（3）缺乏完整模型在独立正常及病变测试集上的无偏疾病AUC；LODO使用训练正常，固定评分置换不等于置换重训。（4）分析探索性，spot p未完全考虑空间相关，多数外部检验未多重校正。（5）种子不完全配对、早停依赖激活目标、NO-ALL无临床符号约束，不能唯一归因。（6）只检验一个架构与通路；随机重训另用协议且仅大小匹配。MRP-ST阈值尚未验证。',
'本文是审计方法学':'临床转化需要独立评价。当前审计未建立诊断、分级或预后工具。后续须预先确定任务与分析计划、多中心多组织学招募独立患者、评价配准及平台差异并与适当基线比较。样本量应按精度或功效计算，不能任意指定切片数。只有明确用途的模型建立外部效度和校准后，才适合开展决策影响研究。',
'我们提出了面向':'本研究建立七变体标签依赖审计，并在4张宫颈病变切片的场景中示范。排序依赖监督与通路正则，切片内结构保留。GeoMx提示分离，但未建立评分方向迁移或临床效用。指标定义、终点可比性及结果溯源是审计可解释性的必要条件。',
'适用性。':'适用性。MRP-ST面向独立样本少、空间观测多的模型。迁移到其他空间组学或影像任务须任务特异验证。≤10病变样本是报告提示，不意味着超过10常规交叉验证就充分。消融及留出需要重训，不能替代外部验证或现有报告指南。',
'伦理声明。':'伦理声明。本研究二次分析公开去标识化数据（GSE208654、GSE217237、GSE316098、GSE268014、GSE281978、GSE308792、TCGA-CESC和10x演示数据），未招募新参与者、实施干预或访问可识别个人信息。原始采集的伦理审批及同意程序见来源研究。',
'数据与代码可得性。':'数据与代码可得性。来源数据见表2的公开编号及10x数据门户；代码库为https://github.com/WorldBMW/pathway-anchored-risk-benchmark。随附数值源数据包包含切片级审计、外部汇总、穷举及Monte Carlo检验、来源文件清单。原始数据和第三方权重遵守其原始访问条款。',
}
def chinese(doc):
 settext(doc.paragraphs[0],TITLE_CN)
 for k,v in CN.items():change(doc,k,v)
 # Match the main English display-item count.
 for t in list(doc.tables):
  if len(t.columns)==5 and '代表' in t.rows[0].cells[1].text:t._element.getparent().remove(t._element)
 for p in list(doc.paragraphs):
  if p.text.startswith('表 1  与现有'):remove(p)
 for a,b in [('表 1 对','补充表 S4 对'),('表 2','表 1'),('表 3','表 2'),('表2','表1'),('表3','表2'),('校准符号','评分方向'),('标签记忆审计','标签依赖审计'),('成熟组织学分级体系','浸润模式分类体系'),('完全无监督','目标关闭'),('完全无标签','目标关闭'),('P001–P030','实际纳入的30张切片')]:rep(doc,a,b)
 for p in paragraphs(doc):
  if '0.07' in p.text and 'Moran' in p.text and len(p.text)<100:settext(p,'切片内Moran差约0.009；未建立增益')
 t=doc.tables[-1]
 for row in t.rows[1:]:
  key=row.cells[0].text.replace('-','')
  if key in A:settext(row.cells[-1].paragraphs[0],{'FULL':'全部损失激活','NOIMAGE':'切片内I及τ接近FULL','NOHPV':'关闭通路正则','NONEIGH':'关闭邻域项','NOPROG':'关闭有序目标；保留代理','NOSUP':'关闭两类目标；保留通路项','NOALL':'关闭目标及通路项；方向未锚定'}[key])
 t=doc.tables[0]
 for i,s in enumerate(['极少病变切片的大量空间观测可造成乐观评估。','样本内拟合不能建立泛化；空间相关、队列效应和终点不一致影响解释。','组合七变体目标去除、切片级统计、留出及外部对照，区分拟合后对齐与泛化，并提出五项报告建议。','评估小样本空间模型的研究者、审稿人及病理团队。'],1):settext(t.cell(i,1).paragraphs[0],s)
def save_final(doc,name):
 if name=='Manuscript_JBI.docx':english(doc)
 elif name=='Manuscript_JBI_中文.docx':chinese(doc)
 elif name=='CoverLetter_JBI.docx':cover(doc)
 elif name=='Highlights_JBI.docx':highlights(doc)
 for a,b in [('; the axis is not a proliferation surrogate',''),('；说明该轴不是增殖替代指标',''),('at disease n=4','with four diseased sections'),('Z. Wang, W. Yang, et al., ST-Align','Z. Wang, R. Yu, ST-Align'),('arXiv:2411.16793. https://iclr.cc/virtual/2026/10014521.','https://openreview.net/forum?id=EWzPy228cM.'),('Commun. Biol. 9 (2025) 46','Commun. Biol. 9 (2026) 46'),('187–220','187–202'),('Marker validity was confirmed in independent cervical scRNA-seq','Marker-set coherence was assessed in independent cervical scRNA-seq'),('mean within-lineage score correlation','mean within-marker-set gene correlation'),('across lineages (gap','between marker sets (gap')]:rep(doc,a,b)
 if name=='Manuscript_JBI.docx':
  change(doc,'Figure 1.','Figure 1. Cervical histologic context and model architecture. A: original H&E images from normal ectocervix (P020; GSE217237), HSIL (HSIL_HPV; GSE208654), invasive squamous-cell carcinoma (cervical_cancer_1; 10x Genomics), and invasive adenocarcinoma (CA_HPV_2; GSE208654). Lower images show the boxed regions. Images are cropped and resized for display without color adjustment; physical magnifications differ. H&E provides histologic context, not spatial validation of model scores. B: expression and frozen UNI morphology features enter a cross-modal model with score, proxy, and reconstruction outputs. The six training terms are defined in Section 3.2.')
  change(doc,'Figure 3.','Figure 3. Same-cohort ablation across seeds. A: equal-weight mean within-section Moran’s I across 35 sections. B: section-level Kendall τ-b. Bars and error bars show the mean and sample SD across optimization seeds; dots show each seed. Counts indicate seeds with nominal one-sided fixed-score permutation p<0.05 (B=10,000); complete values and exact enumeration are in Supplementary Table S1. These are alignment diagnostics, not null-refitted tests of generalization. Legacy globally centered statistics appear in Supplementary Table S3.')
  change(doc,'Figure 5.','Figure 5. External evaluation. A: section-level Moran’s I for model scores, untrained KEGG panels, and expression-matched random panels; dots represent sections. The dashed zero line approximates the within-section coordinate-permutation null. B: GeoMx core scores; the six-patient paired p=0.03125 is primary, while displayed cores share patients. C: descriptive mean core scores by Silva invasion pattern; these are not independent patient-level estimates. D: held-out disease percentiles relative to 31 training normals. Supervised and target-free protocols differ in both ordinal and proxy targets; pathway regularization remains active. SCC 1, ADC and SCC 2 denote CA_HPV_1, CA_HPV_2 and cervical_cancer_1, respectively.')
 elif name=='Manuscript_JBI_中文.docx':
  change(doc,'图 1','图 1 宫颈组织学背景与模型架构。A：原始H&E图像来自正常外宫颈（P020；GSE217237）、HSIL（HSIL_HPV；GSE208654）、浸润性鳞癌（cervical_cancer_1；10x Genomics）及浸润性腺癌（CA_HPV_2；GSE208654）。下排为框选区域；仅裁剪及缩放显示，未调整颜色，物理放大倍数不同。H&E提供组织学背景，不构成模型评分的空间验证。B：表达与冻结UNI形态特征进入跨模态模型，输出评分、代理及重建；六项训练损失见3.2节。')
  change(doc,'图 3','图 3 同队列多种子消融。A：35张切片内Moran’s I的等权均值。B：切片级Kendall τ-b。柱与误差线为优化种子的均值及样本标准差，圆点为单个种子；上方计数为单侧固定评分置换p<0.05的种子数（B=10,000）。完整数值及穷举见补充表S1。该检验诊断评分对齐，不是置换重训的泛化检验；旧全局中心化指标见补充表S3。')
  change(doc,'图 5','图 5 外部评估。A：头颈部各切片模型、KEGG面板及表达匹配随机面板的Moran’s I，圆点为切片；零线近似切片内坐标置换零分布。B：GeoMx组织芯评分，主要推断来自六对患者（p=0.03125），芯共享患者。C：Silva浸润模式的描述性芯均值。D：留出疾病相对31张训练正常切片的百分位；两方案同时改变有序和代理目标，均保留通路项。SCC 1、ADC、SCC 2依次为CA_HPV_1、CA_HPV_2、cervical_cancer_1。')
 for p in doc.paragraphs:
  if name=='Manuscript_JBI.docx':
   for prefix,number in [('Mean per-spot risk rises',2),('Removing ordinal supervision reduced',3),('High-risk spots are enriched',4),('GeoMx: group separation',5)]:
    if p.text.startswith(prefix) and f'(Figure {number})' not in p.text:settext(p,p.text.rstrip('.')+(f' (Figure\u00a0{number}; Table\u00a02).' if number==3 else f' (Figure\u00a0{number}).'))
  if p.style.name=='Caption':
   for r in p.runs:r.font.color.rgb=RGBColor(0,0,0);r.font.size=Pt(10);r.italic=False;r.bold=False
  elif p.style.name=='Normal' and len(p.text)>45 and len(p.runs)==1 and (p.runs[0].font.size is None or p.runs[0].font.size.pt<=11):
   p.runs[0].bold=False
 if name=='Manuscript_JBI_中文.docx':
  rep(doc,'标志物有效性','标记集一致性');rep(doc,'标志物有效性在独立宫颈scRNA-seq','标记集一致性在独立宫颈scRNA-seq');rep(doc,'中确认：谱系内平均相关','中评估：标记集内平均基因相关');rep(doc,'跨谱系','标记集间')
  rep(doc,'谱系内评分相关','标记集内基因相关');rep(doc,'谱系间','标记集间');rep(doc,'标记有效性','标记集一致性')
 for t in doc.tables:
  if len(t.rows)<=9:
   for i,row in enumerate(t.rows):
    for c in row.cells:
     for p in c.paragraphs:p.paragraph_format.keep_with_next=i<len(t.rows)-1
 doc.save(P/name)
 print('Saved',name)
def run_source(filename,changes=None,main=True):
 src=(ROOT/'scripts'/filename).read_text(encoding='utf-8-sig')
 for a,b in (changes or []):
  if a not in src:raise ValueError(('source pattern missing',filename,a[:70]))
  src=src.replace(a,b)
 ns={'__file__':str(ROOT/'scripts'/filename),'__name__':'__main__' if main else filename[:-3]}
 exec(compile(src,str(ROOT/'scripts'/filename),'exec'),ns)
 return ns
def figures():
 injection="""_new = json.load(open('results/jbi_submission_audit/numbers.json', encoding='utf-8'))
_audit = json.load(open('results/jbi_submission_audit/rank_perm_recomputed.json', encoding='utf-8'))
abl_summary = _new['ablation']
for _v in abl:
    for _r in abl[_v]:
        _r['moran'] = next(x['moran_within_mean'] for x in _audit['variants'][_v] if x['seed']==_r['seed'])
NUM['ablation'] = abl_summary"""
 # Regenerate data plots directly; the shared historical figure directory is untouched.
 run_source('gen_figures_v10.py',[
 ('"07_CSBJ_v10"','"08_JBI"'),
 ("NUM['ablation'] = abl_summary",injection),
 ("with open('results/paper_v10_numbers.json', 'w'","with open('papers/08_JBI/_qa_20260924/figure_build_numbers.json', 'w'"),
 ('ax.set_ylabel("Moran\'s I (spatial coherence)")','ax.set_ylabel("Mean within-section Moran\'s I")'),
 ('ax.set_title("A. Spatial coherence"','ax.set_title("A. Within-section spatial autocorrelation"'),
 ('p = s.get(\'tau_perm_p_max\', s[\'tau_p_max\'])',"p = s['tau_perm_p_max']"),
 ('lab = "n.s.\\n(p=%.3f)" % p','lab = "p=%.3f-%.3f" % (s["tau_perm_p_min"],p)'),
 ('B. Cross-platform rank discrimination; sign inverted vs training','B. GeoMx core scores (exploratory; cores share patients)'),
 ('tumor {tum.mean():.2f} < benign {ben.mean():.2f}, MWU p={gx[\'core_level\'][\'benign_vs_tumor_p\']:.1e}','patient-paired p=0.03125; core MWU p={gx[\'core_level\'][\'benign_vs_tumor_p\']:.1e}'),
 ('C. Not a risk grade (core level)','C. Descriptive pattern/risk groups (cores share patients)'),
 ("label='Fully unsup. (no labels)'","label='Target-free (pathway term retained)'"),
 ("label='Semi-sup. (3 diseases labeled)'","label='Ordinal + proxy targets active'"),
 ('plt.savefig(os.path.join(OUT, "Fig3_ablation_multiseed.png")); plt.close()','plt.savefig(os.path.join(OUT, "Fig3_ablation_multiseed.png")); plt.savefig(os.path.join(OUT, "Fig3_ablation_multiseed.pdf")); plt.close()'),
 ('plt.savefig(os.path.join(OUT, "Fig4_biology_correlation.png")); plt.close()','plt.savefig(os.path.join(OUT, "Fig4_biology_correlation.png")); plt.savefig(os.path.join(OUT, "Fig4_biology_correlation.pdf")); plt.close()'),
 ('plt.savefig(os.path.join(OUT, "Fig2_risk_violin.png")); plt.close()','plt.savefig(os.path.join(OUT, "Fig2_risk_violin.png")); plt.savefig(os.path.join(OUT, "Fig2_risk_violin.pdf")); plt.close()'),
 ('plt.savefig(os.path.join(OUT, "Fig5_external_transfer.png")); plt.close()','plt.savefig(os.path.join(OUT, "Fig5_external_transfer.png")); plt.savefig(os.path.join(OUT, "Fig5_external_transfer.pdf")); plt.close()'),
 ])
 run_source('gen_fig1_framework.py',[
 ('"07_CSBJ_v10"','"08_JBI"'),('"per-spot linear"','"128 - 64 - 1 MLP"'),
 ('"HPV-pathway head"','"Proxy probe"'),('"alignment to the", "108 KEGG-HPV genes"','"binary target from", "detected panel genes"'),
 ('["linear", r"','["128 - 256 - 4,991", r"'),
 ('YSMALL - BH / 2 - 3.6','YSMALL - BH / 2 - 1.6'),
 ('"var", "hpv", "sup"','"var", "hpv", "sup"'),
 ('on the Visium hexagonal lattice','in the batch; var = score variance floor; sup = score + probe'),
 ])
 run_source('gen_fig1_histology_v10.py',[('"07_CSBJ_v10"','"08_JBI"')])
 # Draw a concise graphical abstract using plotting primitives; no generated tissue data.
 import matplotlib.pyplot as plt
 from matplotlib.patches import FancyBboxPatch,FancyArrowPatch
 fig,ax=plt.subplots(figsize=(13.28,5.31),dpi=300);ax.set(xlim=(0,1),ylim=(0,1));ax.axis('off')
 ax.text(.5,.955,'Auditing label dependence in small-sample spatial transcriptomics',ha='center',fontsize=17,weight='bold',color='#203864')
 ax.text(.5,.90,'A pathway-anchored benchmark in cervical lesions',ha='center',fontsize=12,color='#555555')
 blocks=[(.015,'1  SMALL-SAMPLE PROBLEM','#eaf0fa'),(.35,'2  CONTROLLED AUDIT','#f3f0f8'),(.685,'3  BOUNDED CONCLUSIONS','#eef4ef')]
 for x,title,fc in blocks:
  ax.add_patch(FancyBboxPatch((x,.14),.30,.69,boxstyle='round,pad=0.01',fc=fc,ec='#203864',lw=1.5))
  ax.text(x+.15,.78,title,ha='center',fontsize=12,weight='bold',color='#203864')
 texts=[
 "35 cervical sections\n31 normal; 4 diseased\n46,820 correlated spots\n\nIn-sample disease AUC = 1.00\n\nMany spots do not add\nindependent disease samples.",
 "Seven variants; multiple seeds\n\nOrdinal targets off\n        ↓\nProxy targets off\n        ↓\nPathway regularizer off\n\nSection-level tests and hold-outs",
 "Disease-ordering Kendall τ\n0.46 → 0.25 → 0.21 → −0.11\n\nWithin-section structure persists\nNo established morphology gain\nGeoMx score direction reverses\n\nFive-item MRP-ST proposal\nNo established clinical utility"]
 for (x,_,_),txt in zip(blocks,texts):ax.text(x+.018,.705,txt,va='top',fontsize=11.5,linespacing=1.45,color='#222222')
 for x in [.325,.66]:ax.add_patch(FancyArrowPatch((x,.5),(x+.019,.5),arrowstyle='-|>',mutation_scale=18,color='#203864',lw=2))
 ax.text(.5,.065,'Target dependence is not proof of memorization; spatial coherence is not clinical validation.',ha='center',fontsize=11.5,style='italic',color='#555555')
 fig.subplots_adjust(left=.01,right=.99,bottom=.01,top=.99)
 for ext in ['png','pdf','tiff']:
  kw={'pil_kwargs':{'compression':'tiff_lzw'}} if ext=='tiff' else {}
  fig.savefig(P/f'GraphicalAbstract_JBI.{ext}',dpi=300,**kw)
 plt.close(fig)
def supplementary(g):
 doc,C,h1,h2,body,ctr,table,fig,emit=g.make_manuscript()
 emit('Supplementary Information',16,True);emit(TITLE,12,True)
 emit('Mingwei Bai, Zhengmin Gu, Ye Zhang, Lang Guo',11)
 body('This supplement uses the same result files and definitions as the revised manuscript. Numerical source data and file hashes accompany the submission.')
 h1('Supplementary Table S1. Section-level disease ordering')
 body('Current saved NPZ scores are aggregated to one mean per section. Monte Carlo tests are one-sided for positive τ (B=10,000, seed 20260924); exact p enumerates all 209,440 arrangements. Scores remain fixed. These are alignment diagnostics, not label-refitted model tests. Seeds are optimization repeats, not independent biological samples.')
 rows=[]
 for v,rr in D['variants'].items():
  for x in rr:rows.append([v,x['seed'],f"{x['tau']:.4f}",f"{x['p_perm_B10000']:.5f}",f"{x['p_exact']:.6f}"])
 table(['Variant','Seed','Kendall τ-b','Monte Carlo p','Exact p'],rows)
 body('FULL seeds are 0, 3, and 4; other three-seed arms use 0, 1, and 2. NO-PROG uses 0–3. Some older permutation summaries used different scores and were not reused. Asymptotic SciPy p-values are two-sided whereas the permutation test is one-sided; their numerical differences are not evidence that one is necessarily anti-conservative.')
 h1('Supplementary Table S2. MRP-ST reporting proposal')
 table(['Item','Recommended reporting'],[
 ['R1: target removal','Report all active losses and nested ordinal/proxy/pathway ablations. Include several seeds and disclose unmatched seeds or stopping rules. Do not interpret dependence alone as memorization.'],
 ['R2: permutation','Use independent samples rather than spots. Distinguish fixed-score alignment diagnostics from label permutation with complete refitting. Report direction, B, seed, and p-value calculation.'],
 ['R3: random genes','Prefer size- and expression-matched controls. Identify trained-model controls separately from untrained panel means and disclose mismatches.'],
 ['R4: hold-outs','Fit preprocessing on training samples only. Report whether normal controls are also held out and whether folds/donor sections share information.'],
 ['R5: evidence','Separate spatial structure, ordering, discrimination, calibration, and outcomes. State uncertainty and clinical limitations.']])
 body('The suggested trigger of ≤10 diseased samples or ≤5 in a class is unvalidated and does not define sufficient sample size. Supported means evidence for the stated estimand/design, suggestive means exploratory or unstable evidence, and unsupported means not established rather than proven absent. In this study, spatial autocorrelation is documented; target dependence and external GeoMx separation are suggestive; morphology benefit and clinical utility are not established. External section-level panel comparisons do not demonstrate donor-level replication. Random-panel retraining was size-matched only, a departure from the preferred design.')
 h1('Supplementary Table S3. Sensitivity to spatial estimand')
 table(['Variant','Within-section I','Legacy pooled I','Seeds'],[
 [v,f"{a['moran_mean']:.3f} ± {a['moran_sd']:.3f}",f"{a['moran_global_mean']:.3f} ± {a['moran_global_sd']:.3f}",','.join(a['seeds'])] for v,a in A.items()])
 body('Within-section I centers the score within each section and averages the 35 section statistics with equal weight. The SD shown is across seed-specific section means, not across independent patients. Legacy pooled I centers over all spots; even if neighbor edges stay within sections, section-mean shifts enter its numerator and denominator. It cannot be treated as a pure within-section measure. FULL minus NO-IMAGE is 0.009 for the primary estimand versus 0.072 for the legacy estimand. No inferential morphology-benefit claim is made.')
 h1('Supplementary Note S1. Frozen UNI feature autocorrelation')
 s=json.loads((R/'uni_embed_moran_s1.json').read_text(encoding='utf-8'))
 body('The saved feature analysis uses local principal components after section-specific z-standardization. Across 35 sections, pooled within-section PC1 Moran’s I is 0.618; the median section PC1 statistic is 0.579 (IQR 0.491–0.726; range 0.246–0.948). The within-section permutation null is centered near −0.0007. This establishes spatial structure in the frozen features, not clinical information or incremental prediction. Unique patch counts check gross extraction failure but do not independently establish registration accuracy.')
 h1('Supplementary Note S2. GeoMx hierarchy and descriptive stratification')
 body('The final expression-only score file was joined to GEO SOFT labels by ROI identifier and checked against the supplied labeled CSV. It covers 186 ROIs, 93 cores and seven patients, with six paired tumor–benign patients. Tumor scores are lower in all six pairs (two-sided paired Wilcoxon p=0.03125). The score is uncalibrated; direction reversal is not a formal probability-calibration result.')
 body('Silva core means are A −14.25, B −12.07 and C −12.73. In epithelial ROIs they are −16.18, −12.52 and −12.13; stromal values are −12.32, −11.62 and −13.32. The A-low epithelial pattern follows the original training score orientation for better-prognosis tissue, not the externally reversed tumor–benign orientation. This inconsistency further limits a monotonic clinical-risk interpretation. High/Low core means are −12.55/−13.68 (p=0.100). Pattern and compartment tests reuse patient observations and lack comprehensive multiplicity adjustment; all are descriptive. Balanced epithelial/stromal ROI counts do not exclude compositional confounding.')
 h1('Supplementary Note S3. External head-and-neck controls')
 body('The model score has mean section-specific Moran’s I 0.76/0.67 in GSE268014/GSE281978, versus near-zero within-section coordinate-permutation nulls (B=1,000, p≤0.001). Plain KEGG panels give 0.71/0.62 and expression-decile-matched random panels approximately 0.58/0.48. KEGG exceeds random means in 8/8 and 6/6 sections (nominal two-sided sign p=0.008/0.031). These section-based tests assume section independence and do not account for donor pairing; therefore donor-level replication is not claimed.')
 table(['Contrast','p','Interpretation'],[
 ['Malignant versus non-malignant','0.195','Not established at section level'],
 ['Tumor front versus core','0.055','Not significant; direction mostly reversed'],
 ['HPV status: GSE268014 / GSE281978','0.200 / 0.133','Small groups; no established discrimination'],
 ['Primary versus lymph-node metastasis','1.000','No detected difference; not equivalence']])
 h1('Supplementary Table S4. Position relative to existing methods')
 table(['Method family','Primary purpose','Role in this study'],[
 ['HE2RNA / ST-Net','Image-to-expression prediction','Context; not a direct benchmark of disease risk'],
 ['OmiCLIP / SPADE / SEAL / MINT / ST-Align / STPath','Representation or expression prediction','Motivate evaluation across modalities'],
 ['scGPT / Novae / CellCharter','Representation and spatial organization','Context for separating structure from endpoints'],
 ['MRP-ST proposal','Audit target dependence and evaluation units','One worked architecture/disease example; independent validation needed']])
 h1('Supplementary Table S5. Cervical sections')
 raw=json.loads((R/'results_FULL_seed0.json').read_text(encoding='utf-8'))
 counts=raw['per_sample_spots']
 table(['Section ID','Severity','Spots'],[
 [Path(k).name, 'HSIL' if 'HSIL' in k else 'Cancer' if ('CA_HPV' in k or 'cervical_cancer_1' in k) else 'Normal',str(v)] for k,v in counts.items()])
 body('Counts sum to 46,820 spots across 35 sections. Section names are the analyzed IDs, not an assumed contiguous P001–P030 series.')
 doc.save(P/'Supplementary_Information_JBI.docx')
def fixed_margins(cell,top=80,start=100,bottom=80,end=100):
 pr=cell._tc.get_or_add_tcPr();mar=pr.find(qn('w:tcMar'))
 if mar is None:mar=OxmlElement('w:tcMar');pr.append(mar)
 for side,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
  x=OxmlElement('w:'+side);x.set(qn('w:w'),str(val));x.set(qn('w:type'),'dxa');mar.append(x)
def main():
 figures()
 run_source('jbi_final_figures.py')
 src=(ROOT/'scripts/gen_jbi.py').read_text(encoding='utf-8')
 src=src.replace('"paper_v10_numbers.json"','"jbi_submission_audit/numbers.json"')
 g=types.ModuleType('jbi_base');g.__file__=str(ROOT/'scripts/gen_jbi.py')
 exec(compile(src,g.__file__,'exec'),g.__dict__)
 g._set_cell_margins=fixed_margins
 g.save=save_final
 g.build_manuscript();g.build_manuscript_cn();g.build_highlights();g.build_cover_letter()
 supplementary(g)
if __name__=='__main__':main()
