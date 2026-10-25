from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
import shutil,zipfile,json,hashlib
P=Path('papers/08_JBI');S=P/'Submission_JBI_20260925';U=S/'Upload'
source=Document(P/'Manuscript_JBI.docx');title=source.paragraphs[0].text
coi=next(p.text for p in source.paragraphs if p.text.startswith('Declaration of competing interest.'))
def blank(heading):
 d=Document(P/'Manuscript_JBI.docx')
 for e in list(d._element.body):
  if not e.tag.endswith('}sectPr'):d._element.body.remove(e)
 p=d.add_paragraph(heading);p.paragraph_format.space_after=Pt(14)
 for r in p.runs:r.bold=True;r.font.size=Pt(14)
 p=d.add_paragraph(title);p.paragraph_format.space_after=Pt(12)
 for r in p.runs:r.bold=True;r.font.size=Pt(11)
 d.add_paragraph('Mingwei Bai, Zhengmin Gu, Ye Zhang, Lang Guo')
 return d
d=blank('Statement of Significance');t=deepcopy(source.tables[0]._element);d._element.body.insert(len(d._element.body)-1,t)
d.save(P/'Statement_of_Significance_JBI.docx')
d=blank('Declaration of competing interests');d.add_paragraph(coi.split('. ',1)[1]);d.save(P/'Declaration_of_Competing_Interests_JBI.docx')
for n in ['Statement_of_Significance_JBI.docx','Declaration_of_Competing_Interests_JBI.docx']:shutil.copy2(P/n,U/n)
for p in [S/'投稿文件与核对说明.txt',P/'投稿文件与核对说明_20260925.txt']:
 s=p.read_text(encoding='utf-8-sig');s+='\n投稿系统截图所需的独立文件已补齐：\nStatement of Significance → Statement_of_Significance_JBI.docx（由正文意义声明提取）。\nGraphical abstract → GraphicalAbstract_JBI.tiff。\nManuscript → Manuscript_JBI.docx。\nDeclaration of competing interests → Declaration_of_Competing_Interests_JBI.docx（沿用正文无利益冲突声明；需作者确认真实性）。\n';p.write_text(s,encoding='utf-8-sig')
checks=[{'file':str(f.relative_to(S)).replace('\\','/'),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()} for f in sorted(S.rglob('*')) if f.is_file() and f.name!='SHA256_manifest.json']
(S/'SHA256_manifest.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
with zipfile.ZipFile(P/'JBI_Submission_20260925.zip','w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for f in sorted(S.rglob('*')):
  if f.is_file():z.write(f,Path(S.name)/f.relative_to(S))
print('Added two standalone declarations; updated submission ZIP.')
print(coi)
