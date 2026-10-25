#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""_qa_v15.py — v15 稿件 QA：禁字符、表数、字数、v15 新增内容。"""
import re
from pathlib import Path

from docx import Document

D = Path(__file__).resolve().parent.parent / "papers" / "20_MRP-ST-Repair_Patterns"
d = Document(D / "Manuscript_Patterns.docx")
ps = [p.text for p in d.paragraphs]
full = "\n".join(ps)
for t in d.tables:
    for r in t.rows:
        full += "\n" + " | ".join(c.text for c in r.cells)

WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9'+/-]*")
print("em-dash", full.count("\u2014"), "| literal **", full.count("**"),
      "| main tables", len(d.tables))
print("words(all):", len(WORD.findall(full)))

ii = [i for i, x in enumerate(ps) if x.strip().upper() == "INTRODUCTION"][0]
jj = [i for i, x in enumerate(ps)
      if x.strip().upper().startswith("EXPERIMENTAL PROCEDURES")][0]
body = "\n".join(ps[ii:jj])
print("main text words (Introduction -> Experimental Procedures):",
      len(WORD.findall(body)))

checks = [
    # v14 内容仍在
    "Control: the channel is opened by coupling",
    "Table S17", "Table S16", "44. [Dataset] Axelrod", "GSE316098",
    # v15 新增：摘要层三级证据
    "Validated in one domain; two further cohorts test predictions, not utility",
    "Validation is end-to-end in spatial transcriptomics alone",
    "no accepted repair improved a downstream endpoint, which the framework predicts in advance",
    "Not tested", "structural mappings only, for which we claim no transfer",
    # v15 新增：算子非贡献
    "We therefore make no claim that the adversarial operator used here outperforms",
    "The deliverable of this paper is the auditing protocol, not the repair operator",
    "The instrument does not depend on the operator",
    # v15 新增：partial removal
    "Passing the gate means the measured group dependence has been brought below a calibrated threshold",
    "Do not read acceptance as erasure",
    # v15 新增：Prop1 退化
    "Degenerate regime, to be read before any accept verdict is used",
    "Check that the baseline is above the null before reading an accept verdict",
    # v15 新增：group-count placement
    "That interval is a placement for this configuration and not a transferable cut-off",
    "placements for this configuration, not transferable cut-offs",
    # v15 新增：线性探针
    "LDI measures the group identity that a",
    "nonlinear probe can decode",
    # v15 新增：统计
    "The cohort sizes are fixed by what is publicly available",
]
for c in checks:
    n = full.count(c)
    print(f"  {'OK ' if n >= 1 else 'MISS'} {n}  {c[:70]}")

# 禁止的过度宣传措辞
for bad in ["validated across multiple domains", "validated in multiple domains",
            "improves downstream performance", "improves prediction performance",
            "our repair improves", "fully removed", "group identity is removed"]:
    n = full.lower().count(bad.lower())
    print(f"  {'OK ' if n == 0 else 'WARN'} {n}  banned: {bad}")

# 参考文献条数只数 REFERENCES 段，否则把协议里的编号步骤也算进去
ri = [i for i, x in enumerate(ps) if x.strip().upper().startswith("REFERENCES")][0]
print("ref count:", len([p for p in ps[ri:] if re.match(r"^\d+\. ", p.strip())]))
print("SI table objects:", len(Document(D / "Supplemental_Information_Patterns.docx").tables))
for p in Document(D / "TitlePage_Patterns.docx").paragraphs:
    if "Manuscript statistics" in p.text or "approximately" in p.text:
        print("  TP:", p.text[:200])
