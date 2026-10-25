"""Patch the CURRENT v28 rather than rebuilding from the earlier v27.

Preserves tables, figures and layout. All original artifacts are backed up once.
"""
from pathlib import Path
import re,copy,json,sys,shutil
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
sys.stdout.reconfigure(encoding='utf-8')
P=Path('papers/20_MRP-ST-Repair_Patterns');Q=Path('.qa/patterns_v29')
for n in ['Manuscript_Patterns.docx','Supplemental_Information_Patterns.docx','CoverLetter_Patterns.docx','TitlePage_Patterns.docx','CoverLetter_Patterns.md','TitlePage_Patterns.md','SUBMISSION_READINESS.md']:
 dst=P/(Path(n).stem+'_v29pre_20261025'+Path(n).suffix)
 if not dst.exists():shutil.copy2(P/n,dst)
t=(P/'MRP-ST-Repair_manuscript_Patterns_v28.md').read_text(encoding='utf-8')
original=t
changes=[]
def replace_paragraph(prefix,new):
 global t
 pat=re.compile(r'(?m)^'+re.escape(prefix)+r'[^\n]*(?:\n(?!\n)[^\n]*)*')
 hits=list(pat.finditer(t))
 if len(hits)!=1:raise ValueError((prefix,len(hits)))
 old=hits[0].group();t=t[:hits[0].start()]+new+t[hits[0].end():]
 changes.append((old,new))
repo='https://github.com/WorldBMW/pathway-anchored-risk-benchmark';doi='https://doi.org/10.5281/zenodo.23640979'
release=repo+'/releases/tag/v1.0.0'
ref='44. [Software and manuscript snapshot] WorldBMW. WorldBMW/pathway-anchored-risk-benchmark: v1.0 - MRP-ST Audit Evidence Package (v1.0.0). Zenodo (2026). '+doi+'.'
availability='A manuscript snapshot and an associated code snapshot are publicly archived at '+doi+'[44], with the linked repository at '+repo+' (release v1.0.0; commit 73ad0f1710953651c43055afc693f7a18ce93c6f). This deposit is a historical snapshot: its code archive does not contain the cross-domain audit, v28 reanalysis, or main-figure generation scripts used for the present revision. It must not be treated as the complete reproducibility archive for the current results. The additional analysis code, result tables and file-level checksum manifest are provided in the accompanying revision package; public deposition of that package is required for a complete version-matched archive.'
replace_paragraph('**Code.**','**Code.** '+availability)
replace_paragraph('The public source accessions used for the transcriptomic analyses','The public source accessions used for the transcriptomic analyses are GSE208654, GSE217237, GSE268014, GSE278936 and GSE316098, together with the stated 10x Genomics example. HAR and Gas Sensor Array Drift are obtained from the UCI repository. '+availability+' Source licences and participant-data restrictions remain those of the original deposits; no new participant data were collected.')
replace_paragraph('**Group-preserving null, recomputed.**','**Within-group label permutation sensitivity.** We compared unrestricted label permutation with labels shuffled independently within each recorded group (B = 2,000 probe refits). The restricted scheme preserves the observed group-by-label counts and asks about association beyond that structure, conditional on within-group exchangeability; it is not permutation of whole independent groups. Baseline permutation scales are approximately 0.058 versus 0.058 in HAR, 0.060 versus 0.046 in gas drift, 0.081 versus 0.050 for GeoMx risk, 0.077 versus 0.048 for Silva and 0.083 versus 0.087 for cell type. In the 30 reanalysis cells of these multi-unit cohorts, the archived ratio (τ_LOO − 0.001)/σ̂ remains above 1.645 using either scale. This ratio is not centered on the restricted null mean, which was not stored, so its numerical increase is not evidence of greater statistical significance. Empirical permutation p values, where exchangeability is defensible, address the specified conditional null; the normal-reference gate is descriptive. Where labels are constant within group, within-group shuffling is the identity, σ̂ = 0 and this ratio is undefined. This limits that permutation scheme, not the possibility of group-level inference under a different scientifically justified null.')
replace_paragraph('**Identity decoding, separately from label utility.**','**Identity structure, separately from label utility.** The reanalysis reports nearest-neighbour same-group hit rate divided by the archived empirical prior Σ(n_g/N)² and distance-based AUROC for same-group versus different-group pairs. These are geometry summaries, not open-set group classifiers or privacy guarantees. The empirical prior includes self-pairs, whereas the nearest-neighbour search excludes self; a random-neighbour reference would instead be Σn_g(n_g − 1)/[N(N − 1)]. We retain the archived ratio for traceability and do not use it as a calibrated significance test. Pairwise distances also share observations and do not provide independent pair replicates. From λ = 0 to 2, the archived nearest-neighbour ratio changes from 12.7–13.8 to 4.8–6.3 in HAR, 5.14–5.15 to 4.70–4.84 in gas drift and 5.65–5.71 to 5.43–5.56 in GeoMx risk. These values differ from the logistic group-classifier lift used in Figure 5 and earlier tables. Both measurements show retained group structure under their own definitions; neither certifies identity erasure.')
replace_paragraph('**Nominal labels and coding sensitivity.**','**Nominal labels and coding sensitivity.** One-versus-rest ridge probes under recorded group holdout give macro AUROC 0.96–0.97 in HAR and 0.88–0.94 in gas drift, with balanced accuracy 0.80–0.86 and 0.67–0.84. These are conditional representation audits. Twenty random recodings give τ_LOO SD 0.073–0.112 and an observed range up to 0.49, while the nominal macro AUROC is invariant to relabelling. Thus nominal scoring remedies the coding dependence of the utility readout, but does not validate the earlier τ-based cross-domain ordering or feasibility interpretation. At λ = 2, HAR balanced accuracy decreases in all three seeds; gas-drift macro AUROC and balanced accuracy both decrease in all three seeds. No consistent utility improvement is demonstrated.')
replace_paragraph('**Outer-loop refitting.**','**Outer-loop encoder refitting.** The archived HAR comparison retrains the encoder and probe over 30 subject holdouts at one seed and two doses. However, clipping percentiles and global standardization were fitted before the holdouts, so the analysis is not fully nested preprocessing. It reports the mean of per-subject Kendall correlations (0.802 versus 0.806 at λ = 0; 0.776 versus 0.804 at λ = 2), whereas the main τ_LGO pools out-of-fold predictions across subjects. These estimands are distinct. Comparing the difference with between-fold SD is not an equivalence test; folds share training data. These exploratory results therefore do not establish absence of optimism or validate the complete pipeline. A corrected train-only preprocessing runner is supplied separately, with pooled and mean-fold statistics reported distinctly.')
replace_paragraph('**Donor and patient resolution.**','**Grouping assumptions and patient sensitivity.** The cervical runner merges CA_HPV_1 and CA_HPV_2 through a sample-name regular expression, giving 35 sections over 34 assumed groups. A verified source patient-identifier map is not supplied by that operation. We therefore report the existing counts under this grouping assumption rather than claiming a reconstructed patient map. Under it, the observed τ_LOO − τ_LGO ranges from −0.007 to 0.073 across nine runs. The GeoMx analysis enumerates 21 possible merges of two of seven case identifiers, reflecting the source count of six patients. It applies probe holdout to the preprocessed expression matrix, not the repaired learned embeddings: the τ_LGO ranges 0.519–0.614 for risk, 0.530–0.625 for Silva and 0.706–0.708 for cell type are raw-feature grouping sensitivity only. This neither identifies the actual shared patient nor establishes invariance of repair conclusions to that pairing.')
replace_paragraph('The representation is generally learned before the probe holdouts, so most estimates','The representation is generally learned before probe holdout, so most estimates are conditional audits rather than generalization estimates of a full training pipeline. The archived outer-refit exercise still fits preprocessing across all subjects and compares a mean-fold statistic rather than pooled τ_LGO. It provides neither a fully nested estimate nor an equivalence test. A corrected train-only preprocessing runner is provided, but the transcriptomic experiments remain conditional; independent external evaluation and correctly grouped outer fitting are required for stronger deployment claims.')
replace_paragraph('The remediation runs above replace the corresponding caveats','The additional analyses narrow specific uncertainties but do not replace the remaining methodological limitations. Within-group label shuffling does not test the same null as unrestricted or whole-group permutation, and it requires its own exchangeability assumptions. The stored scale ratios are not centered on a restricted-null mean, while the GeoMx scheme does not preserve paired compartments within ROIs. Degeneracy under within-group-constant labels does not preclude group-level inference with a different null. The nearest-neighbour and pair-distance summaries are descriptive, and high-stakes uses need stronger decoding probes. Nominal metrics address label coding for utility, but the ordinal results elsewhere remain encoding-dependent.')
replace_paragraph('The group identifiers still bound the conclusions in one specific way.','The group identifiers still bound the conclusions. The cervical 34-group map is generated by name-based merging rather than verified patient metadata. The resolution ladder instead uses 35 section groups and cannot be pooled with it as donor-disjoint evidence. GeoMx has seven specimens from six patients, but its actual patient-to-case map is unresolved here. The 21-pair enumeration probes raw-feature holdout sensitivity only, while the learned-embedding analyses remain grouped by specimen. Neither that enumeration nor within-case label shuffling establishes patient-disjoint inference or preserves every ROI-level dependency.')
replace_paragraph('**Group-preserving null and the remediated statistics.**','**Restricted permutation and additional audit statistics.** The v28 restricted null shuffles labels independently within each recorded group, retaining group-by-label counts. It is a conditional within-group randomization scheme, not whole-group permutation. B = 2,000 and seed 20261010 were used for probe refitting and empirical tail counts. The stored LDI_group is (τ_LOO − 0.001)/σ̂_group; it is not centered using a restricted-null mean and is not interpreted as a Gaussian Z statistic. σ̂_group = 0 is reported when the scheme has no label variation; no ratio is defined then. For group-constant labels, a scientifically justified permutation of group-level label assignments may be appropriate, subject to independent-group exchangeability; its validity must be assessed rather than inferred from the degenerate within-group scheme. The archived nearest-neighbour ratio uses Σ(n_g/N)², while the random-neighbour prior with self excluded is Σn_g(n_g − 1)/[N(N − 1)]. One-versus-rest macro AUROC and balanced accuracy use recorded-group holdout. The patient-pair sensitivity enumerates merges on preprocessed raw GeoMx features, not repaired embeddings. The original outer-refit script fits preprocessing globally and averages foldwise τ; the corrected runner fits it on outer training data and reports both pooled and mean-fold τ.')
replace_paragraph('**Symbols and interpretation.**','**Symbols and interpretation.** τ_in is in-sample rank association; τ_LOO holds out one analysis unit, whereas τ_LGO or τ_LDO holds out its recorded group. LDI = (τ_LOO − 0.001)/σ̂ uses the archived unrestricted unit-permutation scale and does not certify identity removal. R = τ_LOO/τ_in is unit-holdout retention and Δ = τ_LOO − τ_LGO is the paired holdout difference. σ̂_group is a standard deviation of the within-group permutation distribution, not itself an index; the corresponding ratio is undefined at zero variance and is not centered on that restricted null. Figure 5 and earlier tables use logistic group-classifier accuracy divided by 1/G. The additional v28 analysis uses a separate nearest-neighbour same-group ratio and pair-distance AUROC; those values must not replace the earlier classifier values in their legends. Training seeds quantify optimization spread on fixed cohorts, not independent population replication. Mean-fold and pooled out-of-fold τ are distinct statistics.')
replace_paragraph('**Table 7. Archived unit-level null','**Table 7. Unrestricted and within-group label-permutation scales in the reanalysis subset.** Three seeds and two doses per multi-unit task were examined, together with nine cervical archived cells; the complete original dose grids were not rerun under the restricted scheme. κ is measured on the sampled audit units rather than the full cohort. Baseline σ̂ values are averaged across three seeds; ratio ranges span seeds and λ ∈ {0, 2}. LDI_group retains the archived floor 0.001 rather than a restricted-null mean, so it is a descriptive scaled association, not a null-centered Z score. “Gate unchanged” refers only to this normal-reference dependence ratio for cells with nondegenerate distributions, not to a valid group-level test or all acceptance axes. Cervical within-group shuffling has zero variance and the ratio is unavailable. Standard deviations remain defined at zero; only division by zero is undefined.')
# Change table headers consistently without rewriting numerical records.
t=t.replace('σ̂ group-preserving','σ̂ within-group').replace('LDI group-preserving','Scaled association (within-group SD)').replace('Gate decision changed','Reference crossing changed')
t=t.replace('Identity leakage in group-structured data, and why post-hoc repair cannot fix it','Projection outcomes in the tested representations')
t=t.replace('**Figure 1. Identity leakage in group-structured data, and why repairing a trained model afterwards\ncannot work.** (A) The auditable loop: detect leakage with a calibrated statistic under leave-one-group-out;','**Figure 1. Auditing the tested projection intervention.** (A) The workflow schematic; its association index is computed under unit-level probe holdout, with group holdout reported separately;')
# Citation with real deposited metadata; do not invent a five-author Zenodo creator list.
if '\n44. ' not in t:t=t.rstrip()+'\n'+ref+'\n'
for line in list(t.splitlines()):
 if line.startswith('| Software |') and 'github.com/WorldBMW' in line:
  cells=[c.strip() for c in line.strip('|').split('|')];cells[-1]=doi+' [44]; historical snapshot at '+repo+'; current-revision audit scripts supplied separately';t=t.replace(line,'| '+' | '.join(cells)+' |')
# Correct inaccurate AI-use declaration for code-assisted scientific work.
ai='During preparation, AI-assisted tools were used for language editing and code review. OpenAI Codex also assisted in drafting and checking analysis scripts. Numerical results were obtained by executing the documented scripts; they were not invented as narrative values. Script versions, inputs and outputs are provided for independent verification. The authors remain responsible for the scientific claims, code and submitted content.'
replace_paragraph('During the preparation of this manuscript','During the preparation of this manuscript, '+ai[0].lower()+ai[1:]) if 'During the preparation of this manuscript' in t else None
out=P/'MRP-ST-Repair_manuscript_Patterns_v29.md';out.write_text(t,encoding='utf-8')

def plain(s):return s.replace('**','').replace('*','').replace('`','').strip()
def settext(p,text):
 text=plain(text)
 if p.runs:
  p.runs[0].text=text
  for r in p.runs[1:]:r.text=''
 else:p.add_run(text)

d=Document(P/'Manuscript_Patterns.docx');miss=[]
for old,new in changes:
 a=plain(old)
 hits=[p for p in d.paragraphs if p.text.strip()==a]
 if not hits:
  prefix=a[:65];hits=[p for p in d.paragraphs if p.text.strip().startswith(prefix)]
 if len(hits)==1:settext(hits[0],new)
 else:miss.append({'prefix':a[:100],'matches':len(hits)})
repls={
 'Identity leakage in group-structured data, and why post-hoc repair cannot fix it':'Projection outcomes in the tested representations',
 'group-preserving null':'within-group label-permutation reference',
 'σ̂ group-preserving':'σ̂ within-group',
 'LDI group-preserving':'Scaled association (within-group SD)',
 'Gate decision changed':'Reference crossing changed',
}
for p in d.paragraphs:
 s=p.text
 if s.startswith('Figure 1. Identity leakage'):
  s=s.replace('Figure 1. Identity leakage in group-structured data, and why repairing a trained model afterwards cannot work.','Figure 1. Auditing the tested projection intervention.').replace('detect leakage with a calibrated statistic under leave-one-group-out','report unit-level label association alongside group holdout')
 for a,b in repls.items():s=s.replace(a,b)
 if p.text!=s:settext(p,s)
for tab in d.tables:
 for row in tab.rows:
  for c in row.cells:
   for p in c.paragraphs:
    s=p.text
    if 'github.com/WorldBMW' in s:s=doi+' [44]; historical snapshot at '+repo+'; current audit scripts supplied separately'
    for a,b in repls.items():s=s.replace(a,b)
    if s!=p.text:settext(p,s)
hits=[p for p in d.paragraphs if p.text.startswith('43. ')]
assert len(hits)==1
np=OxmlElement('w:p');hits[0]._p.addnext(np);p=Paragraph(np,hits[0]._parent);p.add_run(ref)
assert len(d.inline_shapes)==5 and len(d.tables)==7
d.save(P/'Manuscript_Patterns_v29.docx')
if miss:(Q/'unmatched_patches.json').write_text(json.dumps(miss,indent=2),encoding='utf-8')
print('Patched manuscript; unmatched',miss)
# Supplemental edits preserve all numbers but remove incorrect inference.
s=Document(P/'Supplemental_Information_Patterns.docx')
for p in s.paragraphs:
 st=p.text
 st=st.replace('这是「组级推断不可行」的直接证据','这仅表示组内置换方案退化，并不排除其他组级推断')
 st=st.replace('group-preserving','within-group label-permutation')
 st=st.replace('p = 0.000','p < 0.001')
 if st!=p.text:settext(p,st)
s.add_heading('Interpretation of the additional v28 analyses',1)
s.add_paragraph('The restricted permutation shuffles labels within recorded groups, not entire independent groups. Its zero variance under group-constant labels invalidates that particular ratio only; other group-level nulls may remain available. The archived ratio is not centered on the restricted null mean. The cervical grouping is a name-based assumption, not a verified source-patient mapping. The 21 GeoMx pairings are evaluated on preprocessed expression features, not learned repaired embeddings. The archived outer-refit analysis fitted preprocessing globally and averaged per-fold correlations. Fold SD is not an equivalence bound. Figure 5 uses the original logistic identity classifier; the newer nearest-neighbour ratio is a separate descriptive measurement with a different chance denominator. The Zenodo record 10.5281/zenodo.23640979 is a public historical snapshot and does not contain the complete current audit implementation.')
s.save(P/'Supplemental_Information_Patterns_v29.docx')
# A concise, outward-facing cover letter, keeping author-status claims conditional in the report.
cl='''# Cover letter

25 October 2026

Dear Editors of Patterns,

Please consider “Auditing label dependence and sibling-unit inflation in group-structured representations” as a Research Article.

The study addresses evaluation of representations learned from repeated observations within people, specimens and measurement batches. Its relevance to Patterns lies in a shared data-science problem: unit-level association, performance on unseen groups, direct identity structure and retained task utility can disagree. We report those quantities separately using spatial transcriptomic controls and exploratory wearable-activity and gas-sensor audits.

The main contribution is an auditable comparison of these measurements and their failure modes. A controlled simulation shows high unit-level association despite deterioration in group-level transfer. A fixed-scale interval calculation exposes a conflict in demanding both low association and retention of the same association; it is not a guarantee of an attainable intervention. The repair operator is established, suppression is partial, and no consistent downstream benefit is demonstrated. We make no claim of clinical utility, privacy protection or superiority over untested deconfounding families.

Additional analyses report nominal-label metrics, recoding sensitivity, within-group permutation scales, grouping assumptions and an exploratory encoder-refit comparison. They do not resolve all patient identifiers, establish fully nested validation for every cohort, or turn the association threshold into a group-level significance test. These limits are stated alongside the results.

Related manuscript. We disclose a related manuscript by the same group that also develops label-dependence indexing and uses overlapping public cervical accessions. It is not relied on as evidence in the present paper. A copy and its actual editorial status should accompany submission; independently recomputing a value does not by itself establish that results or text do not overlap.

Archive. A public manuscript and associated earlier code snapshot are available at https://doi.org/10.5281/zenodo.23640979 and https://github.com/WorldBMW/pathway-anchored-risk-benchmark. The frozen repository version does not contain the complete audit and reanalysis scripts used here. The corresponding current-revision code, result records and checksum manifest are prepared as a separate reproducibility package; the existing DOI is cited as a historical snapshot, not as a claim that the present analysis is fully reproduced by that deposit.

AI-assisted language editing and code review, including assistance with drafting analysis scripts, are disclosed in the manuscript. The authors remain responsible for the submitted content. Author approval, exclusive-submission status and related-manuscript status require confirmation before upload.

Sincerely,

Lang Guo, on behalf of the authors

School of Pharmacy, Shenyang Medical College, Shenyang 110034, China

guolang1016@symc.edu.cn
'''
(P/'CoverLetter_Patterns_v29.md').write_text(cl,encoding='utf-8')
sys.path.insert(0,str(Path(__file__).resolve().parent));from _md2docx import convert
convert(str(P/'CoverLetter_Patterns_v29.md'),str(P/'CoverLetter_Patterns_v29.docx'))
tp=(P/'TitlePage_Patterns.md').read_text(encoding='utf-8').replace('43 references','44 references')
tp+='\n**Public historical snapshot:** '+doi+'. Current-revision scripts are supplied separately because the linked frozen code archive is incomplete for these analyses.\n'
(P/'TitlePage_Patterns_v29.md').write_text(tp,encoding='utf-8');convert(str(P/'TitlePage_Patterns_v29.md'),str(P/'TitlePage_Patterns_v29.docx'))
print('Saved v29 source, preserved 5 figures / 7 tables, and synchronized companion files')
