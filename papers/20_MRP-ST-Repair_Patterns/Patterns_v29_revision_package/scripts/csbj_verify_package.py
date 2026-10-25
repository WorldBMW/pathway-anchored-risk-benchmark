from pathlib import Path
import json,zipfile,re,hashlib
import numpy as np
from docx import Document
from lxml import etree
from PIL import Image
R=Path('G:/project4_cervix_multimodal');O=R/'papers/12_CSBJ_audit_20261010';E=O/'evidence';Q=O/'qa';Q.mkdir(exist_ok=True)
a=json.loads((E/'corrected_probe.json').read_text(encoding='utf-8'));b=json.loads((E/'reproduced_probe.json').read_text(encoding='utf-8'))
for g,rr in a['groups'].items():
 for r,s in zip(rr,b['groups'][g]):
  for k in ['legacy_no_intercept','intercept_sensitivity']:
   for metric in ['tau','p_refit']:assert np.allclose(r[k][metric],s[k][metric],atol=1e-12)
  assert r['legacy_no_intercept']['q_BH_21']==s['legacy_no_intercept']['q_BH_21']
report={'portable_reproduction':'PASS: 21 primary + 21 intercept probe estimates/p-values agree; BH q agrees','renderer':'NOT COMPLETED: LibreOffice/soffice executable unavailable','documents':{},'figures':{}}
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
for p in O.glob('*.docx'):
 if p.name.startswith('~$'):continue
 d=Document(p);t='\n'.join(x.text for x in d.paragraphs);root=etree.fromstring(zipfile.ZipFile(p).read('word/document.xml'));bad=[]
 for tbl in root.xpath('//w:tbl',namespaces=ns):
  widths=[int(v) for v in tbl.xpath('./w:tblGrid/w:gridCol/@w:w',namespaces=ns)]
  if sum(widths)!=9360:bad.append(widths)
 assert not bad,(p,bad)
 assert '[[' not in t and '[@' not in t,p
 report['documents'][p.name]={'paragraphs':len(d.paragraphs),'tables':len(d.tables),'images':len(d.inline_shapes),'native_math_objects':len(root.xpath('//m:oMath',namespaces=ns)),'table_geometry':'PASS','unresolved_source_tokens':False}
for p in (O/'figures').glob('*.png'):
 with Image.open(p) as im:report['figures'][p.name]={'pixels':im.size,'dpi':im.info.get('dpi')}
main=Document(O/'Manuscript_CSBJ.docx');assert len(main.tables)==3 and len(main.inline_shapes)==7
text='\n'.join(p.text for p in main.paragraphs);abstract=text.split('Abstract\n')[1].split('\nKeywords:')[0];report['abstract_words']=len(abstract.split())
assert '−0.092' not in text and '0.551' not in text and 'D_joint' not in text and 'We do not propose' not in text
assert 'LDI=3.26' in text and '−0.105' in text and '0.497' in text and 'MRP-ST is defined by five mandatory clauses' in text
report['known_pending']=['author declaration confirmation','public version release and real Zenodo DOI','code/data license selection','Word page rendering review']
(Q/'verification_report.json').write_text(json.dumps(report,indent=2),encoding='utf-8');print(json.dumps(report,indent=2))
