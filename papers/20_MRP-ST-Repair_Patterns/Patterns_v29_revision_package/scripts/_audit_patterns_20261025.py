import json,re,sys,hashlib
from pathlib import Path
from docx import Document
sys.stdout.reconfigure(encoding='utf-8')
R=Path('results'); P=Path('papers/20_MRP-ST-Repair_Patterns')
for f in ['repair_e4_hnscc_stable.json','repair_prostate_gse278936_stable.json','unit_ladder_cervical.json','within_group_label_geomx.json','crossdomain_audit.json','dump_cervix.json']:
 a=json.loads((R/f).read_text(encoding='utf-8-sig'))
 print('\nFILE',f, 'keys',list(a))
 if 'scan' in a:
  for k,v in a['scan'].items():
   if isinstance(v,dict): print(k,{x:v[x] for x in v if re.search('tau|ldi|sigma|auc',x)})
 if 'summary' in a:
  for k,v in a['summary'].items(): print(k,{x:v[x] for x in v if not isinstance(v[x],(dict,list))})
 if f=='dump_cervix.json': print('samples',a['samples'])
for f in ['Manuscript_Patterns.docx','CoverLetter_Patterns.docx','TitlePage_Patterns.docx','Supplemental_Information_Patterns.docx']:
 d=Document(P/f); print('\nDOC',f,'paras',len(d.paragraphs),'tables',len(d.tables),'images',len(d.inline_shapes))
 if f.startswith('Supplemental'):
  for p in d.paragraphs:
   if 'Table S17' in p.text or 'permutation' in p.text.lower(): print(p.text)
t=(P/'MRP-ST-Repair_manuscript_Patterns_v26.md').read_text(encoding='utf-8')
for h in ['### The sibling-unit trap','### Decision boundary I','### Cross-domain audit']:
 s=t.index(h); e=t.find('\n### ',s+1)
 print('\nSECTION\n',t[s:e if e>0 else t.index('## DISCUSSION')])
