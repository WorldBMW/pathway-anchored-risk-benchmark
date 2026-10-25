from pathlib import Path
import shutil,sys
from docx import Document
from docx.shared import RGBColor
sys.stdout.reconfigure(encoding='utf-8')
P=Path('papers/20_MRP-ST-Repair_Patterns')
f=P/'MRP-ST-Repair_manuscript_Patterns_v27.md';t=f.read_text(encoding='utf-8')
old='cell type is a nominal label, so the sign of its τ depends on the arbitrary epithelium/stroma coding and the magnitude is the interpretable quantity.'
t=t.replace(old,'Cell type is binary; these results remain descriptive under the recorded unit-permutation reference and case-level grouping.')
f.write_text(t,encoding='utf-8')
d=Document(P/'Manuscript_Patterns.docx')
for p in d.paragraphs:
 if old in p.text:p.text=p.text.replace(old,'Cell type is binary; these results remain descriptive under the recorded unit-permutation reference and case-level grouping.')
d.save(P/'Manuscript_Patterns.docx')
s=Document(P/'Supplemental_Information_Patterns.docx')
for p in s.paragraphs:
 if p.text.startswith('GSE316098 is a NanoString'):
  p.text='GSE316098 contains 186 epithelial/stromal compartment measurements from 93 ROIs across seven specimens from six patients. The archived analysis groups by seven case identifiers; it does not establish patient-disjoint holdout. The matched-feature label arms are risk, Silva pattern and cell type. Their baseline kappa values are 0.2248, 0.2728 and 0.0000; baseline tau_LOO is 0.750, 0.797 and 0.709; sigma is 0.0803, 0.0768 and 0.0835; baseline Delta is +0.3736, +0.3955 and +0.0075. Baseline identity lift is 6.55x in all arms. Mean group-level retention under dosing is 0.806, 0.717 and 0.989. These are descriptive contrasts: label choice also changes the task and signal structure. The archived dual gate uses tau_LOO for both axes and accepts zero of twelve dosed cells in every arm. Unit-label permutation does not establish a valid group-level 5% test, and paired compartments from the same ROI require an appropriate restricted null.'
 if p.text.startswith('The cervical cohort, the spot embedding'):
  p.text='This resolution control partitions each source section into K pooled blocks, for K = 1, 3 and 8. The archived grouping variable is the section identifier (G = 35), not the 34-donor mapping of the main cervical archive. At K = 1, unit and section-group holdouts coincide by construction. At K > 1, units share their section label. The identity probe has no repeated test group at K = 1 and reports zero lift; at K > 1 its decoding is measured on repeated section groups. Baseline Delta is zero at both K = 1 and K = 8, so repeated units do not guarantee inflation. The declining unit-permutation sigma is resolution-dependent and does not create additional independent donors. Values are the original archived run outputs.'
 if p.text.startswith('Table S') or 'Heading' in p.style.name:
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
s.save(P/'Supplemental_Information_Patterns.docx')
read=P/'SUBMISSION_READINESS.md';bak=P/'SUBMISSION_READINESS_v26pre_20261025.md'
if not bak.exists():shutil.copy2(read,bak)
read.write_text('''# Patterns 投稿状态 v27

2026-10-25 独立复核后的状态：**尚不满足可直接投稿的完整证据要求**。

请优先阅读 REVIEW_REPORT_v27_20261025.md。该报告取代本文件旧版的“全部通过”“唯一风险是代码链接”“可保证初筛”等判断。

本轮已修订四份 Word，保留 *_v27pre_20261025.docx 备份。当前主稿源为 MRP-ST-Repair_manuscript_Patterns_v27.md，图形摘要使用 figures/graphical_abstract_v27.tif。旧 v16–v26 稿及旧中文稿不作为同步版本上传。

仍需完成：组保持的统计推断；宫颈与 GeoMx 的供体/患者映射；名义多分类评价；公共代码可访问性与环境重现；受限请求导致尚未完成的官方指南与文献核验；作者对相关稿件及投稿声明的事实确认。详见独立报告的证据和处置表。

结构和指定归档数值检查见 AUDIT_v27_checks.json，不代表全部原始实验已重跑，也不代表期刊一定录用。
''',encoding='utf-8')
print('Corrected supplemental interpretations and superseded stale submission-readiness claims')
