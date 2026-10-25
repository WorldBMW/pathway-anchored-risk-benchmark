"""Convert a markdown manuscript into a formatted .docx (Patterns submission)."""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches


def parse_inline(paragraph, text):
    """Render **bold**, *italic*, `code` into runs."""
    tokens = re.split(r'(\*\*.+?\*\*|\*.+?\*|`.+?`)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            r.font.name = 'Consolas'
        else:
            paragraph.add_run(tok)


def add_table(doc, rows):
    header = rows[0]
    body = rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        r = p.add_run(h.strip())
        r.bold = True
        r.font.size = Pt(9)
    for row in body:
        cells = t.add_row().cells
        for i, c in enumerate(row):
            if i >= len(cells):
                break
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            txt = c.strip().replace('**', '')
            r = p.add_run(txt)
            r.font.size = Pt(9)
    # narrow margins inside cells
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
    return t


def convert(md_path, docx_path, title=None):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    # Join wrapped lines whose bold span is broken across a line break, so that
    # "**Figure 1. ...\n...**" renders as bold text instead of literal asterisks.
    merged = []
    buf = ''
    for ln in lines:
        s = ln.strip()
        if buf:
            buf = buf + ' ' + s
            if s.count('**') % 2 == 1:
                merged.append(buf)
                buf = ''
            continue
        if s.startswith('**') and s.count('**') % 2 == 1:
            buf = s
            continue
        merged.append(ln)
    if buf:
        merged.append(buf)
    lines = merged

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    for s in doc.sections:
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped or stripped == '---':
            i += 1
            continue

        # collect a markdown table
        if stripped.startswith('|'):
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                raw = lines[i].strip().strip('|')
                cells = raw.split('|')
                if all(set(c.strip()) <= set('-: ') for c in cells):
                    i += 1
                    continue
                rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
                doc.add_paragraph()
            continue

        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(stripped[2:])
            r.bold = True
            r.font.size = Pt(16)
            r.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        if stripped.startswith('## '):
            h = doc.add_heading(level=1)
            r = h.add_run(stripped[3:])
            r.font.color.rgb = RGBColor(0, 0, 0)
            r.font.name = 'Times New Roman'
            i += 1
            continue

        if stripped.startswith('### '):
            h = doc.add_heading(level=2)
            r = h.add_run(stripped[4:])
            r.font.color.rgb = RGBColor(0, 0, 0)
            r.font.name = 'Times New Roman'
            i += 1
            continue

        if stripped.startswith('**Figure') or stripped.startswith('**Table'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            parse_inline(p, stripped)
            i += 1
            continue

        # Numbered items (highlights, contribution list, references). We keep the
        # number as literal text rather than relying on Word's automatic list
        # numbering: python-docx's "List Number" style carries no per-paragraph
        # numPr, so every numbered block in the document shares one counter and
        # the reference list would restart or continue unpredictably. Literal
        # numbers render identically everywhere and cannot desynchronise from
        # the in-text [n] citations.
        if re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            parse_inline(p, stripped)
            i += 1
            continue

        # Explicit bracket numbering "[1]" is emitted verbatim.
        if re.match(r'^\[\d+\]\s', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            parse_inline(p, stripped)
            i += 1
            continue

        if stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            parse_inline(p, stripped[2:])
            i += 1
            continue

        p = doc.add_paragraph()
        parse_inline(p, stripped)
        i += 1

    doc.save(docx_path)
    print('wrote', docx_path)


if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
