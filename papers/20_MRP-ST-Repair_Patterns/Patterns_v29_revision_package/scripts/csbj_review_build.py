"""Build a scientifically corrected CSBJ submission draft without changing source results."""
from pathlib import Path
import re,json,shutil,hashlib,zipfile,sys
ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'papers/13_CSBJ_v2_20261014'
OUT=SRC/'CSBJ_submission_checked_20261014'
OUT.mkdir(exist_ok=True)
QA=SRC/'submission_review'

def section(s,heading,body):
    pat=r'(?ms)^'+re.escape(heading)+r'\n.*?(?=^### |^## |\Z)'
    s,n=re.subn(pat,lambda m:heading+'\n'+body.strip()+'\n\n',s,count=1)
    assert n,heading
    return s

def revise():
    s=(SRC/'Manuscript_CSBJ.md').read_text(encoding='utf-8-sig')
    title='MRP-ST: Auditing supervision dependence and representation information in spatial transcriptomic models'
    s=re.sub(r'^# .*', '# '+title,s,count=1)
    s=s.replace('Ye Zhang², Lang Guo³*','Ye Zhang², Zhiping Ni¹, Lang Guo³*')
    s=s.replace('Ye Zhang: 13332457636@163.com; Lang Guo','Ye Zhang: 13332457636@163.com; Zhiping Ni: nizhiping@jm360.cn; Lang Guo')
    s=section(s,'## Abstract', '''Spatial transcriptomic scores can reflect imposed supervision, expression structure or sample identity, which require different validation experiments. We present MRP-ST, a five-clause reporting protocol covering supervision ablations, statistical units, score association, label-free representation probing and endpoint-specific external checks. A reference model is audited using known-signal simulations and 35 cervical sections. Removing ordinal supervision reduces mean score–severity Kendall tau-b from 0.460 to 0.251. A label-free frozen representation retains seed-unstable association: mean leave-one-section-out probe tau-b is 0.270, with one of three seeds significant under refitted label permutations. These probes remain transductive in covariates and do not establish patient-level generalization. Identity-confounded simulations illustrate how label-trained encoders can produce near-maximal leave-one-out probe correlations despite lacking an injected shared disease program. In a dose–response grid, mean probe correlation falls from 0.704 to 0.182 at 90 samples and from 0.457 to 0.177 at 35 samples, with a nonmonotone weak-signal response at the larger scale. Grid results are descriptive because the stored permutation procedure did not refit label-dependent predictions. Illustrative readouts from SpaGCN, STAGATE and GraphST use different readout procedures from the reference probe and do not establish comparative performance. External cervical GeoMx scores show a reversed tumor–benign direction across six paired patients; spatial organization in head-and-neck cohorts does not establish phenotype discrimination. MRP-ST provides an executable framework for stating what each analysis supports while separating supervision effects, representation association and biological transport.

Keywords: spatial transcriptomics; model auditing; supervision ablation; representation probing; permutation testing; reproducibility''')
    # Replace novelty and causal overclaims with supported positioning.
    s=s.replace('but none determines what information a spatial score has actually learned.','provide the reporting context for a task-specific audit of spatial scores.')
    a=s.index('We introduce MRP-ST,');b=s.index('\n\n## 2.',a)
    s=s[:a]+'''We introduce MRP-ST, a Minimum Reporting Protocol for Spatial Transcriptomic model audits. The protocol organizes five experiments that distinguish supervision-sensitive score orientation, annotation-associated representation information and endpoint-specific transport. It complements TRIPOD+AI, REFORMS and DOME [@tripod;@reforms;@dome] by specifying these experiments for spatial transcriptomic scores. We report a descriptive permutation-standardized score association statistic, explicit supervision ablations, refitted probes of label-free embeddings and controlled simulations. Synthetic cohorts test known injected programs and identity structure; cervical spatial transcriptomics supplies a worked demonstration. A score can lose its disease ordering after objective removal while its representation retains information. Conversely, a probe of embeddings trained on all sample labels can retain the label axis even when the probe itself leaves a sample out. The proposed protocol is a task-specific framework rather than a validated community consensus standard, and the current simulation grid does not establish its population error rates.'''+s[b:]
    s=s.replace('### 2.1. Audit objects and statistical units','### 2.1. Experimental and technical design: audit objects and statistical units')
    # Do not invent publication status of the reference model.
    s=s.replace('an in-house architecture developed by the authors for this programme of work and not previously published; it is described here in full so that the audit can be reproduced, and its status is declared to the journal.','an in-house architecture developed by the authors for this programme of work. Its implementation and objectives are specified here; related manuscripts must be disclosed separately in the submission record.')
    a=s.index('The Label-Dependence Index (LDI)');b=s.index('\n\nReal-data FULL',a)
    s=s[:a]+'''The earlier working name Label-Dependence Index (LDI) is replaced here by the standardized score association statistic Z_score, because a fixed-score permutation reference cannot identify a training-time dependence mechanism:

$$ Z_{score}=\\frac{\\overline{\\tau}_{FULL}-\\mu_{perm}}{\\sigma_{perm}}. $$

The descriptive reference distribution permutes complete section labels against the fixed FULL seed-0 scores (B=10,000, permutation seed 20261005), yielding mean 0.001, SD 0.141 and an empirical 95th percentile of 0.236. Standardizing the mean FULL correlation of 0.460 gives Z_score=3.26. This combines an across-seed observed summary with a single-seed fixed-score reference and is reported only as a descriptive statistic, not a calibrated significance test. The fitted scores were trained with these labels; permuting labels without retraining that model is not a test of learning, generalization or causal label dependence. The Gaussian reference value 1.645 is not exactly equivalent to the empirical percentile. C3 therefore relies on the stated supervision channels and ablation contrasts, rather than this statistic alone (Figure 8).

The NOALL scalar head has a large offset and a very small between-section spread. Its computed tau changes with aggregation precision, making its rank ordering numerically sensitive. The historical summary is −0.105 (seed SD 0.084), whereas recomputation from the same stored arrays in float64 gives approximately −0.092 (SD 0.095). These are numerical summaries rather than evidence of an absent or inverted biological axis. Loss of scalar ordering does not imply absence of information in the embeddings. Positive affine transformations preserve ranks mathematically, but finite-precision aggregation can alter nearly tied sample means; a small spread-to-offset ratio alone does not prove that every ordering difference is rounding error.'''+s[b:]
    s=s.replace('using fixed thresholds (0.35 upper, 0.15 lower, α=0.05) that were set before the cervical analysis and applied unchanged to every cohort.','using exploratory thresholds (0.35 upper, 0.15 lower, α=0.05). The repository does not establish prospective registration or selection independently of these analyses.')
    s=s.replace('§3.1 measures this rule\'s error rates directly on known-signal cohorts.','§3.1 reports descriptive behavior on known-signal cohorts; valid population error-rate estimation requires additional independent cohorts and correctly refitted permutation inference.')
    s=s.replace('### 2.5. Known-signal simulations: dose–response design and operating characteristics','### 2.5. Known-signal simulations and descriptive dose–response design')
    s=s.replace('Probe p-values use the same refitted permutation procedure as the cervical audit (B=10,000).','The archived grid calls scripts/probe_lodo.py, whose permutation loop shuffles labels against unchanged fitted predictions. These grid p-values are not refitted permutation tests and are excluded from inferential conclusions and evidence grades. Complete grid embeddings were unavailable in the local review snapshot, so grid-wide corrected p-values could not be reconstructed.')
    s=s.replace('A readout that orders severity on CTRL cohorts can only do so by memorizing donor identity.','No shared label-associated program is injected. Supervised in-sample separation may reflect identity memorization, while chance association remains possible in finite generated cohorts.')
    a=s.index('The grading rule of Table 2 is a decision procedure');b=s.index('\n\nThe cervical example',a)
    s=s[:a]+'''We distinguish a run-level correlation threshold from the cohort-level C4 rule. The run-level threshold has descriptive sensitivity and control crossing fractions:

$$ \\mathrm{sens}_{run}(\\theta)=\\frac{\\#\\{L_1\\ runs:\\hat\\tau\\ge\\theta\\}}{\\#L_1\\ runs},\\qquad \\mathrm{cross}_{CTRL}(\\theta)=\\frac{\\#\\{CTRL\\ runs:\\hat\\tau\\ge\\theta\\}}{\\#CTRL\\ runs}. $$

At each scale there are three independently generated cohorts and three optimizer runs per cohort. All nine runs are shown, but they are clustered repetitions rather than nine independent Bernoulli trials. We report counts without run-level Clopper–Pearson intervals. The cohort-level C4 rule also requires valid refitted permutation results across seeds, which are unavailable for this grid; therefore no cohort grades, nominal type-I control or calibrated protocol error rates are inferred from it. The 0.35 threshold and swept ranges are exploratory.

With binary labels and no ties in predictions, the attainable tau-b under perfect cross-class separation is

$$ \\tau_{max,\\ no\\ score\\ ties}=\\sqrt{\\frac{n_0 n_1}{\\binom{n_0+n_1}{2}}}, $$

giving 0.707 (50/40), 0.457 (31/4), 0.717 (12/9) and 0.425 (19/2). This is conditional on untied predictions, not a universal bound imposed by binary labels. Predictions tied within each class can attain tau-b=1. Rounded published-method values and possible tied readouts therefore cannot be interpreted against this formula without checking prediction ties. Exact equality for an untied prediction vector means all cross-class pairs are concordant; a rounded near-equality alone does not establish perfect separation.'''+s[b:]
    s=s.replace('Four lesions define the limited disease diversity; n=4 is not the total sample count in correlation tests.','Only four sections have lesion annotations; all 35 sections enter section-level association analyses.')
    s=section(s,'### 2.6. Unified audit of published embedding methods','''Three published unsupervised spatial embedding methods, SpaGCN [@spagcn], STAGATE [@stagate] and GraphST [@graphst], provide illustrative readouts on selected simulated cohorts. They are not trained to solve the same cross-section supervised task as the reference model. In scripts/benchmark_st_methods.py, SpaGCN uses histology=False, l=0.5, 50 PCs, Louvain initialization and resolution 1.0; STAGATE uses radius 150, 1,000 epochs and learning rate 0.001; GraphST uses 600 epochs and output width 64. STAGATE and GraphST are clustered by k-means with k=min(15,max(3,floor(sqrt(n_spots)))), n_init=10 and random_state=42. The square root is taken over spots within a section, not over sections. The implementation is specified rather than asserted to reproduce every published default.

The pathway panel is loaded from docs/hpv_genes.json; a hard-coded fallback is used if loading fails. Cluster mean raw-count panel expression defines each cluster value, and its maximum gives a section-level scalar. This differs from the reference probe's pooled embedding readout and is sensitive to library size, pathway coverage and the selected panel. Exact gene lists and fallback behavior are included with the code. Method embeddings may be label-free while this readout remains biologically selected using the injected program; the resulting task is not a blind biological discovery test.

The stored function named lodo_from_sample_risks actually calculates direct tau-b between fixed section risks and annotations, then permutes annotations 2,000 times (random seed 0). It neither fits a linear probe nor performs leave-one-section-out prediction. Such a fixed-score permutation test is appropriate only for a label-independent prespecified scalar under exchangeability. We independently recompute these permutation results from the archived risk vectors. These p-values are unadjusted exploratory results across the 18 method-by-cohort comparisons.

The selected subsets contain 21 sections (12 normal + 9 diseased at the main scale; 19 + 2 at the matched scale), from generator seed 20261006 only. The reference column reports correlations of leave-one-out ridge predictions from NOALL embeddings on these same sections, averaged over three optimizer seeds. Its archived subset p-values are not supported by a verified refitting implementation and are omitted. Single method runs, different readouts and multiplicity preclude comparative grades or claims of superiority. Systematic spatial-domain benchmarking is a different task [@stbench].''')
    s=s.replace('C3 Label dependence. Report τ̄, the permutation-null calibration, LDI and the θ-gate decision, together with the numeric precision used to compute the statistic.','C3 Supervision dependence. Declare every label-derived objective, report τ̄ and supervision-removal contrasts, and state numeric precision and rank sensitivity. A fixed-score permutation association statistic is descriptive when training used the same labels; it cannot by itself establish causal dependence or generalization.')
    s=s.replace('Probes of label-trained transductive embeddings must be reported as leakage-contaminated upper bounds, not as held-out estimates (§3.2).','Probes of label-trained transductive embeddings must be identified as contaminated diagnostics rather than held-out estimates (§3.2). They are not guaranteed mathematical upper bounds on valid predictive performance.')
    a=s.index('Clauses C3 and C4 are the operative core:');b=s.index('\n\n## 3.',a)
    s=s[:a]+'''Clauses C3 and C4 distinguish changes in score orientation from association retained in a label-free representation. Table 2 provides an exploratory grading rule for the latter. The grading rule is used only when valid refitted p-values are available; descriptive correlation thresholds are reported separately. Figures 8 and 9 show numerical association references and contamination controls, without interpreting either as a calibrated label-dependence test.'''+s[b:]
    s=s.replace('### 3.1. Measured operating characteristics of the grading rule','### 3.1. Descriptive dose–response and correlation-threshold crossings')
    s=section(s,'### 3.1. Descriptive dose–response and correlation-threshold crossings','''[[TABLE4]]

Mean leave-one-out probe tau-b was 0.704 (run SD 0.009), 0.276 (0.154), 0.125 (0.122) and 0.182 (0.104) at L1–L4 of the 90-section scale; at the matched 35-section scale it was 0.457 (0.000), 0.394 (0.095), 0.286 (0.096) and 0.177 (0.148) (Table 3; Figure 2). The larger-scale curve decreases from L1 to L3 and rebounds at L4; the matched-scale means decrease monotonically. Run SD mixes between-generator and within-generator variation and does not establish statistical equivalence between adjacent levels. L1 is near the untied-prediction separation value: eight of nine main-scale runs attain 0.707 and one is 0.678; all nine matched-scale runs attain 0.457.

The CTRL means are −0.030 (SD 0.111) and +0.042 (0.102). At the exploratory correlation threshold 0.35, nine of nine L1 runs at each scale exceed the threshold and none of the 18 pooled CTRL runs does. Swept thresholds with identical observed separation lie in [0.20,0.675] and [0.20,0.45]. These ranges describe this finite grid, not confidence intervals or prospectively validated error rates. Reused optimizer runs do not increase the number of independent generated cohorts. Stored significance counts are withdrawn because label permutations did not refit the probe; neither the reported nominal 5% false-positive compatibility nor cohort support grades are retained.''')
    s=section(s,'### 3.2. The control arm dissociates memorization, encoder leakage and representation information','''On CTRL cohorts, FULL heads have in-sample tau-b ranges of 0.415–0.643 (mean 0.490) at 90 sections and 0.162–0.280 (mean 0.237) at 35 sections (Figure 3). Leave-one-out probes of label-trained transductive embeddings yield 0.692–0.707 and 0.457, respectively, while label-free embeddings yield −0.030 (SD 0.111) and +0.042 (0.102). The mean differences are approximately 0.73 and 0.41 tau units.

The held-out probe section was not held out from supervised encoder training: its own annotation entered encoder optimization. These correlations consequently illustrate contamination despite exclusion from the linear probe fit. Exposure of all covariates alone is transductive learning; exposure of held-out labels is the leakage mechanism. The differences are observed contrasts, not a causal decomposition of leakage magnitude or valid generalization estimates. Archived FULL-embedding permutation p-values are omitted because a frozen representation trained with original labels does not supply a null model for permuted-label training. Label-free grid significance also requires correctly refitted probes. The earlier identity-confounded cohorts in Supplementary Table S1 provide a separately archived, correctly refitted analysis.''')
    s=s.replace('### 3.3. Published methods inherit the same difficulty gradient','### 3.3. Illustrative published-method readouts on selected subsets')
    s=section(s,'### 3.3. Illustrative published-method readouts on selected subsets','''[[TABLE5]]

At L1, direct pathway readouts of SpaGCN and GraphST correlate positively with the main-scale annotations (tau-b 0.677 and 0.491); GraphST has tau-b 0.425 at the matched scale. STAGATE yields −0.173 and −0.045. We report recomputed fixed-score permutation p-values in Table 4, with unadjusted inference clearly separated from the reference ridge correlations. The latter average 0.717 and 0.425 at L1. The two readout types share sections and a statistic but not an evaluation algorithm, so these values do not establish relative model performance.

SpaGCN produces the same rounded matched-scale readout correlation, 0.291, at L1, L3 and L4; GraphST produces 0.336 at both L3 and L4. Equal correlations can result from identical rank orderings, tied predictions or discreteness and do not prove identical risk vectors or readout saturation. STAGATE has a negative main-scale L4 correlation of −0.465. Since the scalar is the maximum raw-count panel expression, its sign is biologically oriented by the chosen panel rather than by an arbitrary embedding sign. The negative association may reflect clustering, library size or sampling; this single-run experiment does not identify the cause. One-sided positive-association p-values do not test association in either direction. Multiplicity and different readouts constrain all comparisons (Figure 4).''')
    s=s.replace('The calibrated label dependence was LDI=3.26 against θ=1.645, so the gate is triggered and the ordering is label-dependent; D_ordinal was 0.209.','The descriptive standardized fixed-score association is Z_score=3.26; D_ordinal is 0.209. Supervised ordering and the ablation contrast document sensitivity to the ordinal objective, without using Z_score as a causal test.')
    s=s.replace('it is null-compatible rather than a negative measurement.','it is numerically sensitive and is not treated as evidence of a null biological association.')
    a=s.index('The transductive scope of the audit is quantified');b=s.index('\n\nMoran\'s I remained',a)
    s=s[:a]+'''Two distinct endpoints are available in the cervical data. For section-level lesion discrimination (four diseased versus 31 normal), the in-cohort FULL score AUC is 1.00 and mean panel expression AUC is 0.75. These are descriptive resubstitution results. A separate morphology-to-expression baseline predicts panel expression in leave-one-section-out folds and uses the held-out section's own median panel expression to define a spot-level proxy. This is neither laboratory HPV status nor section-level lesion classification. Mean fold AUC is 0.497 (SD 0.021) for ridge and 0.497 (0.020) for MLP, across 16 and 33 evaluable folds, respectively; single-class or constant-prediction folds are excluded. The fold-specific median differs from the pooled median used by the reference model, and preprocessing was not fully confined to training folds. These metrics cannot be subtracted from the lesion AUC or described as a drop from 1.00 to chance by the same model. The same baseline has mean expression-reconstruction Pearson correlation −0.0145 across 35 folds. We do not have patient-held-out disease AUC for the full reference pipeline.

Diagnostic probes on all stored ablation embeddings return mean leave-one-section-out correlations of approximately 0.420–0.462 for several target-informed variants (Figure 9). Their archived p-values shuffled labels against unchanged predictions and are not used here. In addition, variants differ in contamination channels: FULL and NOPROG retain expression-derived proxy supervision, FULL also retains ordinal supervision, and NOSUP retains pathway alignment while removing both explicit targets. NOSUP cannot be described as having trained on all ordinal labels. The NOALL primary probe above is supported by a correctly refitted permutation analysis; its SD is 0.116, whereas the diagnostic script's differing aggregation gives 0.109. The primary specification is used for inferential reporting.

Fixed-score permutation references remain stable across 20 combinations of nominal alpha and permutation seed, but this is numerical reference stability rather than validation of a label-dependence decision. NOALL has a median absolute score around 1.4×10⁵ and between-section spread around 0.298; NOSUP has spread around 0.359 and magnitude around 0.304. These diagnose distinct numerical regimes, without proving all NOALL ranks are floating-point noise or validating a dependence gate.'''+s[b:]
    # Rewrite Discussion/Conclusions to remove unsupported claims left across sections.
    s=section(s,'## 4. Discussion','''MRP-ST organizes complementary questions about spatial scores: which objectives orient a scalar, what information is recoverable from a label-free representation and which endpoints remain interpretable outside the source cohort. These distinctions complement general reporting guidance [@tripod;@reforms;@dome]. The protocol's five clauses are a proposed task-specific framework, not evidence that all necessary validation has been completed.

The identity-confounded simulations demonstrate a failure of validation boundaries. A sample excluded from ridge fitting can still have contributed its label to encoder training. Near-separation in such embeddings is a contaminated diagnostic rather than transferable information. Transductive exposure of covariates must be stated separately from held-out-label exposure. Label-free encoders remove direct target supervision, but pathway selection, architecture selection and acquisition confounding still affect interpretation.

The grid's correlation curves are useful descriptive evidence of signal difficulty. They do not validate protocol error rates: optimizer runs are clustered within three generated cohorts per cell, and the stored permutation procedure did not refit the label-dependent probe. The full grid requires reanalysis from embeddings with corrected permutations and more independent cohorts. The earlier three-scenario evidence package and cervical primary probes do use algebraically refitted permutation inference. These distinct evidence families are mapped in Supplementary Materials.

Published-method illustrations have additional limitations. They use single method runs and a raw-count cluster-max pathway readout rather than the reference model's ridge probe. Their panel is selected using the injected program, making them tests of a specific readout rather than blind biological discovery. Correlations shared across levels do not demonstrate identical outputs, and negative pathway-score correlations cannot be explained solely by an arbitrary embedding sign. Multiplicity and mismatched readouts prevent comparative rankings [@stbench].

In cervical tissue, source-confounded annotations and four lesion sections limit disease attribution. Frozen NOALL embeddings contain seed-unstable section-label association, conditional on transductive covariates and label exchangeability. A significance result in one seed does not establish patient-level prediction. The fixed-score statistic Z_score measures apparent association, not the fraction of information imposed by supervision. Fully inductive evaluation must refit preprocessing, encoder, stopping rules and readout within a patient-held-out boundary.

Biological associations reuse expression input and therefore do not independently establish pathway mechanism. External GeoMx tumor–benign contrasts have a consistent reversed direction across six paired patients. Head-and-neck spatial organization and phenotype separation are separate endpoints. TCGA gene screens are not external application of the same spatial predictor. Likewise, lesion AUC and spot-level expression-proxy AUC measure different tasks; no full-model out-of-sample disease AUC is available here.

The reproducibility package makes these boundaries inspectable: variant definitions, pooled representations for the primary refitted probes, source scripts and result extracts are supplied. Full model retraining requires source datasets and foundation-model weights. Missing complete grid embeddings and absent historical training commits limit reproduction of some stored results. Prospective validation of the protocol and any clinical application remains future work.''')
    s=section(s,'## 5. Conclusions','''MRP-ST provides a structured audit of supervision-sensitive scores, label-free representation association and endpoint-specific transport. In the cervical demonstration, removing ordinal supervision reduces mean score tau-b from 0.460 to 0.251, while a label-free transductive representation retains seed-unstable probe association (mean 0.270; one of three seeds significant under refitted permutations). Identity-confounded simulations illustrate contamination when held-out probe labels entered encoder training. The dose–response grid and third-party readouts are descriptive and do not establish calibrated error rates or comparative superiority. External score reversal and weak phenotype discrimination restrict the biological interpretation that can be transported. Reporting these boundaries makes supported computational claims distinguishable from untested generalization and clinical claims.''')
    s=s.replace('Competing interests: Mingwei Bai is affiliated','Competing interests: Mingwei Bai and Zhiping Ni are affiliated')
    # Actual CRediT role cannot be inferred from the request to add an author.
    s=s.replace('Lang Guo: conceptualization, methodology, supervision, project administration and writing—review and editing.','Zhiping Ni: data curation, resources and writing—review and editing. Lang Guo: conceptualization, methodology, supervision, project administration and writing—review and editing.')
    s=s.replace('All analysis code, model configurations and the complete audit evidence package (pooled representations, result extracts, SHA-256 checksums and corrected-analysis scripts) are openly available at','The existing analysis repository and v1.0 release are available at')
    s=s.replace('this submission corresponds to release v1.0 of that repository','the original audit archive is identified by release v1.0 of that repository')
    s=s.replace('releases/tag/v1.0).','releases/tag/v1.0). This revision includes additional corrections and supplemental evidence that require versioned deposition with the submitted files. The release is not asserted to contain complete grid embeddings or every result of this revision.')
    s=s.replace('queue-level','cohort-level').replace('leave-one-donor-out','leave-one-section-out')
    s=s.replace('The pathway panel is loaded from docs/hpv_genes.json; a hard-coded fallback is used if loading fails.','The pathway panel is loaded from docs/hpv_genes.json, which contains 40 genes in the reviewed source snapshot, rather than the reference model’s nominal 108-gene panel; a 40-gene hard-coded fallback is used if loading fails.')
    s=s.replace('Only the middle seed was significant, including after adjustment across 21 primary tests.','Only the middle seed was significant, including after adjustment across 21 primary tests. Independent reproduction from the packaged arrays matches all 21 archived primary tests when their original array dtype is retained. Float64 preprocessing gives cervical correlations of 0.215, 0.397 and 0.200 and refitted p-values of 0.1041, 0.0027 and 0.1126; the one-significant-seed conclusion is unchanged. This sensitivity concerns preprocessing arithmetic in addition to scalar-head aggregation.')
    s=s.replace('The stored primary representation','The stored primary representation')
    s=s.replace('Four-head attention combines the modalities;','The fusion module uses four-head attention with a single image key/value token and an expression-query residual followed by layer normalization. With one key, the softmax attention weight is trivially one, so this implementation does not learn spatial or multi-patch attention weights;')
    s=s.replace('Sources are GEO GSE208654 and GSE217237 plus a public 10x Genomics cervical-cancer section.','Sources are GEO GSE208654 and the GSE217237 SuperSeries (normal ectocervical Visium samples from subseries GSE290350, GSE299708 and GSE300112), plus a public 10x Genomics cervical-cancer section. The normal comparator cohort includes HIV, contraceptive and microbiome study contexts and is not a homogeneous healthy reference population.')
    (OUT/'Manuscript_CSBJ.md').write_text(s,encoding='utf-8')
    return title

def module():
    # Reuse deterministic table and equation helpers without executing evidence-copy side effects.
    source=(ROOT/'scripts/csbj_build_package_v2.py').read_text(encoding='utf-8')
    source=source.replace("O=R/'papers/13_CSBJ_v2_20261014';OLD=", "O=R/'papers/13_CSBJ_v2_20261014/CSBJ_submission_checked_20261014';OLD=")
    source=source.replace("E=O/'evidence';F=O/'figures';E.mkdir", "E=R/'papers/13_CSBJ_v2_20261014/evidence';F=O/'figures';E.mkdir")
    source=source.replace('\npackage_evidence()\n','\n')
    source=source.replace("REF=read(E/'references_verified.json')","REF=read(R/'papers/13_CSBJ_v2_20261014/submission_review/references_corrected.json')")
    # no original-module main block
    source=source[:source.index("if __name__=='__main__':")]
    g={'__file__':str(ROOT/'scripts/csbj_build_package_v2.py'),'__name__':'csbj_review_helpers'}
    exec(compile(source,str(ROOT/'scripts/csbj_build_package_v2.py'),'exec'),g)
    return g

def build():
    import numpy as np
    import matplotlib.pyplot as plt
    from scipy.stats import kendalltau
    title=revise();g=module()
    g['plt'].rcParams['font.family']='DejaVu Serif'
    # Correct native Word equations, keeping original nine-equation positions.
    g['EQS'][4]=[[g['mr']('Z_score = '),g['frac']('τ̄_FULL − μ_perm','σ_perm')]]
    g['EQS'][7]=[[g['mr']('sens_run(θ) = # { L1 runs: τ̂ ≥ θ } / # L1 runs')],[g['mr']('cross_CTRL(θ) = # { CTRL runs: τ̂ ≥ θ } / # CTRL runs')]]
    g['EQS'][8]=[[g['mr']('τ_max, no score ties = '),g['rad'](g['fraction']([g['mr']('n₀ n₁')],'C(n₀+n₁, 2)'))]]
    # Recompute appropriate fixed-score permutations for third-party readouts.
    bm=g['BM'];tests={};B=10000
    cached=OUT/'third_party_fixed_score_checks.json'
    if cached.exists():tests=json.loads(cached.read_text(encoding='utf-8'))
    for cohort,cc in ([] if tests else bm['groups'].items()):
        for method,r in cc['methods'].items():
            names=list(r['sample_risks']);x=np.array([r['sample_risks'][n] for n in names]);y=np.array([2 if n.startswith('CA_') else 0 for n in names])
            obs=float(kendalltau(x,y).statistic);rng=np.random.default_rng(20261014)
            ge=sum(kendalltau(x,rng.permutation(y)).statistic>=obs-1e-12 for _ in range(B))
            tests[cohort+'|'+method]={'tau':obs,'p':(ge+1)/(B+1),'n_perm':B,'n':len(x),'n_unique_risks':len(np.unique(x))}
    (OUT/'third_party_fixed_score_checks.json').write_text(json.dumps(tests,indent=2),encoding='utf-8')
    g['T3R'][0][2]='Repeatable positive association under this exploratory rule; not patient-level validation'
    g['T4H']=['Level (fraction f)','90 sections mean τ (SD)','35 sections mean τ (SD)']
    g['T4R']=[[r[0],r[1],r[4]] for r in g['T4R']]
    g['T4T']='Table 3. Descriptive probe correlations on the simulation grid. Nine optimizer runs per cell are nested within three generated cohorts. Run SD is descriptive. Stored grid p-values did not refit predictions and are excluded; no evidence grades or population error rates are assigned. Separation reference values assume no prediction ties.'
    g['T5H']=['Subset','SpaGCN τ (p)','STAGATE τ (p)','GraphST τ (p)','Reference τ̄ (SD)']
    cohorts=['group_A_g20261006_L1','group_A_g20261006_L3','group_A_g20261006_L4','group_A_g20261006_L1_small','group_A_g20261006_L3_small','group_A_g20261006_L4_small']
    rows=[]
    for c,old in zip(cohorts,g['T5R']):
        vals=[tests[c+'|'+m] for m in ['SpaGCN','STAGATE','GraphST']]
        rr=g['OSP'][c];rows.append([old[0]]+[f"{v['tau']:.3f} ({v['p']:.4f})" for v in vals]+[f"{np.mean(rr['taus']):.3f} ({np.std(rr['taus'],ddof=1):.3f})"])
    g['T5R']=rows
    g['T5T']='Table 4. Different readouts on the same selected 21-section subsets. Published-method columns use direct pathway scalar association and 10,000 fixed-score label permutations, recomputed here; all p-values are unadjusted exploratory values. Reference correlations are leave-one-out ridge probe means (three optimizer seeds); subset probe p-values are omitted because refitting could not be verified. These results are not comparable method grades.'
    # Alter table widths for three columns.
    code=(ROOT/'scripts/csbj_build_package_v2.py').read_text(encoding='utf-8')
    fn=code[code.index('def build_main():'):code.index("if __name__=='__main__':")]
    fn=fn.replace('[2050,1500,850,1150,1500,850,1460]','[3120,3120,3120]')
    exec(compile(fn,'review_build_main','exec'),g)
    # All final main figures are regenerated; no contaminated significance is plotted.
    g['figures']()
    save=g['savefig'];S=g['S'];COL=g['COL'];P=g['P'];REG=g['REG'];CFP=g['CFP'];OSP=g['OSP']
    fig,axs=plt.subplots(1,2,figsize=(9,3.8))
    for size,c,m,label in [('main',COL[0],'o','90 sections'),('small',COL[1],'s','35 sections')]:
        rr=[S['detection_rate'][size][lv] for lv in ['L1','L2','L3','L4']]+[S['type_i_error_control'][size]['CTRL']]
        axs[0].errorbar(range(5),[v['mean_tau'] for v in rr],yerr=[v['sd_tau'] for v in rr],marker=m,color=c,capsize=3,label=label)
        roc=S['threshold_roc'][size];axs[1].plot(roc['threshold'],roc['sensitivity'],color=c,marker=m,markevery=5,label=label+' L1 crossings');axs[1].plot(roc['threshold'],roc['one_minus_specificity'],ls='--',color=c,label=label+' CTRL crossings')
    axs[0].set(xticks=range(5),xticklabels=['L1','L2','L3','L4','CTRL'],ylabel='Probe tau-b (mean ± run SD)',title='A  Descriptive signal response');axs[0].axhline(0,color='gray',lw=.7);axs[0].legend(fontsize=8)
    axs[1].set(xlabel='Exploratory tau threshold',ylabel='Observed run crossing fraction',title='B  Threshold crossings');axs[1].axvline(.35,color='black',ls=':');axs[1].legend(fontsize=7)
    fig.tight_layout();save(fig,'Fig2_Operating_characteristics')
    fig,axs=plt.subplots(1,2,figsize=(9,3.9),sharey=True)
    for ax,suffix,label in zip(axs,['','_small'],['A  90 sections','B  35 sections']):
        cs=['group_C_g2026100'+str(i)+'_CTRL'+suffix for i in [6,7,8]]
        vv=[[r['full_head_insample_tau'] for c in cs for r in S['per_cohort'][c]['runs']],[t for c in cs for t in CFP[c]['taus']],[r['tau_lodo'] for c in cs for r in S['per_cohort'][c]['runs']]]
        for j,v in enumerate(vv):ax.bar(j,np.mean(v),color=COL[j],alpha=.3);ax.scatter(j+np.linspace(-.2,.2,len(v)),v,color=COL[j],marker=['D','o','s'][j],s=25)
        ax.set(xticks=range(3),xticklabels=['FULL head\nin-sample','FULL embedding\ncontaminated probe','NOALL embedding\nlabel-free probe'],title=label);ax.axhline(0,color='gray',lw=.7);ax.tick_params(axis='x',labelsize=8)
    axs[0].set_ylabel('Tau-b (descriptive; no inference)');fig.tight_layout();save(fig,'Fig3_Dissociation')
    fig,axs=plt.subplots(1,2,figsize=(9.6,3.8),sharey=True)
    for ax,cc,label in zip(axs,[cohorts[:3],cohorts[3:]],['A  From 90-section cohorts','B  From 35-section cohorts']):
        for j,m in enumerate(['SpaGCN','STAGATE','GraphST']):ax.plot(np.arange(3)+(j-1)*.07,[tests[c+'|'+m]['tau'] for c in cc],marker=['o','s','^'][j],color=COL[j],ls='none',label=m+' scalar')
        ax.errorbar(np.arange(3)+.2,[OSP[c]['mean_tau'] for c in cc],yerr=[np.std(OSP[c]['taus'],ddof=1) for c in cc],fmt='D',color='black',capsize=3,label='Reference ridge (3 seeds)');ax.set(xticks=range(3),xticklabels=['L1','L3','L4'],title=label);ax.axhline(0,color='gray',lw=.7)
    axs[0].set_ylabel('Tau-b (different readout procedures)');axs[0].legend(fontsize=7);fig.tight_layout();save(fig,'Fig4_Methods_benchmark')
    ldi=json.loads((ROOT/'results/csbj_ldi_validated.json').read_text(encoding='utf-8'));fig,axs=plt.subplots(1,2,figsize=(9,3.5))
    inj=ldi['injected_calibration'];axs[0].plot(inj['alphas'],inj['tau_curve'],'o-',ms=3);axs[0].set(xlabel='Rank-mixture weight (illustrative)',ylabel='Tau-b',title='A  Constructed rank association')
    vals=[np.mean([r['tau'] for r in rr]) for rr in REG.values()];errs=[np.std([r['tau'] for r in rr],ddof=1) for rr in REG.values()]
    axs[1].errorbar(range(7),vals,yerr=errs,fmt='o',capsize=3);axs[1].set(xticks=range(7),xticklabels=list(REG),ylabel='Mean score tau-b ± seed SD',title='B  Descriptive ablation association');axs[1].tick_params(axis='x',rotation=35,labelsize=8);axs[1].axhline(0,color='gray',lw=.7)
    fig.tight_layout();save(fig,'Fig8_LDI_calibration')
    local=json.loads((ROOT/'results/csbj_local_controls.json').read_text(encoding='utf-8'));fig,ax=plt.subplots(figsize=(8.5,3.8))
    probes=local.get('lodo_contamination',local.get('probe_contamination',{}))
    if not probes:
        for key,val in local.items():
            if isinstance(val,dict) and 'FULL' in val and isinstance(val['FULL'],dict) and 'tau_lodo_mean' in val['FULL']:probes=val
    names=list(probes)
    for j,v in enumerate(names):
        ts=[r['tau_lodo'] for r in probes[v]['per_seed']];ax.scatter(j+np.linspace(-.1,.1,len(ts)),ts,s=35,color=COL[0] if v=='NOALL' else 'gray')
    ax.set(xticks=range(len(names)),xticklabels=names,ylabel='Diagnostic leave-one-section-out tau-b',title='Contamination channels differ across variants; p-values omitted');ax.axhline(0,color='gray',lw=.7);fig.tight_layout();save(fig,'Fig9_Local_controls')
    # Replace workflow and graphical abstract misleading error-rate statements.
    for name,graphical in [('Fig1_Audit_workflow',False),('Graphical_Abstract',True)]:
        fig,ax=plt.subplots(figsize=(9,3.2 if graphical else 4.2));ax.axis('off')
        ax.text(.5,.91,'MRP-ST: Match the claim to its validation experiment',ha='center',fontsize=14,weight='bold',transform=ax.transAxes)
        boxes=[('Remove supervision','Declare label-derived losses\nCompare ablation scores'),('Probe label-free embeddings','Refit probes under permutation\nDeclare transductive boundary'),('Check external endpoints','Separate direction and organization\nfrom phenotype discrimination')]
        for j,(h,b) in enumerate(boxes):
            x=.015+j*.335;ax.add_patch(g['FancyBboxPatch']((x,.3),.3,.42,boxstyle='round,pad=.01',transform=ax.transAxes,fc='#EEF5F7',ec=COL[0]));ax.text(x+.15,.61,h,ha='center',fontsize=9.2,weight='bold',transform=ax.transAxes);ax.text(x+.15,.44,b,ha='center',fontsize=8.5,transform=ax.transAxes)
        ax.text(.5,.13,'Sections are the statistical units. Simulation-grid results are descriptive.',ha='center',fontsize=10,transform=ax.transAxes);save(fig,name)
    caps=[
      'Figure 1. Proposed MRP-ST workflow. Supervision-removal contrasts, refitted label-free probes and external endpoints answer distinct questions. The cervical analysis pools sections and remains transductive in covariates. No clinical validation or calibrated protocol error rate is inferred.',
      'Figure 2. Descriptive simulation grid. A: mean leave-one-out probe correlation and SD across nine runs per cell, nested within three generated cohorts. B: observed L1 and CTRL correlation-threshold crossing fractions. These are exploratory counts; archived grid p-values did not refit the probe and are excluded. No independent-run binomial intervals are assigned.',
      'Figure 3. Identity-confounded CTRL diagnostics. FULL heads fit in-sample; probes of FULL embeddings are contaminated by held-out-label exposure during encoder training; NOALL probes are label-free but transductive in covariates. All dots are descriptive optimizer runs and bars are means. Stored p-values are omitted. No mathematical upper-bound or causal leakage decomposition is asserted.',
      'Figure 4. Different readouts on identical selected 21-section subsets from one generator seed. Published methods use direct cluster-max pathway expression scores; the reference column uses leave-one-out ridge probes of pooled NOALL embeddings, with three-seed SD. No markers encode significance or comparative evidence grades. Fixed-score permutation results for the published-method readouts are provided in Table 4.',
      'Figure 5. Cervical section-level audit. A: score–severity correlations and optimizer-seed dispersion. NOALL scalar ranks are sensitive to aggregation precision. B: primary NOALL probe correlations and correctly refitted one-sided permutation p-values (10,000 permutations). Seeds reuse the same 35 sections; only seed 1 passes the 21-test BH correction. These are conditional transductive analyses.',
      'Figure 6. Expression-program orientation. Correlations reuse expression input and are descriptive. Raw NOALL seed orientations are shown without label-based sign reversal; they do not establish independent pathway mechanism.',
      'Figure 7. External endpoints. A: FINAL-checkpoint patient/tissue means in six paired GeoMx patients; all tumor-minus-benign differences are negative (exact two-sided signed-rank p=0.03125). B: spatial organization in two head-and-neck cohorts and separate phenotype tests. Organization does not establish discrimination or probability calibration.',
      'Figure 8. Descriptive association diagnostics, not calibration of causal label dependence. A: constructed mixture of random and label ranks illustrates association as mixture weight changes. B: stored ablation score–severity means with optimizer-seed SD. NOALL is numerically sensitive. Z_score is a fixed-score standardized association summary and does not supply a valid test when the encoder was trained on the same labels.',
      'Figure 9. Diagnostic probe correlations by cervical variant. Stored probes exclude each section from ridge fitting but not from encoder training. FULL and other target-informed variants retain different supervision channels; NOSUP removes ordinal and proxy losses but retains alignment. No archived diagnostic p-values are interpreted, because their permutations did not refit predictions. The primary, correctly refitted NOALL result appears in Figure 5.'
    ]
    g['CAPTIONS']=[(name,cap) for (name,_),cap in zip(g['CAPTIONS'],caps)]
    # Initials plus family names follow the journal's reference instructions.
    metadata=json.loads((QA/'reference_metadata_audit.json').read_text(encoding='utf-8'))
    def reftext(k):
        r=g['REF'][k];a=metadata[k].get('metadata',{}).get('author',[])
        if a:
            authors=[]
            for x in a:
                if x.get('given'):authors.append(' '.join(t[0]+'.' for t in re.split(r'[\s-]+',x['given']) if t)+' '+x.get('family',''))
                else:authors.append(x.get('name',x.get('family','')))
            names='; '.join(authors)
        else:names='; '.join(r['authors'])
        return f"{names}. {r['title']}. {r['journal']}. {r['year']};{r.get('volume','')}:{r.get('page','')}. {r['url']}"
    g['reftext']=reftext
    # BibTeX should include complete volume/pages as well as correct author metadata.
    print('BUILD',g['build_main'](),flush=True)
    (OUT/'References.bib').write_text('\n\n'.join('@article{'+k+',\n'+',\n'.join('  '+field+' = {'+str(value)+'}' for field,value in [('author',' and '.join(r['authors'])),('title',r['title']),('journal',r['journal']),('year',r['year']),('volume',r.get('volume','')),('pages',r.get('page','')),('doi',r.get('doi','')),('url',r['url'])])+'\n}' for k,r in g['REF'].items()),encoding='utf-8')
    supplementary(g,title)
    return g

def supplementary(g,title):
    from docx.shared import Pt
    d=g['setup']();d.add_paragraph('Supplementary Materials',style='Title');d.add_paragraph(title);d.add_paragraph('Mingwei Bai, Zhengmin Gu, Ye Zhang, Zhiping Ni and Lang Guo')
    sections=[
     ('S1. Evidence families and validation boundaries','The earlier three-scenario design contains strong-shared (A), weak-shared (B) and identity-confounded (C) cohorts at 90 and 35 sections, plus the cervical cohort. These seven cohorts, each with three optimizer seeds, form the 21-test primary refitted permutation family. They are distinct from the later L1–L4/CTRL grid, whose p-values shuffled labels against unchanged predictions. Table S1 supplies the valid primary family. The dose–response grid is descriptive, and its complete per-run embeddings were not available in the reviewed local snapshot.'),
     ('S2. Reference architecture and objectives','The architecture, losses, recorded weights, sample-local batching, preprocessing and seven ablations are specified in Section 2.2 and Table 1. The real-data launcher specifies 40 epochs; learning rate 0.001 and batch size 256 are current inherited defaults, not recovered training-time settings. No historical source commit is asserted. Explicit labels and expression-derived proxy targets must be distinguished from pathway alignment. NOSUP excludes ordinal and proxy objectives; NOALL additionally excludes alignment.'),
     ('S3. Simulation designs','The earlier scenario A has shared signal amplitude 2.5 and batch amplitude 0.3; B has signal 0.20, batch 1.5, nonresponder fraction 0.35 and program-gene presence probability 0.5; C has donor fingerprint amplitude 1.2 and batch amplitude 0.5, with no injected shared disease signal. The earlier generator seed is 20261006. The later grid varies nonresponder fraction 0,0.25,0.50,0.75 at fixed amplitude 0.10 and program-gene presence probability 0.80; generator seeds are 20261006–20261008. Each design reuses each generated cohort across three optimizer seeds. Generated cohorts have 1,500 spots per sample and 3,000 genes; stored fitted simulation runs retain 30,000 spots after the global cap. Placeholder images do not test realistic morphology.'),
     ('S4. Refitted label permutations','For each held-out section, calculate training-only standardization, ridge penalty 1 and a leave-one-out prediction operator H with zero diagonal. Predictions for each complete label permutation are H y_perm. The encoder and covariates remain fixed; this is appropriate only for a label-independent representation under exchangeability. Use 10,000 permutations, seed 20261010, and (count+1)/(B+1). The intercept sensitivity is separately reported. The family contains 21 tests; BH q-values address that family but correlated seeds are not independent biological replications.'),
     ('S5. Numerical precision and association summaries','NOALL scalar means are sensitive to float32 versus float64 aggregation. Historic mean tau −0.105 differs from recomputed values near −0.092. These are not biological effect estimates with validated uncertainty. All 21 archived primary tests reproduce when their original array dtype is retained. Float64 standardization and operator construction give cervical taus 0.215,0.397,0.200 with refitted p=0.1041,0.0027,0.1126; the one-significant-seed conclusion is unchanged. Z_score replaces the previous misleading LDI name: it standardizes mean FULL tau against a fixed FULL-seed-0 permutation reference and is descriptive only. Supervised encoder retraining under permuted labels would be required for a corresponding model-training null. The binary-label separation formula assumes untied score predictions; within-class ties can raise tau-b to 1.'),
     ('S6. Published-method readouts','SpaGCN, STAGATE and GraphST run separately within sections. Their raw-count pathway means use the 40 genes in docs/hpv_genes.json (or a 40-gene fallback), not the reference model’s nominal 108-gene panel. K-means k depends on n_spots rather than section count. The archived function named lodo_from_sample_risks computes direct fixed-score tau, not leave-one-out prediction. This revision recalculates fixed-score permutation p-values with 10,000 draws and seed 20261014; results are unadjusted exploratory values. Reference subset p-values remain unverified and are excluded. No method receives a C4 multi-seed grade from a single published-method run.'),
     ('S7. Separate cervical AUC tasks','Section-level lesion AUC (FULL 1.00; raw panel mean 0.75) is in-cohort discrimination of four lesions and 31 normal sections. The morphology baseline predicts expression and measures spot-level proxy AUC within each left-out section; its proxy uses that section’s median panel expression. Ridge mean fold AUC 0.4967 uses 16 evaluable folds and MLP 0.497 uses 33; excluded folds and fold-specific thresholds limit interpretation. These endpoints do not estimate a FULL-model disease AUC drop. The same reconstruction baseline has mean Pearson −0.0145 across 35 sections.'),
     ('S8. External endpoints','GSE316098 GeoMx FINAL scores are averaged within patient and tissue type before paired inference. Six of seven patients provide both tissues; corrected patient/core annotations are used. Expression-only input matches 4,788 of 4,991 genes with zero filling. Silva core statistics ignore within-patient dependence and remain descriptive. HNSCC spatial organization and HPV-status discrimination are separate tests; an 11-gene TCGA-CESC screen is contextual and not model-level prognosis.'),
     ('S9. Reproduction and availability','The existing repository release v1.0 is public. The checked revision must be deposited with an immutable identifier before submission so the released source and evidence correspond to the actual manuscript. The reviewed package includes pooled arrays for the 21-test primary family, result extracts, a refitted-probe verifier, corrected reference metadata and checksums. Raw datasets and foundation-model weights are not bundled. Complete later-grid embeddings are unavailable locally; this prevents corrected grid-wide significance, rather than just full GPU retraining.')]
    for h,p in sections:d.add_heading(h,1);d.add_paragraph(p)
    rows=[]
    for cohort,rr in g['P']['groups'].items():
        for r in rr:
            a=r['legacy_no_intercept'];b=r['intercept_sensitivity'];rows.append([cohort,r['seed'],r['n_samples'],f"{a['tau']:.4f}",f"{a['p_refit']:.4f}",f"{a['q_BH_21']:.4f}",f"{b['tau']:.4f}",f"{b['p_refit']:.4f}"])
    g['table'](d,['Cohort','Seed','n','τ','p refit','q BH','τ + int.','p + int.'],rows,[1900,450,450,1150,1350,1350,1360,1350],'Table S1. Primary refitted probe family. A/B/C are the earlier scenarios, not the L1–L4 grid.')
    mapping=[['A / A_small','Earlier strong shared program','Primary 21-test probe family'],['B / B_small','Earlier weak shared program','Primary 21-test probe family'],['C / C_small','Earlier identity-confounded design','Primary 21-test probe family'],['cervical','35-section demonstration','Main Figure 5 / Section 3.4'],['L1–L4 / CTRL','Later three-generator dose grid','Descriptive Figure 2 / Table 3'],['method subset','21 sections from generator 20261006','Different readouts; Table 4 / Figure 4']]
    g['table'](d,['Evidence identifier','Design','Manuscript mapping'],mapping,[2200,3000,4160],'Table S2. Evidence-to-text mapping.')
    rows=[[r['patient'],f"{r['benign_mean']:.4f}",f"{r['tumor_mean']:.4f}",f"{r['difference']:.4f}"] for r in g['G']['pairs']]
    g['table'](d,['Patient','Benign','Tumor','Difference'],rows,[2100,2420,2420,2420],'Table S3. GeoMx patient-level pairs.')
    rows=[[a,role,'https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc='+a] for a,role in [('GSE208654','Cervical lesions / normal'),('GSE217237','Normal ectocervix'),('GSE316098','Cervical GeoMx'),('GSE268014','HNSCC spatial cohort'),('GSE281978','HNSCC spatial cohort'),('GSE308792','Single-cell context')]]
    g['table'](d,['Accession','Role','Source'],rows,[1500,2800,5060],'Table S4. GEO accessions. A public 10x cervical section is additionally included in the source cohort.')
    d.save(OUT/'Supplementary_Materials_CSBJ.docx')
    d=g['setup']();d.add_paragraph(title,style='Title')
    md=(OUT/'Manuscript_CSBJ.md').read_text(encoding='utf-8').split('## Abstract')[0]
    for p in md.split('\n\n')[1:]:d.add_paragraph(p.strip())
    d.save(OUT/'Title_Page_CSBJ.docx')
    letter='''Dear Editors,

Please consider our manuscript, "'''+title+'''", as a Research Article for the General section of Computational and Structural Biotechnology Journal.

MRP-ST proposes five experiments for interpreting spatial transcriptomic scores: supervision removal, declaration of statistical units, score-association diagnostics, refitted probing of label-free representations and endpoint-specific external analyses. In 35 cervical sections, removing ordinal supervision changes mean score tau-b from 0.460 to 0.251; a label-free transductive representation retains seed-unstable association (mean probe tau-b 0.270; one of three seeds significant). Controlled simulations illustrate contamination when held-out probe labels entered encoder training. External paired GeoMx scores reverse direction relative to source-cohort severity.

The study does not establish calibrated protocol error rates, comparable rankings of published methods or patient-level diagnostic performance. The simulation-grid correlations are descriptive because its archived label permutations did not refit the probe. Section-level lesion AUC and spot-level expression-proxy AUC are reported as separate tasks. The supplemental package supplies correctly refitted primary probe tests, evidence mapping, source scripts and numerical limitations.

The original repository release v1.0 is public at https://github.com/WorldBMW/pathway-anchored-risk-benchmark/releases/tag/v1.0. This corrected revision requires a matching versioned deposit before submission.

OpenAI Codex assisted code generation, statistical checking, drafting, editing and figure preparation. Authors retain responsibility for reviewing the final content. All plots were generated from numerical results; no AI-generated experimental images are presented.

Author approval, exclusive-submission status, related manuscripts, funding and competing-interest declarations must be confirmed by the author team before this letter is submitted. Complete author ORCIDs remain to be supplied. Fourth author Zhiping Ni contributed data curation, resources and manuscript review and editing. No approval or declaration is made on any author's behalf.

Corresponding author: Lang Guo, School of Pharmacy, Shenyang Medical College, 146 Huanghe North Avenue, Shenyang 110034, China; guolang1016@symc.edu.cn; +86 13555707001.

Authors: Mingwei Bai (baimingwei@jm360.cn); Zhengmin Gu (guzm@cmu1h.com); Ye Zhang (13332457636@163.com); Zhiping Ni (nizhiping@jm360.cn); Lang Guo (guolang1016@symc.edu.cn).

Sincerely,
Lang Guo, on behalf of the author team
'''
    (OUT/'Cover_Letter_CSBJ.md').write_text(letter,encoding='utf-8');d=g['setup']();d.add_paragraph('Cover letter - author confirmation required',style='Title')
    for p in letter.split('\n\n'):d.add_paragraph(p)
    d.save(OUT/'Cover_Letter_CSBJ.docx')

if __name__=='__main__':build()
