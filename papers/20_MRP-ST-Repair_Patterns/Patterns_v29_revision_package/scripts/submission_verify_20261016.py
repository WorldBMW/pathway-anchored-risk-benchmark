from pathlib import Path
import sys,json,re,zipfile,hashlib,shutil,io
R=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(R/'.qa/submission_deps'))
import fitz
import numpy as np
from scipy.stats import kendalltau
from docx import Document
from PIL import Image,ImageDraw
OUT=R/'papers/投稿复核_20261016'; QA=R/'.qa/submission_reaudit_20261016'

def numeric_audit():
    ev=R/'papers/14_BiB_20261016/evidence';saved=json.loads((ev/'corrected_probe.json').read_text())
    records=[]
    for group,rows in saved['groups'].items():
        stem='pooled_cervical' if group=='cervical' else 'pooled_'+group
        for row in rows:
            fp=ev/(stem+'_seed'+str(row['seed'])+'.npz')
            z=np.load(fp,allow_pickle=True);X=z['X'];y=z['y'].astype(float);n=len(y);p=X.shape[1];H=np.zeros((n,n))
            for i in range(n):
                tr=np.arange(n)!=i;mu=X[tr].mean(0);sd=X[tr].std(0)+1e-8
                xt=(X[tr]-mu)/sd; xv=(X[i]-mu)/sd
                # Dual ridge is equivalent and avoids inverting a 128-by-128 matrix.
                H[i,tr]=np.linalg.solve(xt@xt.T+np.eye(n-1),xt@xv)
            pred=H@y
            tau=float(kendalltau(pred,y).statistic);target=row['legacy_no_intercept']['tau']
            err=float(np.max(np.abs(pred-np.array(row['legacy_no_intercept']['predictions']))))
            rng=np.random.default_rng(20261010);ii,jj=np.triu_indices(n,1);ge=0
            for start in range(0,10000,100):
                Y=np.stack([rng.permutation(y) for _ in range(100)]);P=Y@H.T;a=np.sign(P[:,jj]-P[:,ii]);b=np.sign(Y[:,jj]-Y[:,ii]);den=np.sqrt((a*a).sum(1)*(b*b).sum(1));tv=np.divide((a*b).sum(1),den,out=np.zeros(len(P)),where=den>0);ge+=int((tv>=tau-1e-12).sum())
            pv=(ge+1)/10001;savedp=row['legacy_no_intercept']['p_refit']
            records.append({'group':group,'seed':row['seed'],'n':n,'dtype':str(X.dtype),'tau_recomputed':tau,'tau_saved':target,'p_refit_recomputed':pv,'p_refit_saved':savedp,'p_match':abs(pv-savedp)<1e-10,'max_prediction_error':err,'tau_match':abs(tau-target)<1e-10})
    baseline=json.loads((R/'results/baseline_hpv_auc.json').read_text())['per_regressor'][0]
    a=np.array([x['auc'] for x in baseline['per_fold'] if x['auc'] is not None])
    report={'ridge_probes':records,'all_tau_match':all(r['tau_match'] for r in records),'all_p_match':all(r['p_match'] for r in records),'proxy_baseline':{'valid_folds':len(a),'attempted_folds':len(baseline['per_fold']),'mean_recomputed':float(a.mean()),'sd_recomputed':float(a.std()),'stored_mean':baseline['mean_auc'],'target':'within-section median-thresholded expression proxy, not disease classification'},'scope':'Independent recomputation of LOO predictions, tau and 10,000 refitted permutations for all 21 primary saved pooled representations, preserving stored float32 feature arithmetic and adding 1e-8 to training-fold standard deviations. Proxy AUC summary also recomputed. Model training not rerun. Float64 sensitivity changes two cervical seed taus slightly in opposite directions (0.21463 vs 0.20736; 0.39652 vs 0.40380), leaving their mean unchanged.'}
    (QA/'numeric_audit.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print('Numerical audit:',report['all_tau_match'],report['proxy_baseline'],flush=True)

def anonymized_evidence():
    src=OUT/'Audit_evidence.zip';dest=OUT/'Interdisciplinary Sciences Computational Life Sciences/Audit_evidence_blinded.zip'
    patterns=[r'https?://github\.com/WorldBMW/[^\s"\)]+',r'WorldBMW',r'Mingwei Bai',r'Zhengmin Gu',r'Ye Zhang',r'Zhiping Ni',r'Lang Guo',r'JiMei Technology Group Co\., Ltd\.',r'guolang1016@symc\.edu\.cn',r'baimingwei@jm360\.cn',r'guzm@cmu1h\.com',r'nizhiping@jm360\.cn',r'13332457636@163\.com',r'[A-Z]:[/\\][^\s"\n\r]*']
    manifest=[]
    with zipfile.ZipFile(src) as a,zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as b:
        for name in a.namelist():
            if name=='SHA256SUMS.txt':continue
            raw=a.read(name)
            if name.endswith(('.json','.py','.txt','.csv')):
                t=raw.decode('utf-8-sig')
                for p in patterns:t=re.sub(p,'[withheld for double-blind review]',t,flags=re.I)
                raw=t.encode('utf-8')
            b.writestr(name,raw);manifest.append(hashlib.sha256(raw).hexdigest()+'  '+name)
        b.writestr('SHA256SUMS.txt','\n'.join(manifest)+'\n')

def inspect_docs():
    report=[]
    for path in OUT.rglob('*.docx'):
        d=Document(path);t='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for tab in d.tables for row in tab.rows for c in row.cells)
        xml=zipfile.ZipFile(path).read('word/document.xml').decode()
        markers=re.findall(r'\[\[.*?\]\]|\{\{FIG:|\[@[\w;@]+\]|\$\$',t)
        identities=[x for x in ['Mingwei Bai','Zhengmin Gu','Ye Zhang','Zhiping Ni','Lang Guo','WorldBMW','jm360.cn','symc.edu.cn'] if x.lower() in t.lower()]
        report.append({'file':str(path.relative_to(OUT)),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'images':len(d.inline_shapes),'equations':xml.count('<m:oMath>'),'unresolved_markers':markers,'identities':identities if 'blinded' in path.name else [],'tracked_changes':bool(re.search(r'<w:(?:ins|del)\b',xml)),'comments':bool(re.search(r'<w:commentRangeStart\b',xml))})
    assert all(not r['unresolved_markers'] and not r['identities'] and not r['tracked_changes'] and not r['comments'] for r in report)
    for r in report:
        if r['file'].split('\\')[-1].startswith('Manuscript_'):assert r['tables']==3 and r['images']==4 and r['equations']==5,r
    (QA/'document_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Document checks:',len(report),'passed',flush=True)

def rasterize():
    pages=[]; allimgs=[]
    for pdf in (QA/'render').rglob('*.pdf'):
        doc=fitz.open(pdf);folder=pdf.parent/pdf.stem;folder.mkdir(exist_ok=True)
        text='\n'.join(p.get_text() for p in doc)
        if pdf.name.startswith('Supplementary_GPB'):shutil.copy2(pdf,OUT/pdf.parent.name/pdf.name)
        for i,p in enumerate(doc):
            target=folder/f'page-{i+1}.png';p.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(target)
            blocks=[b for b in p.get_text('blocks') if b[4].strip()]
            outside=[b[4][:100] for b in blocks if b[0]<25 or b[1]<25 or b[2]>p.rect.width-25 or b[3]>p.rect.height-25]
            pages.append({'file':str(pdf.relative_to(QA/'render')),'page':i+1,'words':len(p.get_text('words')),'blocks_outside_25pt_margin':outside})
            allimgs.append((pdf.parent.name+'/'+pdf.stem+f' p{i+1}',target))
        print('Raster:',pdf.parent.name,pdf.stem,len(doc),flush=True)
    sheets=QA/'contacts';sheets.mkdir(exist_ok=True)
    for start in range(0,len(allimgs),9):
        sheet=Image.new('RGB',(1836,2475),'#dddddd');draw=ImageDraw.Draw(sheet)
        for j,(label,p) in enumerate(allimgs[start:start+9]):
            im=Image.open(p);im.thumbnail((600,785));x=j%3*612;y=j//3*825;sheet.paste(im,(x,y+38));draw.text((x+6,y+4),label[:75],fill='black')
        sheet.save(sheets/f'contact-{start//9+1:02}.jpg',quality=90)
    (QA/'page_audit.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Pages',len(pages),'sheets',(len(allimgs)+8)//9,flush=True)

if __name__=='__main__':
    if '--numeric' in sys.argv:numeric_audit()
    elif '--raster' in sys.argv:rasterize()
    else:anonymized_evidence();inspect_docs()
