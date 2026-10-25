from pathlib import Path
import json,re,shutil,hashlib,zipfile,sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];SRC=ROOT/'papers/13_CSBJ_v2_20261014';OUT=SRC/'CSBJ_submission_checked_20261014';QA=SRC/'submission_review'
sys.path.insert(0,str(ROOT/'scripts'))
from csbj_review_build import module
g=module()

def update_author():
    old='Lang Guo: conceptualization, methodology, supervision, project administration and writing—review and editing. Zhiping Ni: contribution statement requires author confirmation before submission.'
    new='Zhiping Ni: data curation, resources and writing—review and editing. Lang Guo: conceptualization, methodology, supervision, project administration and writing—review and editing.'
    old_cover='Complete author ORCIDs and the contribution statement for fourth author Zhiping Ni remain to be supplied.'
    new_cover='Complete author ORCIDs remain to be supplied. Fourth author Zhiping Ni contributed data curation, resources and manuscript review and editing.'
    wording={
      'One descriptive contrast and one calibrated index are defined.':'A descriptive ablation contrast and a standardized association summary are reported.',
      "Synthetic cohorts provide the known-signal positive controls on which the protocol's error rates are measured.":'Synthetic cohorts provide known-signal positive controls for describing probe behavior.',
      '2.6. Unified audit of published embedding methods':'2.6. Illustrative published-method readouts',
      'rather than this statistic alone (Figure 8).':'rather than this statistic alone.',
      'Figures 8 and 9 show numerical association references and contamination controls, without interpreting either as a calibrated label-dependence test.':'Numerical association references and contamination controls are descriptive, without interpreting either as a calibrated label-dependence test.',
      '(Figure 9). Their archived p-values':'Their archived p-values',
      'variants Their archived p-values':'variants. Their archived p-values',
      'The NOALL reading is reported as an interval because that head collapses to a near-constant offset and its rank ordering is not identifiable (§2.3); it is numerically sensitive and is not treated as evidence of a null biological association.':'The NOALL scalar ordering is numerically sensitive (§2.3); its historical summary is not treated as a stable effect estimate or evidence of a null biological association.',
      'Markus Ojala; Gemma C. Garriga.':'M. Ojala; G. C. Garriga.',
    }
    diagnostic_sentence='Figures 8 and 9 summarize descriptive association diagnostics and contamination contrasts. '
    def clean_text(text):
        text=re.sub('(?:'+re.escape(diagnostic_sentence)+')+',diagnostic_sentence,text)
        tail='Prospective validation of the protocol and any clinical application remains future work.'
        if tail in text and diagnostic_sentence not in text:text=text.replace(tail,diagnostic_sentence+tail)
        return text
    for name in ['Manuscript_CSBJ.md','Manuscript_CSBJ_resolved.md','Cover_Letter_CSBJ.md','References_verified.txt']:
        p=OUT/name;s=p.read_text(encoding='utf-8').replace(old,new).replace(old_cover,new_cover)
        for a,b in wording.items():s=s.replace(a,b)
        s=clean_text(s)
        p.write_text(s,encoding='utf-8')
    for name in ['Manuscript_CSBJ.docx','Cover_Letter_CSBJ.docx']:
        doc_path=OUT/name;d=Document(doc_path)
        for para in d.paragraphs:
            if old in para.text or old_cover in para.text:para.text=para.text.replace(old,new).replace(old_cover,new_cover)
            for a,b in wording.items():
                if a in para.text:para.text=para.text.replace(a,b)
            cleaned=clean_text(para.text)
            if cleaned!=para.text:para.text=cleaned
            if re.match(r'^\[\d+\] ',para.text):
                for run in para.runs:run.font.size=Pt(9)
        if name=='Manuscript_CSBJ.docx':
            # Keep two figures with their legends per page; eliminate orphan legend heading.
            pics=[p for p in d.paragraphs if p._p.xpath('.//w:drawing')]
            for p in list(d.paragraphs):
                if p._p.xpath('.//w:br[@w:type="page"]') and not p.text.strip():p._p.getparent().remove(p._p)
                if p.text=='Figure legends':p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
            for i,p in enumerate(pics):p.paragraph_format.page_break_before=(i>0 and i%2==0)
            for p in d.paragraphs:
                if p.style.name=='Title':
                    for border in p._p.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
            for border in d.styles['Title'].element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
            for table in d.tables:
                for para in table.rows[0].cells[0].paragraphs:para.paragraph_format.keep_with_next=True
                for cell in table.rows[0].cells:
                    for para in cell.paragraphs:para.paragraph_format.keep_with_next=True
        d.save(doc_path)
    sup=OUT/'Supplementary_Materials_CSBJ.docx';d=Document(sup)
    d.tables[0].cell(0,1).text='k'
    d.save(sup)

def report():
    d=g['setup']()
    for n in ['Normal','Title','Heading 1','Heading 2','Caption']:
        rf=d.styles[n].element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),'宋体' if n in ['Normal','Caption'] else '黑体')
    d.add_paragraph('CSBJ 稿件逐项检查与修订报告',style='Title')
    d.add_paragraph('检查日期：2026年10月14日；对象：papers/13_CSBJ_v2_20261014。原始稿、结果和已发布 v1.0 均保留；修订材料存放在 CSBJ_submission_checked_20261014。')
    sections=[
      ('审查结论','原稿存在实质性统计与解释错误，不能仅做英文润色后直接投稿。修订稿已撤下无有效证据支持的推断，保留能复现的主分析，并重新生成配图、参考文献和补充材料。当前版本可供作者终审；尚不能宣称已完成所有投稿条件。作者 ORCID、全体作者批准、独家投稿与相关稿件状态尚未收到确认；修订版需对应新的不可变归档标识。若保留原稿“校准误差率”的核心卖点，必须补齐完整模拟网格嵌入并用真正重拟合置换复算，而不能沿用既有显著率。'),
      ('作者与单位','作者顺序已按用户确认改为 Mingwei Bai¹、Zhengmin Gu²、Ye Zhang²、Zhiping Ni¹、Lang Guo³*。倪志平为第四作者，邮箱 nizhiping@jm360.cn，单位 JiMei Technology Group Co., Ltd., Shenyang 110002, China。其贡献与第三作者相同：data curation、resources、writing—review and editing。Lang Guo 保留为通讯作者，原姓名、单位、地址、电话和邮箱未擅改。截图中的姓名拼写与现有项目记录一致，但不能通过项目文件独立确认个人法定拼写、单位邮编和联系信息；须以作者本人确认信息为准。两位公司作者的公司归属已同时披露。'),
      ('方法：置换检验的实现与文字冲突','scripts/probe_lodo.py 的 lodo_probe 先拟合原标签预测，然后只打乱评价标签，没有在每次置换下重新拟合预测。scripts/sim_power_grid.py 调用此函数，故网格的显著率及据此判级不等于正文所称的重拟合检验。scripts/csbj_local_controls.py 的诊断 p 值也存在这一问题。修订稿删除这两类 p 值的推断用途，并删除网格基于显著率的支持等级。已提供独立复算代码。完整后来网格嵌入在本地与已发布 tag 中均不完整，不能伪造补算后的误差率。'),
      ('方法：LDI 的解释错误','原 LDI=(tau_FULL−固定评分置换均值)/置换SD，仅是固定评分关联的标准化量。训练标签已进入 FULL 编码器，此时打乱评价标签不能检验训练时标签依赖，更不能证明泛化。跨种子均值与单种子置换参考混用也不是严格校准。已改称描述性 Z_score=3.26，移除因果判断门限；标签监督来源和消融差异才是监督依赖的直接证据。高斯1.645与经验tau分位点0.236也不作为完全等价判据。'),
      ('方法：重复运行和有限样本公式','九次运行是三个生成队列乘三个优化种子，不能按九次独立二项试验给“精确”误差率区间。已改为描述性计数，不再宣称名义5%误差率控制或总体灵敏度。二标签 tau-b 公式 sqrt(n0*n1/C(n,2)) 仅适用于预测无并列的条件；类内并列的预测可以达到tau-b=1。修订公式明确这一条件，删除通用标签上界与仅凭舍入值判断完美分离的说法。阈值被表述为探索性，未发现可验证的事前注册记录。'),
      ('内容：两类 AUC 不可直接比较','AUC=1.00 来自35切片中4病变/31正常的队列内疾病判别；AUC≈0.497 来自形态预测表达后，在留出切片内识别表达代理高/低点位，分别是不同模型、不同单位和不同标签任务。代理还采用每切片自身中位数，而参考训练使用池化中位数。已删除“同一模型从1.00坍缩到0.497”的摘要、结论和cover letter说法。正文保留两任务分别说明，不宣称现有FULL模型已有患者留出疾病AUC。'),
      ('内容：第三方方法并非统一探针比较','benchmark_st_methods.py 中名为 lodo_from_sample_risks 的函数实际直接计算预先构造的通路评分tau，不进行留一线性探针。SpaGCN/STAGATE/GraphST读出采用docs/hpv_genes.json的40基因原始计数簇均值最大值；参考模型名义面板为108基因，读出则为128维表征岭探针。k-means簇数依赖点位数而非切片数。修订明确算法、面板和读出差异；移除公平排名、统一C4支持等级和无标签发现主张。第三方固定评分置换可在交换性条件下使用，已独立重算18个方法×队列结果（10,000次；随机种子20261014），作为未校正的探索性p值报告。'),
      ('内容：负相关、相同tau和单token注意力','通路评分有固定表达方向，STAGATE负相关不能仅解释成无监督嵌入任意符号翻转；已改为原因未确定。不同水平相同舍入tau不意味着风险向量完全相同或读出饱和；已收窄解释。源码fusion仅有一个图像key/value token，softmax权重恒为1，不是多点位或多patch空间注意力；方法已如实描述为残差式融合。NOALL大偏置小差值说明数值敏感性，但仅凭比例不足以证明所有排序均为浮点噪声；措辞已修正。'),
      ('主结果复现与精度敏感性','21份早期三场景与宫颈的池化表征可用。保留原数组dtype，独立从每份表征重新计算10,000次完整标签置换，并对每次置换重新拟合岭探针；21/21相关系数和p值均与corrected_probe.json匹配。宫颈原taus=0.207355/0.403797/0.200079，p=0.113489/0.001900/0.112489。float64预处理改为0.214631/0.396521/0.200079，p=0.104090/0.002700/0.112589；均仍只有中间种子显著。已补精度敏感性，不把不同计算精度结果混在同一口径里。'),
      ('数据与外部分析','六个GEO编号已从NCBI实时核验均公开可访问：GSE208654、GSE217237、GSE316098、GSE268014、GSE281978、GSE308792。GSE217237为SuperSeries，正常对照来自多子系列及HIV、避孕和微生物组研究背景，不能视为同质健康对照；正文已补。GeoMx六例患者配对检验、HNSCC空间组织/HPV判别、TCGA11基因筛查分别陈述，避免将基因级结果视为完整模型外部预后验证。未重跑GPU训练、全部原始数据预处理或历史编码器训练；相关结论只在既有证据范围内。'),
      ('参考文献','正文22条引用中21条有DOI，均用Crossref官方记录核对；无DOI的Ojala与Garriga论文已查JMLR原刊。GraphST的16位作者名单在原稿中有系统性错误，已改为出版社记录，首作者Yahui Long、末作者Jinmiao Chen；DOME、SpaGCN等作者与团体作者信息也按官方记录修正。保留全部作者，按作者姓名首字母与姓氏规范输出，检查题名、期刊、年、卷、页码/文章号和DOI。文内引用按首次出现连续编号，无未解析引用键；BibTeX补齐卷页信息。'),
      ('格式、配图与投稿材料','修订摘要约250词内；五张主表按首次出现编号，九张主图均有图注与文内说明，图像从数值生成，提供PDF/SVG/PNG/600dpi TIFF。重新绘制模拟图、方法读出图、关联诊断图及图形摘要，去除无效显著性标记和错误门限标签；主图使用形状/颜色辅助区分。提供英文主稿、Title Page、Supplementary Materials和cover letter。现行CSBJ指南推荐Research Article不超过15,000词、10图和5表，本稿在词数和展示项推荐范围内。指南要求cover letter含作者ORCID、作者批准及独家投稿信息，并提交推荐审稿人姓名/单位/邮箱；这些个人声明不代填。'),
      ('归档与直接投稿条件','GitHub v1.0 release已确认存在（2026-10-14发布），实际tag树包含21份主分析NPZ，不含完整后续模拟网格。它承载原版而非本次修订稿，不能说本次修订已全部公开。已准备本地修订证据包、checksums与复算脚本；发布需采用匹配修订版的不可变版本标识。提交前还需：补作者ORCID或说明无；确认全部作者批准、独家投稿和相关稿件；确认资金/利益冲突/二次分析伦理及数据许可声明；提供无利益冲突的推荐审稿人信息。当前英文修订稿保留这些待核状态，不把它标称已可无条件直接提交。')]
    for h,p in sections:d.add_heading(h,1);d.add_paragraph(p)
    d.add_heading('核验来源与证据位置',1)
    for p in ['CSBJ现行作者指南：https://spj.science.org/page/csbj/for-authors','GraphST原刊：https://www.nature.com/articles/s41467-023-36796-3','JMLR置换检验原刊：https://www.jmlr.org/papers/v11/ojala10a.html','文献官方元数据：submission_review/reference_metadata_audit.json','NCBI实时来源核验：submission_review/geo_live_checks.json','21项主分析复现：submission_review/primary_probe_reproduction.json','float64敏感性：submission_review/primary_probe_float64_sensitivity.json','已发布版本树：submission_review/github_v1_tree.json']:
        d.add_paragraph(p)
    d.save(OUT/'CSBJ_逐项检查与修改报告.docx')

def package():
    E=OUT/'reproducibility';E.mkdir(exist_ok=True);(E/'evidence').mkdir(exist_ok=True);(E/'scripts').mkdir(exist_ok=True)
    for f in (SRC/'evidence').iterdir():
        if f.is_file():shutil.copy2(f,E/'evidence'/f.name)
    for f in QA.iterdir():
        if f.is_file() and f.suffix in ['.json','.txt']:shutil.copy2(f,E/'evidence'/f.name)
    for f in ['csbj_review_verify.py','csbj_review_build.py','csbj_review_finish.py','csbj_build_package_v2.py','csbj_submission_review_20261014.py','benchmark_st_methods.py','probe_lodo.py','sim_power_grid.py','train_st.py','simulate_st_positive_control.py','baseline_hpv_auc_eval.py','csbj_local_controls.py']:
        shutil.copy2(ROOT/'scripts'/f,E/'scripts'/f)
    (E/'docs').mkdir(exist_ok=True);shutil.copy2(ROOT/'docs/hpv_genes.json',E/'docs/hpv_genes.json')
    (E/'src').mkdir(exist_ok=True)
    for f in (ROOT/'src').glob('*.py'):shutil.copy2(f,E/'src'/f.name)
    for f in ['third_party_fixed_score_checks.json','References_verified.txt']:
        shutil.copy2(OUT/f,E/'evidence'/f)
    (E/'requirements.txt').write_text('numpy\nscipy\npython-docx\nmatplotlib\npillow\n',encoding='utf-8')
    (E/'README.md').write_text('''# Checked CSBJ revision, 2026-10-14

This is a locally prepared revision, not a newly published release. The manuscript is an author-review draft pending declarations, ORCIDs and matching deposition. No DOI has been invented.

The 21 pooled primary datasets reproduce the archived corrected_probe results with their original dtype. The primary_probe_reproduction report records this check; primary_probe_float64_sensitivity records the alternate float64 preprocessing. The later dose-response grid lacks complete per-run embeddings in this snapshot. Its stored p-values are invalid for refitted-probe inference and must not be reused as calibrated error rates.

The included original scripts/probe_lodo.py and sim_power_grid.py are retained solely as source provenance and contain the documented fixed-prediction permutation issue. Use scripts/csbj_review_verify.py's operator/audit functions for correct algebraic refitting. For a portable primary verification, call audit(z['X'], z['y']) on each evidence/pooled_*_seed*.npz and compare with corrected_probe.json; do not cast X unless performing the stated precision sensitivity. Permutations: 10000, seed 20261010, ridge penalty 1, train-fold scaling, no intercept. Third-party score checks use fixed label-independent risks, seed 20261014, 10000 permutations and remain exploratory unadjusted results.

Original repository v1.0: https://github.com/WorldBMW/pathway-anchored-risk-benchmark/releases/tag/v1.0 . That tag does not contain complete grid embeddings or this revised manuscript. Figures and wording were corrected locally. The author team must deposit the exact revised source and evidence and confirm redistribution conditions before submission. No raw images or foundation-model weights are bundled.
''',encoding='utf-8')
    # Supply a standalone verifier without project paths or metadata downloads.
    verify=(ROOT/'scripts/csbj_review_verify.py').read_text(encoding='utf-8')
    start=verify.index('def operator(');end=verify.index('\ndef sources():')
    body=verify[start:end].replace('SRC/','BASE/').replace("Q/'primary_probe_reproduction.json'", "BASE/'reproduced_probe_report.json'").replace("Q/'primary_probe_float64_sensitivity.json'", "BASE/'float64_sensitivity.json'")
    body=body.replace("p=json.loads((BASE/'evidence/corrected_probe.json').read_text(encoding='utf-8'))", "p=json.loads((BASE/'evidence/corrected_probe.json').read_text(encoding='utf-8'))")
    portable="from pathlib import Path\nimport json\nimport numpy as np\nfrom scipy.stats import kendalltau\nBASE=Path(__file__).resolve().parent\n"+body+'\nif __name__=="__main__":probes()\n'
    (E/'verify_primary_probes.py').write_text(portable,encoding='utf-8')
    files=[p for p in sorted(OUT.rglob('*')) if p.is_file() and p.name not in ['SHA256SUMS.txt','CSBJ_checked_submission_package.zip']]
    (OUT/'SHA256SUMS.txt').write_text('\n'.join(hashlib.sha256(f.read_bytes()).hexdigest()+'  '+f.relative_to(OUT).as_posix() for f in files)+'\n',encoding='utf-8')
    zpath=SRC/'CSBJ_checked_submission_package_20261014.zip'
    with zipfile.ZipFile(zpath,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for f in OUT.rglob('*'):
            if f.is_file():z.write(f,f.relative_to(OUT))
    print('Package',len(files)+1,'files',round(zpath.stat().st_size/1024**2,1),'MiB',flush=True)

if __name__=='__main__':
    update_author();report()
    if '--package' in sys.argv:package()
