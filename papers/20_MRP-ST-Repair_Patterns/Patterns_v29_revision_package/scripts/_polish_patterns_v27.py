from pathlib import Path
import re,json,sys
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
sys.stdout.reconfigure(encoding='utf-8')
P=Path('papers/20_MRP-ST-Repair_Patterns');f=P/'MRP-ST-Repair_manuscript_Patterns_v27.md';t=f.read_text(encoding='utf-8')
t=t.replace('Saeb, A., Lonini','Saeb, S., Lonini')
t=t.replace('Limitations (iii)','Discussion').replace('(limitation x)','(Discussion)')
t=t.replace('the released software instantiation','the local software instantiation')
t=t.replace('Inflation tracks coupling and nothing else:', 'Inflation differs across label choices:')
t=t.replace('What repair costs exposes the mechanism directly:', 'The descriptive retention comparison is:')
t=t.replace('so removing identity damaged the transferable signal in proportion to how much the label shared with the group, and did not damage it at all where the label carried no group information.', 'This pattern is compatible with coupling-related differences, but changing the label also changes the task, its rank structure and signal strength; coupling alone is not isolated.')
t=t.replace('at kappa = 0 there is no channel to open and no repair to justify','at kappa = 0 the observed sibling-unit difference is much smaller, but is not identically zero')
t=t.replace('and it is a prediction about any cohort, not a description of ours','under this fixed-scale, same-statistic gate convention')
t=t.replace('The boundary exists wherever the argument applies;','The external experiments do not establish the existence of the same boundary;')
t=t.replace('Comparability across cohorts of different size','Sensitivity across cohorts of different size')
# External dependent-label bootstrap and absolute causal interpretations were not validated.
t=t.replace('because these are properties of a cohort rather than of the operator','within these archived configurations')
# Data release paragraph: do not assert public release in an untouched duplicate location.
start=t.index('## DATA AND CODE AVAILABILITY');end=t.index('## ACKNOWLEDGMENTS',start)
t=t[:start]+'''## DATA AND CODE AVAILABILITY

The public source accessions used for the transcriptomic analyses are GSE208654, GSE217237, GSE268014, GSE278936 and GSE316098, together with the stated 10x Genomics example. HAR and Gas Sensor Array Drift are obtained from the UCI repository. Analysis scripts and archived result files are held in the local project. The proposed code repository is https://github.com/WorldBMW/cervix-st-multimodal, with intended release v1.0.0. Both the repository and release returned HTTP 404 in the 25 October 2026 audit. Public code access and an immutable archive remain outstanding; this text must be updated after successful verification. Raw participant data are not newly collected or redistributed by this study. Access restrictions and source licenses must be respected.

'''+t[end:]
f.write_text(t,encoding='utf-8')
# Rebuild deterministically and retain embedded figures.
from _md2docx import convert
convert(str(f),str(P/'Manuscript_Patterns.docx'))
d=Document(P/'Manuscript_Patterns.docx');d.add_page_break();d.add_heading('Figures',1)
for i,n in enumerate(['fig1.png','fig2.png','fig3.png','fig4.png','fig5_crossdomain.png'],1):
 d.add_paragraph(f'Figure {i}');d.add_picture(str(P/'figures'/n),width=Inches(6.3))
 d.add_paragraph('See the corresponding numbered legend. LDI denotes the archived unit-level label-association index; reference crossing does not certify identity removal. The Results interpretation applies to these archived panels.')
 if i<5:d.add_page_break()
for tab in d.tables:
 tab.autofit=True
 tab.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in tab.rows:
  for c in row.cells:
   for p in c.paragraphs:
    for r in p.runs:r.font.size=Pt(8)
d.save(P/'Manuscript_Patterns.docx')
print('Polished v27 and embedded 5 figures')
