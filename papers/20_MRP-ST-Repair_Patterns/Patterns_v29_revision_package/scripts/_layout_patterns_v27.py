from pathlib import Path
import copy,re,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
P=Path('papers/20_MRP-ST-Repair_Patterns')
f=P/'MRP-ST-Repair_manuscript_Patterns_v27.md';t=f.read_text(encoding='utf-8')
a=t.index('**Symbols common');b=t.index('**Figure 1.',a)
t=t[:a]+'''**Symbols and interpretation.** τ is Kendall rank association; τ_in is fitted in-sample, τ_LOO holds out one analysis unit, and τ_LGO or τ_LDO holds out a recorded group. The archived LDI is (τ_LOO − 0.001)/σ̂, where σ̂ is the scale of the unit-level refit permutation reference. It is not a measure of group-identity removal. R = τ_LOO/τ_in is unit-holdout retention; Δ = τ_LOO − τ_LGO is sibling-unit inflation. Direct identity lift is a separate classifier's accuracy divided by 1/G. Lines and error bars summarize training seeds on the same cohort and do not represent independent cohort replication. Figure panels retain their archived experimental configurations; a fixed-scale compatibility calculation is not a prediction that a feasible dose exists.

'''+t[b:]
t=t.replace('6-class ordinal','6-class nominal, numerically coded')
t=t.replace('single NVIDIA GeForce RTX 4080 SUPER, 32 GB','NVIDIA GeForce RTX 4080 SUPER (archived hardware declaration; memory provenance unverified)')
t=t.replace('are released at','are locally recorded with proposed public access at')
t=t.replace('tag v1.0.0); Zenodo archive of the same state, DOI reported in the version of record','intended tag v1.0.0); public access and Zenodo archiving pending verification')
t=t.replace('| Manifest-to-result pipeline (one JSON manifest per run) | this paper | released with the tagged release |','| Manifest-to-result mapping | this paper | completeness not independently verified |')
t=t.replace('**Permutation depth.** σ̂ is estimated at B = 2,000 refit permutations, and the null behind the reported p-value uses the same B.','**Permutation depth.** σ̂ uses B = 2,000 refit permutations. The cross-domain and GeoMx protocol records B = 10,000 for p-values when requested; permutation-depth sensitivity is a separate analysis.')
t=t.replace('Protocol step 9','the practical reporting workflow')
f.write_text(t,encoding='utf-8')
# Preserve source-to-Word synchronization before layout-only adjustments.
exec(Path('scripts/_polish_patterns_v27.py').read_text(encoding='utf-8'))
def black(doc):
 for style in doc.styles:
  if style.type==1 and ('Heading' in style.name or style.name=='Title'):style.font.color.rgb=RGBColor(0,0,0)
def header(tab):
 tab.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for cell in tab.rows[0].cells:
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EEF4');cell._tc.get_or_add_tcPr().append(shade)
def landscape_table(doc,tab):
 portrait=copy.deepcopy(doc.sections[-1]._sectPr)
 def sp(land):
  s=copy.deepcopy(portrait);pg=s.find(qn('w:pgSz'))
  if land:pg.set(qn('w:w'),'16838');pg.set(qn('w:h'),'11906');pg.set(qn('w:orient'),'landscape')
  typ=s.find(qn('w:type'))
  if typ is None:typ=OxmlElement('w:type');s.insert(0,typ)
  typ.set(qn('w:val'),'nextPage');return s
 for before,land in [(True,False),(False,True)]:
  p=OxmlElement('w:p');pr=OxmlElement('w:pPr');pr.append(sp(land));p.append(pr)
  if before:tab._tbl.addprevious(p)
  else:tab._tbl.addnext(p)
 tab.autofit=False
 for row in tab.rows:
  for cell in row.cells:cell.width=Inches(9.4/len(row.cells))
 for col in tab.columns:col.width=Inches(9.4/len(tab.columns))
d=Document(P/'Manuscript_Patterns.docx');black(d)
for i,tab in enumerate(d.tables):
 header(tab)
 if i in [2,3]:pass
 # Table order follows the narrative: Table 3, Table 5, Table 4, Table 6, Table 1, Table 2.
 if len(tab.columns)>=9:landscape_table(d,tab)
d.save(P/'Manuscript_Patterns.docx')
s=Document(P/'Supplemental_Information_Patterns.docx');black(s)
for p in s.paragraphs:
 if p._p.xpath('.//w:drawing'):
  p._p.getparent().remove(p._p)
 elif p.text.startswith('An auditable instrument for identity leakage'):p.text='Auditing label dependence and sibling-unit inflation in group-structured representations'
for sec in s.sections:
 sec.page_width=Inches(11.7);sec.page_height=Inches(8.3)
 sec.left_margin=sec.right_margin=Inches(.65)
 sec.top_margin=sec.bottom_margin=Inches(.6)
for tab in s.tables:
 tab.autofit=False;header(tab)
 for col in tab.columns:col.width=Inches(10.4/len(tab.columns))
 for row in tab.rows:
  for cell in row.cells:
   cell.width=Inches(10.4/len(row.cells))
   for p in cell.paragraphs:
    for r in p.runs:r.font.size=Pt(8)
s.save(P/'Supplemental_Information_Patterns.docx')
print('Applied landscape layouts to wide tables; synchronized statistical legend')
