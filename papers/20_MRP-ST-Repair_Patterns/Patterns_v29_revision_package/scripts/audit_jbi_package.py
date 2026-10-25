"""Structural and content audit for the JBI submission package.

The script is intentionally read-only.  It checks the generated DOCX files,
cross-references, highlights, reference numbering, embedded images, and raster
metadata.  Results are printed as JSON so that the checks are reproducible.
"""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image


ROOT = Path(__file__).resolve().parent.parent
PKG = ROOT / "papers" / "08_JBI"


def all_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def paragraphs_between(doc: Document, start: str, stop: str) -> str:
    capture = False
    out: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text == start:
            capture = True
            continue
        if capture and text == stop:
            break
        if capture and text:
            out.append(text)
    return " ".join(out)


def word_count(text: str) -> int:
    return len(re.findall(r"\b[A-Za-z0-9]+(?:[-'][A-Za-z0-9]+)*\b", text))


def docx_audit(path: Path) -> dict:
    doc = Document(path)
    text = all_text(doc)
    with zipfile.ZipFile(path) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
        names = set(zf.namelist())
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "embedded_media": len(media),
        "sections": len(doc.sections),
        "comments_part": "word/comments.xml" in names,
        "track_change_tags": len(re.findall(r"<w:(?:ins|del)\b", text)),
        "citation_placeholders": re.findall(r"\[@[^\]]+\]", text),
        "pending_markers": sorted(set(re.findall(r"\[(?:pending|PENDING)[^\]]*\]", text))),
        "generic_placeholders": sorted(set(re.findall(r"\[(?:DATE|NAME|AFFILIATION|INSERT|TBD)[^\]]*\]", text, re.I))),
        "replacement_characters": text.count("\ufffd"),
    }


def manuscript_audit(path: Path) -> dict:
    doc = Document(path)
    text = all_text(doc)
    base = docx_audit(path)
    abstract = paragraphs_between(doc, "Abstract", "1. Introduction")
    abstract = abstract.split("Keywords:", 1)[0]
    refs = []
    in_refs = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "References":
            in_refs = True
            continue
        if in_refs and re.match(r"^\[\d+\]", t):
            refs.append(t)
    ref_numbers = [int(re.match(r"^\[(\d+)\]", r).group(1)) for r in refs]
    body_before_refs = text.split("\nReferences\n", 1)[0]
    cited = set()
    for group in re.findall(r"\[(\d+(?:[,–-]\d+)*)\]", body_before_refs):
        group = group.replace("–", "-")
        for token in group.split(","):
            if "-" in token:
                a, b = map(int, token.split("-", 1))
                cited.update(range(a, b + 1))
            else:
                cited.add(int(token))
    significance_words = 0
    for table in doc.tables:
        if table.cell(0, 1).text.strip() == "Statement of Significance":
            significance_words = sum(word_count(table.cell(i, 1).text) for i in range(1, len(table.rows)))
            break
    base.update(
        {
            "abstract_words_including_headings": word_count(abstract),
            "statement_of_significance_words": significance_words,
            "references": len(refs),
            "reference_numbers_sequential": ref_numbers == list(range(1, len(refs) + 1)),
            "uncited_reference_numbers": sorted(set(ref_numbers) - cited),
            "undefined_citation_numbers": sorted(cited - set(ref_numbers)),
            "doi_count": len(re.findall(r"https://doi\.org/\S+", "\n".join(refs))),
            "figure_mentions": {str(i): len(re.findall(fr"\bFigure {i}\b", body_before_refs)) for i in range(1, 6)},
            "table_mentions": {str(i): len(re.findall(fr"\bTable {i}\b", body_before_refs)) for i in range(1, 4)},
            "supplementary_mentions": len(re.findall(r"\bSupplementary (?:Table|Note|Figure)\b", body_before_refs)),
        }
    )
    return base


def highlights_audit(path: Path) -> dict:
    doc = Document(path)
    bullets = [p.text.strip() for p in doc.paragraphs if p.style and p.style.name == "List Bullet"]
    return {
        **docx_audit(path),
        "bullets": [{"text": b, "characters": len(b), "within_85": len(b) <= 85} for b in bullets],
        "count_within_3_to_5": 3 <= len(bullets) <= 5,
    }


def image_audit(path: Path) -> dict:
    with Image.open(path) as image:
        return {
            "pixels": list(image.size),
            "dpi": list(image.info.get("dpi", (0, 0))),
            "mode": image.mode,
        }


def main() -> None:
    figures = {p.name: image_audit(p) for p in sorted((PKG / "Figures").glob("*.png"))}
    out = {
        "manuscript": manuscript_audit(PKG / "Manuscript_JBI.docx"),
        "cover_letter": docx_audit(PKG / "CoverLetter_JBI.docx"),
        "highlights": highlights_audit(PKG / "Highlights_JBI.docx"),
        "supplementary": docx_audit(PKG / "Supplementary_Information_JBI.docx"),
        "graphical_abstract": image_audit(PKG / "GraphicalAbstract_JBI.png"),
        "graphical_abstract_tiff": (image_audit(PKG / "GraphicalAbstract_JBI.tiff")
                                     if (PKG / "GraphicalAbstract_JBI.tiff").exists() else None),
        "figures": figures,
    }
    print(json.dumps(out, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
