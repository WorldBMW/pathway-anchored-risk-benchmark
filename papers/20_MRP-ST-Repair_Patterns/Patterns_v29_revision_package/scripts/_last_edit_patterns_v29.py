from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
P=Path('papers/20_MRP-ST-Repair_Patterns')
src=P/'MRP-ST-Repair_manuscript_Patterns_v29.md';t=src.read_text(encoding='utf-8')
pairs={
 'The cervical cohort is 35 Visium sections from 34 donors':'The cervical archive contains 35 Visium sections assigned to 34 name-derived groups',
 'Two of the sections come from one donor, which is why the section count exceeds the donor count by one.':'The archived grouping rule merges CA_HPV_1 and CA_HPV_2; source patient identifiers are still required to verify whether those sections share a donor.',
}
for a,b in pairs.items():t=t.replace(a,b)
blocks=t.split('\n\n')
for i,b in enumerate(blocks):
 if 'Recomputing the null with a within-group permutation' in b:
  blocks[i]='4. **Declare a task-appropriate null.** Choose a permutation scheme for the actual independent unit and scientific question. Within-group label shuffling preserves group-by-label counts and is only one conditional sensitivity analysis; it can be degenerate and does not provide a universal group-level test. Report empirical null quantities and the conditions needed for inference.'
t='\n\n'.join(blocks);src.write_text(t,encoding='utf-8')
d=Document(P/'Manuscript_Patterns_v29.docx')
for p in d.paragraphs:
 s=p.text
 for a,b in pairs.items():s=s.replace(a,b)
 if 'Recomputing the null with a within-group permutation' in s:
  s='4. Declare a task-appropriate null. Choose a permutation scheme for the actual independent unit and scientific question. Within-group label shuffling preserves group-by-label counts and is only one conditional sensitivity analysis; it can be degenerate and does not provide a universal group-level test. Report empirical null quantities and the conditions needed for inference.'
 if s!=p.text:p.text=s
# Start the landscape section before the complete table caption, eliminating a caption-only page.
for tab in d.tables:
 before=tab._tbl.getprevious()
 if before is not None and before.findall('.//'+qn('w:sectPr')):
  caption=before.getprevious()
  if caption is not None and ''.join(caption.itertext()).startswith('Table 5.'):
   before.getparent().remove(before);caption.addprevious(before)
 for c in tab.rows[0].cells:
  for p in c.paragraphs:p.paragraph_format.keep_with_next=True
d.save(P/'Manuscript_Patterns_v29.docx')
print('Corrected remaining donor-map wording and width-change pagination')
