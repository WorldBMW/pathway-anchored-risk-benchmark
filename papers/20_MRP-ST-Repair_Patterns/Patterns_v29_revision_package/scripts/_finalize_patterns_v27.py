from pathlib import Path
import re,json,sys
from docx import Document
from docx.shared import Inches,Pt
sys.stdout.reconfigure(encoding='utf-8')
P=Path('papers/20_MRP-ST-Repair_Patterns'); f=P/'MRP-ST-Repair_manuscript_Patterns_v27.md'
t=f.read_text(encoding='utf-8')
def section(start,end,new):
 global t
 a=t.index(start); b=t.index(end,a+len(start));t=t[:a]+new.strip()+'\n\n'+t[b:]
section('## GRAPHICAL ABSTRACT','## BIGGER PICTURE','''## GRAPHICAL ABSTRACT

**Graphical abstract.** The audit reports four separate objects: label association under unit holdout, association under group holdout, direct group-identity decoding, and retained utility. LDI uses the archived unit-permutation reference and must not be labeled a group-level leakage certificate. The fixed-scale compatibility bound concerns the same association on both axes. Spatial transcriptomics supplies controls and biological readouts; HAR and gas drift supply exploratory audits only. No downstream benefit was demonstrated.
''')
# Correct the opening Results definitions rather than leave competing definitions.
a=t.index('We formalise the setting'); b=t.index('The cheapest repair',a)
t=t[:a]+'''Groups g = 1 … G contribute repeated analysis units. A representation may retain group-specific features, which a probe can exploit when sibling units remain in its training set. We therefore distinguish τ_in, τ_LOO and τ_LGO as defined in Experimental Procedures. LDI = (τ_LOO − 0.001)/σ̂ and R = τ_LOO/τ_in use the archived unit-level probe; neither alone measures identity removal. Δ = τ_LOO − τ_LGO reports the observed difference between unit and group holdouts.

In the main transcriptomic experiments the encoder acts on spots and the audit generally consumes section means. The head-and-neck data have 16 sections and 16 recorded groups. The cervical archive reports 35 section units and 34 donor groups; its τ_LOO and τ_LGO are therefore not interchangeable. The separate resolution ladder groups by 35 section identifiers. The actual grouping map, not the term “donor” in a label, determines the evaluation target.

'''+t[b:]
section('### A precondition test for the detect-then-accept gate','### Decision boundary I', '''### A fixed-scale compatibility check for the dual gate

For the same post-intervention statistic τ and a fixed σ̂ > 0, the two requirements LDI < 1.645 and τ ≥ cτ_base overlap if and only if cτ_base < 0.001 + 1.645σ̂ (Proposition 1). At c = 0.7 and a negligible null floor, the corresponding ratio threshold is approximately 2.35. This is a necessary compatibility condition for the two numerical requirements, not proof that a particular model or dose reaches the overlap. The result concerns label association rather than group-identity independence.

Table 1 reports baseline ratios and their associated approximate critical retention. Shared-program simulation, cervical, head-and-neck and prostate ratios are approximately 5.8, 2.4, 2.18 and 2.0; HAR and gas drift are 13.8 and 11.0. Because σ̂ is refitted after intervention, a baseline-only exclusion requires a fixed-scale convention or a prespecified upper bound on its drift. The archived external-cohort fixed-versus-refitted comparison changes no dependence-gate decision (Table S14), which is an empirical sensitivity analysis, not a general invariance theorem.

When τ_LOO grades dependence and τ_LGO grades fidelity, the gates involve different statistics and this one-dimensional equivalence does not apply. The archived multi-unit runners use τ_LOO for both gate axes and report group-level retention separately. An empty dual gate can therefore arise from retaining a genuinely strong label program, and should not be interpreted as inability to remove a nuisance identity. Conversely a near-null baseline makes a relative retention requirement uninformative. The spurious simulation arm illustrates this degeneracy (Table S12).
''')
# Remove unsupported cross-cohort causal and universal interpretations while retaining numeric records.
paras=t.split('\n\n');out=[]
for p in paras:
 if p.startswith('**The contrast is the finding.**'):
  p='**Descriptive contrast across cohorts.** The cohorts differ in label semantics, class balance, source and sample size, so their dose responses cannot isolate a causal effect of label–identity coupling. The archived contrasts motivate matched controls rather than establish a general repairability boundary. The section-level index also measures target-label association, not direct group-identity decoding.'
 if p.startswith('**A random-label control tests the observed assignment.**'):
  p='**A random-label control tests the observed assignment.** Shuffling the 35 section-level labels while keeping class counts (31, 1, 3) gives a negative baseline LDI (mean −0.82) and negative values across the recorded doses (range −0.65 to −1.15). This checks sensitivity to label assignment. It breaks associations with both useful signal and nuisance structure and does not isolate a causal identity channel; with unique unit identifiers, a deterministic label-to-unit assignment still exists.'
 if p.startswith('**Decision boundary I reproduced with real data.**'):
  p='**Observed ordering in two datasets.** Identity lift is 10.0× chance in HAR and 9.2× in gas drift. Baseline unit-versus-group inflation is +0.0061 and +0.0494, respectively, alongside coupling κ = 0.0025 and 0.0511. This ordering is consistent with coupling-related inflation, but two datasets with different tasks and distributions cannot identify coupling as its sole cause. The nominal-label rank coding also limits interpretation of the magnitude.'
 if p.startswith('**Proposition 1 held out of domain.**'):
  p='**Exploratory fixed-scale check.** Baseline τ_LOO/σ̂ is 13.8 for HAR and 11.0 for gas drift. Neither dataset has a recorded dual-gate pass among its 12 dosed cells. These gates use unit-level τ_LOO on both axes, and the label association remains strong after intervention. This is consistent with the fixed-scale interval conflict, but does not show that no dose could preserve utility while reducing identity. The independent τ_LGO retention remains at least 0.98 for HAR and 0.89 for gas drift.'
 if p.startswith('**The group-count boundary does not transport numerically.**'):
  p='**No transported group-count threshold is established.** Gas drift has ten batches and retains 0.89–0.93 of its group-level association under the recorded intervention. This differs from the simulation placement at 16–32 groups. No phase grid outside spatial transcriptomics was run, so neither existence nor placement of a comparable external boundary is demonstrated.'
 if p.startswith('**Control: the channel is opened by coupling'):
  p=p.replace('the channel is opened by coupling','label choice and sibling-unit inflation')
 p=p.replace('The contrast is causal','The contrast is descriptive')
 p=p.replace('Repair does not extinguish real signal; it extinguishes group identity.','In these positive-control configurations the intervention preserved the planted signal; identity suppression must be assessed separately.')
 p=p.replace('The baseline model leaked (mean LDI = 2.04, above the gate)','The baseline had label association above the reference (mean LDI = 2.04)')
 p=p.replace('The ratio governs whether a feasible dose exists','The ratio governs fixed-scale interval compatibility')
 p=p.replace('non-empty for every c < c_crit, empty for every c ≥ c_crit','overlapping at the fixed reference scale for c < c_crit and non-overlapping for c ≥ c_crit')
 p=p.replace('the scale of the refit permutation null; the acceptance gate is','the scale of the unit-level refit permutation reference; the recorded dependence gate is')
 p=p.replace('LDI = (τ_LGO − τ_⊥)/σ̂','LDI = (τ_LOO − τ_⊥)/σ̂')
 out.append(p)
t='\n\n'.join(out)
# Remove incorrect general fairness interpretation from the proposition remarks.
a=t.index('**How the conflict scales (Remark 1).**');b=t.index('**Protocol.**',a)
t=t[:a]+'''**Scope of the compatibility calculation.** With a fixed positive baseline and a decreasing null scale, the upper interval bound can fall below the retained-association requirement. The calculation does not identify a causal fairness–utility trade-off. A non-empty interval need not be reachable; a moving post-intervention null or different statistics on the two axes requires a different calculation. The approximate c_crit = 1.645 σ̂/τ_base is retained in tables for correspondence with the archive, while the exact value includes the 0.001 floor.

'''+t[b:]
t=t.replace('adopt the three-axis protocol instead of searching for a dose','report the separate dependence, identity-decoding and utility axes rather than interpreting a threshold as successful removal')
t=t.replace('no accepted repair improved a downstream endpoint where no leakage channel exists','no accepted repair demonstrated a downstream improvement')
t=t.replace('whose grouping variable is the donor','whose recorded grouping variable must be checked against source identifiers')
f.write_text(t,encoding='utf-8')
from _md2docx import convert
convert(str(f),str(P/'Manuscript_Patterns.docx'))
d=Document(P/'Manuscript_Patterns.docx')
# Embed submission figures for a self-contained reviewer copy; original figure files are retained.
d.add_page_break();d.add_heading('Figures',level=1)
for i,name in enumerate(['fig1.png','fig2.png','fig3.png','fig4.png','fig5_crossdomain.png'],1):
 d.add_paragraph(f'Figure {i}')
 d.add_picture(str(P/'figures'/name),width=Inches(6.3))
 d.add_paragraph('See the numbered legend and the archived-results interpretation in Results. LDI labels refer to the implemented unit-level label-association index; they do not certify group-identity removal.')
 if i<5:d.add_page_break()
# Repeat table headers and allow automatic row height rather than clipping.
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
for tab in d.tables:
 tab.autofit=True
 for row in tab.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for run in p.runs:run.font.size=Pt(8)
 hdr=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(hdr)
d.save(P/'Manuscript_Patterns.docx')
# Fix the stale reproduction instructions and avoid claiming a release already exists.
r=Path('REPRODUCE.md');s=r.read_text(encoding='utf-8')
s=s.replace('MRP-ST-Repair_manuscript_Patterns_v15.md','MRP-ST-Repair_manuscript_Patterns_v27.md').replace('scripts/_qa_v15.py','scripts/_qa_patterns_v27.py')
s=s.replace('10x Genomics, Human Cervical Cancer FFPE |','10x Genomics plus GEO GSE208654 and GSE217237 |')
s=s.replace('Released from the tagged GitHub release, archived at Zenodo (DOI in the\nmanuscript\'s Data and code availability statement).','Public release and immutable archiving remain to be verified before submission.')
s=s.replace('Everything reported in the paper is regenerated from public data by the scripts\nlisted below.','The scripts and archived results below support reconstruction of the analyses. Full raw-data reproduction has not been verified in this editorial audit.')
s=s.replace('CPU only (no GPU run is needed for any reported number).','The archived key resources table records GPU hardware for some training runs; exact hardware and dependency provenance require verification.')
s+='\n\n## Interpretation corrected in v27\n\nLDI uses unit-level tau_LOO and a unit-label permutation reference. tau_LGO is separate. The multi-unit dual gate in exp_crossdomain_audit.py uses tau_LOO for both axes; it does not certify group-level acceptance. Group-preserving permutations and class-aware metrics are needed before inferential claims for clustered nominal-label data. The cervical resolution ladder uses 35 section groups; the main cervical archive reports 34 donor groups. GeoMx has 7 specimens from 6 patients, with 186 compartment measurements paired within 93 ROIs.\n'
r.write_text(s,encoding='utf-8')
print('Finalized manuscript and reproduction notes')
